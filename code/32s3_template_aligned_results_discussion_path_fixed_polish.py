# -*- coding: utf-8 -*-
"""
32S3 Template-aligned Results and Discussion path-fixed polish DOCX generator
Project: fluoride coagulation defluoridation ML project

Purpose
-------
Use the real Word manuscript template (AI论文--混凝沉淀除氟空稿.docx) and replace the
placeholder Results and Discussion section with a cleaned, formal manuscript-style
section based on the locked 32M2/32Q2 selective-AD result.

This script does NOT retrain models and does NOT change final metrics.
It also does NOT create a new output root. Outputs follow the original project convention:
    outputs/tables/
    reports/

Important writing boundary
--------------------------
R²=0.747676 belongs to the target-free selective AD retained-subdomain result only.
It must not be reported as all-domain grouped-validation performance.

Expected inputs
---------------
1) AI论文--混凝沉淀除氟空稿.docx
2) outputs/tables/32m2_corrected_key_result_table.csv  (preferred)
   or outputs/tables/32j_main_result_table.csv          (fallback)
3) outputs/figures/194_32i2_evidence_chain_summary_clean.png
   outputs/figures/195_32i2_threshold_plateau_evidence_clean.png
   outputs/figures/196_32i2_retained_chemistry_window_clean.png
   outputs/figures/197_32i2_stability_transfer_boundary_clean.png

Outputs
-------
reports/32s3_results_discussion_template_aligned_path_fixed_polished.docx
reports/32s3_results_discussion_template_aligned_path_fixed_polished_text.txt
reports/32s3_run_log.txt
outputs/tables/32s3_template_aligned_docx_manifest.csv
outputs/tables/32s3_final_text_boundary_check.csv
outputs/tables/32s3_figure_insertion_manifest.csv
outputs/tables/32s3_path_resolution_audit.csv
"""

from __future__ import annotations

import csv
import os
import re
import sys
import copy
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Optional, Tuple

try:
    import pandas as pd
except Exception:
    pd = None

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml.ns import qn
except Exception as e:
    print("[ERROR] python-docx is required. Install it by:")
    print(r"    .\.venv\Scripts\python.exe -m pip install python-docx")
    print("Original error:", e)
    sys.exit(1)


SCRIPT_NAME = "32s3_template_aligned_results_discussion_path_fixed_polish.py"
TEMPLATE_NAME = "AI论文--混凝沉淀除氟空稿.docx"

ACCEPTED_DEFAULT = {
    "result_type": "target-free selective AD retained-subdomain",
    "model": "RepeatedBalanced_Extended + XGBoost",
    "rule": "Zone=Outside; NN_Risk <= 0.814815",
    "r2": 0.747676,
    "rmse": 11.474674,
    "mae": 9.155856,
    "coverage": 0.014664,
    "n": 52,
    "dominant_group": "Paper_010",
}

FIGURES = [
    {
        "figure_no": "图 3-1",
        "source_id": "194",
        "filename": "194_32i2_evidence_chain_summary_clean.png",
        "caption": "选择性适应域保留子域结果的证据链汇总。该图比较了未筛选分组验证、预定义可靠域和选择性适应域保留子域在 R²、RMSE、MAE 和覆盖率上的差异，用于说明最终高 R² 结果属于选择性保留子域，而不是未筛选测试集的整体性能。",
    },
    {
        "figure_no": "图 3-2",
        "source_id": "195",
        "filename": "195_32i2_threshold_plateau_evidence_clean.png",
        "caption": "选择性适应域规则的阈值平台证据。虚线位置对应最终接受的 NN_Risk 阈值，表明该结果是在阈值敏感性分析中形成的高性能平台附近获得，而不是单一孤立阈值造成的偶然峰值。",
    },
    {
        "figure_no": "图 3-3",
        "source_id": "196",
        "filename": "196_32i2_retained_chemistry_window_clean.png",
        "caption": "选择性适应域保留子域的化学条件窗口与适应域组成。该图描述保留样本的关键水质、投加量、pH、干扰离子、混凝剂组成和适应域风险水平，用于界定模型解释适用的化学空间。",
    },
    {
        "figure_no": "图 3-4",
        "source_id": "197",
        "filename": "197_32i2_stability_transfer_boundary_clean.png",
        "caption": "选择性适应域结果的稳定性、迁移支持和解释边界。该图显示不同数据划分和相关模型组合下的保留子域表现，并强调最终结果具有明确的文献来源和覆盖范围限制。",
    },
]

FORBIDDEN_FORMAL_TEXT = [
    "本文件由",
    "脚本",
    "自动生成",
    "正文草稿",
    "建议插入",
    "图件声明边界",
    "图件插入区",
    "结果与讨论图件",
    "not_found_create_new_docx",
]


# ----------------------------- path helpers -----------------------------

def _has_any_figure(fig_dir: Path) -> bool:
    if not fig_dir.exists():
        return False
    for prefix in ["194", "195", "196", "197"]:
        if any(fig_dir.glob(f"{prefix}*.png")) or any(fig_dir.glob(f"{prefix}*.jpg")) or any(fig_dir.glob(f"{prefix}*.jpeg")):
            return True
    return False


def _candidate_score(p: Path) -> Tuple[int, int, str]:
    """Score a possible project root. Higher is better; deeper path wins ties.

    32S2 accidentally selected the download folder because it preferred the Word
    template location. 32S3 separates template discovery from project-root
    discovery and chooses the directory that best matches the real PyCharm
    project structure: src/, outputs/tables/, outputs/figures/, reports/, and
    32-series result files.
    """
    score = 0
    if (p / "src").is_dir():
        score += 80
    if (p / "outputs" / "tables").is_dir():
        score += 50
    if (p / "outputs" / "figures").is_dir():
        score += 50
    if (p / "reports").is_dir():
        score += 30
    if (p / "data").is_dir():
        score += 15
    # Strong evidence from our project files
    for fname in [
        "32j_main_result_table.csv",
        "32m2_corrected_key_result_table.csv",
        "32n_locked_key_result_for_word.csv",
        "32q2_locked_key_result_check.csv",
    ]:
        if (p / "outputs" / "tables" / fname).exists():
            score += 35
    if _has_any_figure(p / "outputs" / "figures"):
        score += 40
    if any((p / "src").glob("32*.py")) if (p / "src").is_dir() else False:
        score += 25
    # Template is useful but should not dominate root detection.
    if (p / TEMPLATE_NAME).exists():
        score += 10
    return score, len(p.parts), str(p)


def find_project_root() -> Path:
    """Detect the real project root without being misled by template location."""
    seeds: List[Path] = []
    if "__file__" in globals():
        here = Path(__file__).resolve()
        seeds.extend([here.parent, here.parent.parent])
    cwd = Path.cwd().resolve()
    seeds.extend([cwd, cwd.parent])

    candidates: List[Path] = []
    def add(p: Path):
        try:
            p = p.resolve()
        except Exception:
            return
        if p.exists() and p.is_dir() and p not in candidates:
            candidates.append(p)

    for seed in seeds:
        add(seed)
        for parent in list(seed.parents)[:5]:
            add(parent)
        # Also inspect one-level child folders because the project may be inside
        # a download directory, e.g. D:/浏览器下载/fluoride_coagulation_ml_project/fluoride_coagulation_ml_project.
        try:
            for child in seed.iterdir():
                if child.is_dir() and child.name not in {".venv", "venv", "__pycache__", "catboost_info", ".git"}:
                    add(child)
                    # one extra level for nested duplicate project folder names
                    try:
                        for gchild in child.iterdir():
                            if gchild.is_dir() and gchild.name not in {".venv", "venv", "__pycache__", "catboost_info", ".git"}:
                                add(gchild)
                    except Exception:
                        pass
        except Exception:
            pass

    scored = [(p, _candidate_score(p)) for p in candidates]
    scored = [x for x in scored if x[1][0] > 0]
    if not scored:
        return cwd
    scored.sort(key=lambda x: x[1], reverse=True)
    return scored[0][0]


def build_path_resolution_audit(root: Path) -> List[Dict[str, object]]:
    rows: List[Dict[str, object]] = []
    candidates: List[Path] = []
    if "__file__" in globals():
        here = Path(__file__).resolve()
        candidates.extend([here.parent, here.parent.parent])
    cwd = Path.cwd().resolve()
    candidates.extend([cwd, cwd.parent])
    for seed in list(candidates):
        for parent in list(seed.parents)[:5]:
            candidates.append(parent)
    unique: List[Path] = []
    for p in candidates:
        try:
            p = p.resolve()
        except Exception:
            continue
        if p.exists() and p.is_dir() and p not in unique:
            unique.append(p)
    for p in unique:
        score, depth, _ = _candidate_score(p)
        rows.append({
            "candidate_path": str(p),
            "score": score,
            "depth": depth,
            "selected": str(p) == str(root),
            "has_src": (p / "src").is_dir(),
            "has_outputs_tables": (p / "outputs" / "tables").is_dir(),
            "has_outputs_figures": (p / "outputs" / "figures").is_dir(),
            "has_reports": (p / "reports").is_dir(),
            "has_template_here": (p / TEMPLATE_NAME).exists(),
        })
    rows.sort(key=lambda r: (int(r["score"]), int(r["depth"])), reverse=True)
    return rows


def ensure_dirs(root: Path) -> Tuple[Path, Path, Path]:
    tables = root / "outputs" / "tables"
    figures = root / "outputs" / "figures"
    reports = root / "reports"
    tables.mkdir(parents=True, exist_ok=True)
    figures.mkdir(parents=True, exist_ok=True)
    reports.mkdir(parents=True, exist_ok=True)
    return tables, figures, reports


def find_file(root: Path, filename: str, preferred_dirs: Optional[List[Path]] = None) -> Optional[Path]:
    preferred_dirs = preferred_dirs or []
    search_dirs = []
    for d in preferred_dirs + [root, root / "reports", root / "outputs" / "tables", root / "outputs" / "figures", root / "data", root / "data" / "raw", root / "data" / "processed", root.parent, Path.cwd()]:
        if d not in search_dirs:
            search_dirs.append(d)
    for d in search_dirs:
        p = d / filename
        if p.exists():
            return p
    # Deliberately avoid recursive scans: project folders may contain large archives or virtual envs.
    return None


# ----------------------------- reading helpers -----------------------------

def read_csv_any(path: Path) -> List[Dict[str, str]]:
    if not path or not path.exists():
        return []
    encodings = ["utf-8-sig", "utf-8", "gbk", "gb18030"]
    for enc in encodings:
        try:
            with path.open("r", encoding=enc, newline="") as f:
                return list(csv.DictReader(f))
        except Exception:
            continue
    return []


def first_existing(paths: List[Optional[Path]]) -> Optional[Path]:
    for p in paths:
        if p and p.exists():
            return p
    return None


def parse_float(value, default=None):
    try:
        if value is None:
            return default
        s = str(value).strip().replace("%", "")
        if s == "" or s.lower() in {"nan", "none", "na"}:
            return default
        return float(s)
    except Exception:
        return default


def load_locked_result(root: Path, tables_dir: Path) -> Tuple[Dict, Optional[Path], str]:
    """Prefer 32m2 corrected key result table; fallback to 32j; final fallback constants."""
    p_m2 = find_file(root, "32m2_corrected_key_result_table.csv", [tables_dir])
    p_32j = find_file(root, "32j_main_result_table.csv", [tables_dir])

    # 32m2 is already corrected and preferred.
    if p_m2:
        rows = read_csv_any(p_m2)
        if rows:
            row = rows[0]
            result = ACCEPTED_DEFAULT.copy()
            # flexible keys
            lower = {str(k).strip().lower(): v for k, v in row.items()}
            def get_any(names, default=None):
                for n in names:
                    if n.lower() in lower:
                        return lower[n.lower()]
                return default
            result["model"] = str(get_any(["model", "model_family", "accepted_model", "method"], result["model"])).strip() or result["model"]
            if result["model"] == "XGBoost":
                result["model"] = ACCEPTED_DEFAULT["model"]
            result["rule"] = str(get_any(["rule", "accepted_rule", "selection_rule"], result["rule"])).strip() or result["rule"]
            result["r2"] = parse_float(get_any(["r2", "R2", "R²"], result["r2"]), result["r2"])
            result["rmse"] = parse_float(get_any(["rmse", "RMSE"], result["rmse"]), result["rmse"])
            result["mae"] = parse_float(get_any(["mae", "MAE"], result["mae"]), result["mae"])
            result["coverage"] = parse_float(get_any(["coverage", "Coverage"], result["coverage"]), result["coverage"])
            result["n"] = int(parse_float(get_any(["n", "N", "n_retained", "retained_n", "N retained"], result["n"]), result["n"]))
            result["dominant_group"] = str(get_any(["dominant_group", "Dominant group", "group"], result["dominant_group"])).strip() or result["dominant_group"]
            return result, p_m2, "selected_from_32m2_corrected_key_result_table"

    if p_32j:
        rows = read_csv_any(p_32j)
        # find row with selective retained subdomain; avoid all-domain row
        best = None
        for row in rows:
            text = " ".join(str(v) for v in row.values()).lower()
            if "selective" in text and "retained" in text:
                best = row
                break
        if best is None and rows:
            best = rows[-1]
        if best:
            result = ACCEPTED_DEFAULT.copy()
            lower = {str(k).strip().lower(): v for k, v in best.items()}
            def get_any(names, default=None):
                for n in names:
                    if n.lower() in lower:
                        return lower[n.lower()]
                return default
            result["model"] = str(get_any(["model", "model_family", "method"], result["model"])).strip() or result["model"]
            if result["model"] == "XGBoost":
                result["model"] = ACCEPTED_DEFAULT["model"]
            result["rule"] = str(get_any(["rule", "selection_rule"], result["rule"])).strip() or result["rule"]
            result["r2"] = parse_float(get_any(["r2", "R2", "R²"], result["r2"]), result["r2"])
            result["rmse"] = parse_float(get_any(["rmse", "RMSE"], result["rmse"]), result["rmse"])
            result["mae"] = parse_float(get_any(["mae", "MAE"], result["mae"]), result["mae"])
            result["coverage"] = parse_float(get_any(["coverage", "Coverage"], result["coverage"]), result["coverage"])
            result["n"] = int(parse_float(get_any(["n", "N", "n_retained", "retained_n"], result["n"]), result["n"]))
            result["dominant_group"] = str(get_any(["dominant_group", "Dominant group", "group"], result["dominant_group"])).strip() or result["dominant_group"]
            return result, p_32j, "selected_from_32j_main_result_table"

    return ACCEPTED_DEFAULT.copy(), None, "fallback_locked_constants"


def safe_rglob(base: Path, patterns: List[str]) -> List[Path]:
    """Small recursive search that skips virtual environments and large archive-like folders."""
    hits: List[Path] = []
    if not base or not base.exists() or not base.is_dir():
        return hits
    skip_names = {".venv", "venv", "__pycache__", "catboost_info", ".git"}
    try:
        for pat in patterns:
            for p in base.rglob(pat):
                parts = set(p.parts)
                if parts & skip_names:
                    continue
                if p.is_file():
                    hits.append(p)
    except Exception:
        pass
    # de-duplicate while keeping order
    out: List[Path] = []
    seen = set()
    for h in hits:
        key = str(h.resolve())
        if key not in seen:
            seen.add(key)
            out.append(h)
    return out


def locate_single_figure(root: Path, figures_dir: Path, item: Dict) -> Optional[Path]:
    """Robust figure locator for Windows/PyCharm project layouts.

    It first tries the exact filename in outputs/figures. If not found, it searches
    common locations and then falls back to prefix/keyword matching. This avoids
    the previous failure where an empty image path became Path('') == '.', causing
    PermissionError when python-docx tried to open the project directory.
    """
    filename = item["filename"]
    source_id = str(item["source_id"])

    # 1) exact search in normal locations
    p = find_file(root, filename, [figures_dir, root / "outputs" / "figures", root, Path.cwd()])
    if p and p.is_file():
        return p

    # 2) exact recursive search, prioritizing outputs/figures then project root
    for base in [figures_dir, root / "outputs", root, root.parent, Path.cwd()]:
        for hit in safe_rglob(base, [filename]):
            if hit.is_file():
                return hit

    # 3) robust prefix search: 194*.png / 195*.png etc.
    candidates: List[Path] = []
    for base in [figures_dir, root / "outputs", root, root.parent, Path.cwd()]:
        candidates.extend(safe_rglob(base, [f"{source_id}*.png", f"{source_id}*.jpg", f"{source_id}*.jpeg"]))

    # 4) keyword scoring to choose the right one if several exist
    keywords_by_id = {
        "194": ["evidence", "chain", "summary"],
        "195": ["threshold", "plateau"],
        "196": ["chemistry", "window"],
        "197": ["stability", "transfer", "boundary"],
    }
    keywords = keywords_by_id.get(source_id, [])
    if candidates:
        def score(path: Path) -> Tuple[int, int, str]:
            name = path.name.lower()
            kw_score = sum(1 for kw in keywords if kw in name)
            clean_score = 1 if "clean" in name else 0
            return (kw_score, clean_score, str(path))
        candidates = sorted(candidates, key=score, reverse=True)
        return candidates[0]

    return None


def locate_figures(root: Path, figures_dir: Path) -> List[Dict]:
    found = []
    for item in FIGURES:
        p = locate_single_figure(root, figures_dir, item)
        rec = item.copy()
        rec["path"] = str(p) if p and p.is_file() else ""
        rec["exists"] = bool(p and p.is_file())
        found.append(rec)
    return found


# ----------------------------- docx helpers -----------------------------

def set_chinese_font(run, size_pt: Optional[float] = None, bold: Optional[bool] = None):
    run.font.name = "Times New Roman"
    try:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    except Exception:
        pass
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def set_paragraph_font(paragraph, size_pt: Optional[float] = None, bold: Optional[bool] = None):
    for run in paragraph.runs:
        set_chinese_font(run, size_pt=size_pt, bold=bold)


def add_paragraph_with_style(doc: Document, text: str = "", style: Optional[str] = None, size_pt: Optional[float] = None, bold: Optional[bool] = None, alignment=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if text:
        r = p.add_run(text)
        set_chinese_font(r, size_pt=size_pt, bold=bold)
    if alignment is not None:
        p.alignment = alignment
    return p


def add_heading_safe(doc: Document, text: str, level: int = 1):
    try:
        p = doc.add_heading(text, level=level)
    except Exception:
        p = add_paragraph_with_style(doc, text, size_pt=14 if level <= 2 else 12, bold=True)
    set_paragraph_font(p, size_pt=14 if level <= 2 else 12, bold=True)
    return p


def add_key_result_table(doc: Document, result: Dict):
    rows = [
        ("结果类型", "无目标泄露的选择性适应域保留子域"),
        ("模型", result["model"]),
        ("筛选规则", result["rule"]),
        ("R²", f"{result['r2']:.6f}"),
        ("RMSE", f"{result['rmse']:.6f}"),
        ("MAE", f"{result['mae']:.6f}"),
        ("Coverage", f"{result['coverage']:.6f}（约 {result['coverage'] * 100:.2f}%）"),
        ("保留样本数", str(result["n"])),
        ("主导文献组", result["dominant_group"]),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "项目"
    hdr[1].text = "内容"
    for cell in hdr:
        for p in cell.paragraphs:
            set_paragraph_font(p, size_pt=10.5, bold=True)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for k, v in rows:
        row = table.add_row().cells
        row[0].text = k
        row[1].text = v
        for cell in row:
            for p in cell.paragraphs:
                set_paragraph_font(p, size_pt=10.5)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table


def add_figure(doc: Document, fig: Dict, width_inches: float = 6.2):
    path_str = str(fig.get("path", "") or "").strip()
    path = Path(path_str) if path_str else None
    if path is not None and path.exists() and path.is_file():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=Inches(width_inches))
    else:
        # Do not call add_picture with an empty path. Path("") equals the current
        # directory on Windows/Python and may raise PermissionError [Errno 13].
        add_paragraph_with_style(doc, f"[缺失图件：{fig['filename']}，请检查 outputs/figures/]", size_pt=10.5, bold=True)
    cap = add_paragraph_with_style(doc, f"{fig['figure_no']} {fig['caption']}", size_pt=10.5)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_results_section_doc(result: Dict, figures: List[Dict]) -> Document:
    """Build a temporary docx containing cleaned formal Results and Discussion only."""
    doc = Document()
    add_heading_safe(doc, "3 Results and Discussion（结果与讨论）", level=1)

    add_heading_safe(doc, "3.1 数据集泛化验证与适应域分析的必要性", level=2)
    add_paragraph_with_style(
        doc,
        "为评价机器学习模型在文献异质性条件下的真实泛化能力，本研究在常规随机划分之外，进一步采用按文献来源分组的验证策略。与随机划分相比，分组验证能够避免同一文献或相近实验系列同时进入训练集和测试集，从而更严格地检验模型对未见文献工况的预测能力。结果表明，未筛选测试条件下的模型性能明显低于随机划分结果，说明混凝沉淀除氟数据不仅存在非线性和多因素耦合特征，还受到文献来源、实验体系和局部工况分布差异的显著影响。因此，仅报告随机划分下的模型精度不足以说明模型具有可靠的工程外推能力，仍需结合适应域分析限定模型的可用范围。",
        size_pt=10.5,
    )

    add_heading_safe(doc, "3.2 预定义可靠域与选择性适应域的对比", level=2)
    add_paragraph_with_style(
        doc,
        "在预定义可靠域分析中，模型尝试通过固定适应域规则筛除高风险样本，以获得更可信的预测子集。然而，固定可靠域并未带来稳定的性能提升，说明单一、静态的适应域阈值难以充分处理本数据集中复杂的跨文献异质性和工况偏移。进一步的选择性适应域分析显示，当模型仅在满足特定风险约束的保留子域内进行评价和解释时，预测误差显著降低，决定系数明显提高。这表明适应域分析的关键并不是简单扩大或缩小预测范围，而是识别模型在当前数据结构下真正具有相对可靠性的局部化学工况空间。",
        size_pt=10.5,
    )

    add_heading_safe(doc, "3.3 选择性适应域保留子域的模型性能", level=2)
    add_paragraph_with_style(
        doc,
        f"最终接受的选择性适应域结果来自 {result['model']}，筛选规则为 {result['rule']}。在该保留子域内，模型取得 R² = {result['r2']:.3f}、RMSE = {result['rmse']:.3f}、MAE = {result['mae']:.3f}，覆盖率为 {result['coverage'] * 100:.2f}%，保留样本数为 {result['n']}。",
        size_pt=10.5,
    )
    add_paragraph_with_style(
        doc,
        "该结果仅表征适应域规则筛选后保留样本子集的预测表现；未筛选测试集结果与筛选后保留子域结果应分开报告，二者不能混用。",
        size_pt=10.5,
    )
    add_key_result_table(doc, result)
    add_figure(doc, figures[0])

    add_heading_safe(doc, "3.4 NN-risk 阈值平台与规则稳定性", level=2)
    add_paragraph_with_style(
        doc,
        "为降低选择性适应域规则被解释为事后挑选结果的风险，本研究进一步对 NN-risk 阈值进行了敏感性分析。结果显示，在接受阈值附近，保留子域 R² 存在相对稳定的高性能区间，而不是由单一孤立阈值产生的偶然峰值。同时，随着阈值放宽，保留样本数和覆盖率逐渐升高，但模型精度并未持续提高，说明更多样本并不必然带来更可靠的外推结果。该现象支持了选择性适应域的核心逻辑：在高度异质的环境实验数据中，模型可靠性需要通过风险约束来界定，而不能简单追求覆盖率最大化。",
        size_pt=10.5,
    )
    add_figure(doc, figures[1])

    add_heading_safe(doc, "3.5 保留子域的化学工况特征与解释边界", level=2)
    add_paragraph_with_style(
        doc,
        f"保留子域样本主要来自 {result['dominant_group']}，并集中于特定混凝剂组成、pH、投加量、初始氟浓度和干扰离子浓度范围内。该结果一方面表明模型在该类化学窗口内能够获得较高预测精度，另一方面也提示模型解释具有明确边界。对于该子域以外的混凝剂类型、极端 pH、不同初始氟负荷或复杂共存离子体系，模型预测仍需谨慎使用。因此，本研究的适应域结果更适合被理解为对特定文献支持工况的可靠识别，而不是对所有混凝沉淀除氟体系的普适性规律归纳。",
        size_pt=10.5,
    )
    add_figure(doc, figures[2])

    add_heading_safe(doc, "3.6 稳定性、迁移支持与应用边界", level=2)
    add_paragraph_with_style(
        doc,
        "稳定性分析进一步表明，选择性适应域结果在部分数据划分和相关模型组合中具有可重复支持，但其覆盖率较低且文献来源集中，因此仍应将其定位为保留子域内的可靠预测证据，不能作为未筛选测试集整体泛化能力的直接证据。在论文表述中，应避免将该保留子域结果直接描述为整个混凝除氟数据集的泛化性能；更严谨的表述是：在无目标泄露的选择性适应域筛选后，模型在保留子域内表现出较高预测精度。该边界声明能够保留结果亮点，同时避免因过度外推导致结论不可靠。",
        size_pt=10.5,
    )
    add_figure(doc, figures[3])

    add_heading_safe(doc, "3.7 本节小结", level=2)
    add_paragraph_with_style(
        doc,
        "综上，混凝沉淀除氟文献数据在跨来源验证中表现出明显的泛化难度，说明该体系存在强烈的数据异质性。通过引入无目标泄露的适应域筛选，本研究识别出一个具有较高预测可靠性的局部保留子域，并在该子域内获得较好的预测结果。该发现为后续工艺参数推荐和模型辅助决策提供了更加谨慎的依据：模型可用于支持特定工况窗口内的预测和解释，但对于适应域之外的水质条件和药剂体系，仍需要补充实验或更多文献数据进行验证。",
        size_pt=10.5,
    )
    return doc


def find_section_paragraph_indices(doc: Document) -> Tuple[Optional[int], Optional[int]]:
    start = None
    end = None
    for i, p in enumerate(doc.paragraphs):
        t = normalize_text(p.text)
        if start is None:
            # match both "3 Results..." and "3. Results..."
            if re.match(r"^3\s*\.?\s*Results\s+and\s+Discussion", t, re.I) or "3 Results and Discussion" in t or "3. Results and Discussion" in t or "3结果与讨论" in t.replace(" ", ""):
                start = i
        elif end is None:
            if re.match(r"^4\s*\.?\s*(Conclusion|结论)", t, re.I) or t.startswith("4. Conclusion") or t.startswith("4 Conclusion") or t.startswith("4. 结论") or t.startswith("4 结论"):
                end = i
                break
    return start, end


def normalize_text(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def replace_results_section(template_path: Path, output_path: Path, section_doc: Document) -> Tuple[bool, str]:
    doc = Document(str(template_path))
    start_idx, end_idx = find_section_paragraph_indices(doc)
    if start_idx is None:
        return False, "Cannot find Results and Discussion section heading in template."
    if end_idx is None:
        return False, "Cannot find Conclusion section heading in template."

    p_elems = [p._p for p in doc.paragraphs]
    start_elem = p_elems[start_idx]
    end_elem = p_elems[end_idx]
    body = doc.element.body

    children = list(body)
    try:
        start_pos = children.index(start_elem)
        end_pos = children.index(end_elem)
    except ValueError:
        return False, "Cannot map section paragraph to body XML elements."

    # Remove old Results and Discussion placeholder content, preserving Conclusion and later content.
    # Remove in reverse order to keep indices valid.
    for elem in children[start_pos:end_pos][::-1]:
        body.remove(elem)

    # Insert new section XML elements immediately before Conclusion.
    insert_before = end_elem
    new_elems = []
    for elem in section_doc.element.body:
        if elem.tag.endswith("}sectPr"):
            continue
        new_elems.append(copy.deepcopy(elem))
    for elem in new_elems:
        insert_before.addprevious(elem)

    # Normalize base style fonts where possible.
    try:
        style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(10.5)
    except Exception:
        pass

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return True, "replaced_template_results_section"


def extract_docx_text(path: Path) -> str:
    try:
        doc = Document(str(path))
        texts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                texts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(texts)
    except Exception:
        return ""


def audit_text(text: str, result: Dict) -> List[Dict[str, str]]:
    rows = []
    checks = [
        ("contains_full_model_name", result["model"] in text, result["model"]),
        ("contains_rule", result["rule"] in text, result["rule"]),
        ("contains_accepted_r2", f"{result['r2']:.3f}" in text or f"{result['r2']:.6f}" in text, f"{result['r2']:.3f} or {result['r2']:.6f}"),
        ("contains_retained_n", str(result["n"]) in text, str(result["n"])),
        ("contains_dominant_group", result["dominant_group"] in text, result["dominant_group"]),
    ]
    for name, ok, expected in checks:
        rows.append({"check": name, "status": "PASS" if ok else "FAIL", "expected": expected})
    for phrase in FORBIDDEN_FORMAL_TEXT:
        rows.append({"check": f"forbidden_phrase_absent::{phrase}", "status": "PASS" if phrase not in text else "FAIL", "expected": "absent"})
    # Avoid a sentence/paragraph that mixes exact R2=0.747676 with all-domain wording.
    mixed = False
    for para in re.split(r"[\n\r]+", text):
        if ("0.747676" in para or "0.748" in para) and ("全域分组验证" in para or "未筛选测试集" in para or "整体性能" in para):
            # allow wording that explicitly says cannot be mixed? In final manuscript, avoid such paragraph entirely.
            mixed = True
            break
    rows.append({"check": "no_exact_high_r2_mixed_with_all_domain_wording", "status": "PASS" if not mixed else "WARN", "expected": "avoid mixing in one paragraph"})
    return rows


def write_csv(path: Path, rows: List[Dict[str, object]]):
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8-sig")
        return
    keys = list(rows[0].keys())
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=keys)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


# ----------------------------- main -----------------------------

def main():
    root = find_project_root()
    tables_dir, figures_dir, reports_dir = ensure_dirs(root)
    path_audit_rows = build_path_resolution_audit(root)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    log: List[str] = []

    def log_print(s=""):
        print(s)
        log.append(str(s))

    log_print("=" * 68)
    log_print("32S3 Template-aligned Results and Discussion path-fixed polish DOCX")
    log_print(f"Run time: {now}")
    log_print(f"Project root: {root}")
    log_print(f"Table output dir: {tables_dir}")
    log_print(f"Figure dir: {figures_dir}")
    log_print(f"Report output dir: {reports_dir}")
    log_print("=" * 68)

    template_path = find_file(root, TEMPLATE_NAME, [root, reports_dir, root / "data", root.parent])
    if not template_path:
        log_print("[ERROR] Real Word template not found.")
        log_print(f"Please put {TEMPLATE_NAME} in the project root or reports/ and run again.")
        (reports_dir / "32s3_run_log.txt").write_text("\n".join(log), encoding="utf-8")
        sys.exit(1)

    result, result_source, selection_source = load_locked_result(root, tables_dir)
    figures = locate_figures(root, figures_dir)

    log_print(f"[INPUT TEMPLATE] {template_path}")
    log_print(f"[KEY RESULT SOURCE] {result_source if result_source else 'fallback constants'}")
    log_print("[LOCKED KEY RESULT]")
    log_print(f"Selective AD retained-subdomain result: {result['model']}")
    log_print(f"Rule: {result['rule']}")
    log_print(f"R2={result['r2']:.6f}, RMSE={result['rmse']:.6f}, MAE={result['mae']:.6f}, Coverage={result['coverage']:.6f}, N={result['n']}, Dominant group={result['dominant_group']}")

    missing_figures = []
    for f in figures:
        log_print(f"[FIGURE] {f['source_id']} exists={f['exists']} path={f['path']}")
        if not f["exists"]:
            missing_figures.append(f["source_id"])

    if missing_figures:
        log_print(f"[WARNING] Missing figure files: {', '.join(missing_figures)}")
        log_print("The DOCX will still be generated with missing-figure placeholders instead of crashing.")

    section_doc = build_results_section_doc(result, figures)
    output_docx = reports_dir / "32s3_results_discussion_template_aligned_path_fixed_polished.docx"
    ok, msg = replace_results_section(template_path, output_docx, section_doc)
    if not ok:
        log_print(f"[ERROR] {msg}")
        (reports_dir / "32s3_run_log.txt").write_text("\n".join(log), encoding="utf-8")
        sys.exit(1)

    final_text = extract_docx_text(output_docx)
    text_path = reports_dir / "32s3_results_discussion_template_aligned_path_fixed_polished_text.txt"
    text_path.write_text(final_text, encoding="utf-8")

    boundary_rows = audit_text(final_text, result)
    fig_rows = []
    for f in figures:
        fig_rows.append({
            "source_id": f["source_id"],
            "figure_no": f["figure_no"],
            "filename": f["filename"],
            "exists": f["exists"],
            "path": f["path"],
            "caption": f["caption"],
        })

    manifest_rows = [
        {"item": "template_docx", "path": str(template_path), "exists": template_path.exists()},
        {"item": "key_result_source", "path": str(result_source) if result_source else "fallback_constants", "exists": bool(result_source and result_source.exists())},
        {"item": "output_docx", "path": str(output_docx), "exists": output_docx.exists()},
        {"item": "output_text", "path": str(text_path), "exists": text_path.exists()},
        {"item": "selection_source", "path": selection_source, "exists": True},
    ]

    write_csv(tables_dir / "32s3_template_aligned_docx_manifest.csv", manifest_rows)
    write_csv(tables_dir / "32s3_final_text_boundary_check.csv", boundary_rows)
    write_csv(tables_dir / "32s3_figure_insertion_manifest.csv", fig_rows)
    write_csv(tables_dir / "32s3_path_resolution_audit.csv", path_audit_rows)

    fail_count = sum(1 for r in boundary_rows if r["status"] == "FAIL")
    warn_count = sum(1 for r in boundary_rows if r["status"] == "WARN")
    log_print("=" * 68)
    log_print(f"[OUTPUT DOCX] {output_docx}")
    log_print(f"[OUTPUT TEXT] {text_path}")
    log_print(f"[OUTPUT TABLE] {tables_dir / '32s3_template_aligned_docx_manifest.csv'}")
    log_print(f"[OUTPUT TABLE] {tables_dir / '32s3_final_text_boundary_check.csv'}")
    log_print(f"[OUTPUT TABLE] {tables_dir / '32s3_figure_insertion_manifest.csv'}")
    log_print(f"[OUTPUT TABLE] {tables_dir / '32s3_path_resolution_audit.csv'}")
    log_print(f"[AUDIT] FAIL={fail_count}; WARN={warn_count}")
    if fail_count == 0:
        log_print("32S3 finished successfully.")
    else:
        log_print("32S3 finished with FAIL rows. Please inspect 32s3_final_text_boundary_check.csv.")
    log_print("Important writing boundary: R2=0.747676 belongs to the selective AD retained-subdomain only.")
    log_print("Please open:")
    log_print(f"    {output_docx}")

    (reports_dir / "32s3_run_log.txt").write_text("\n".join(log), encoding="utf-8")


if __name__ == "__main__":
    main()
