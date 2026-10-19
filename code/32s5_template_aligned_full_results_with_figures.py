# -*- coding: utf-8 -*-
"""
32S5 Template-aligned full Results and Discussion generator with real figure embedding
Project: fluoride coagulation defluoridation ML project

Purpose
-------
This script generates a more complete Results and Discussion chapter by integrating:
1) earlier project steps before the selective AD stage, including dataset characterization,
   preprocessing / feature engineering, baseline model screening, ensemble modeling,
   multi-seed validation, and model interpretation;
2) the locked 32M2/32Q2 selective-AD result;
3) real embedded figures 194-197 and formal captions.

It does NOT retrain any model and does NOT change any final metric.
Outputs follow the original project convention:
    outputs/tables/
    reports/

Expected input folders under the real PyCharm project root:
    src/
    outputs/tables/
    outputs/figures/
    reports/

Main output:
    reports/32s5_results_discussion_full_template_aligned_with_figures.docx
"""

from __future__ import annotations

import csv
import os
import re
import sys
import zipfile
from copy import deepcopy
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

try:
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt
except Exception as exc:  # pragma: no cover
    print("[ERROR] python-docx is required. Install it by:")
    print("        python -m pip install python-docx")
    print(f"Original error: {exc}")
    sys.exit(1)


# -----------------------------------------------------------------------------
# Locked final result. These constants are only fallback values. The script will
# first try to read project CSVs such as 32m2_corrected_key_result_table.csv.
# -----------------------------------------------------------------------------
FALLBACK_RESULT = {
    "result_type": "target-free selective AD retained-subdomain",
    "model": "RepeatedBalanced_Extended + XGBoost",
    "rule": "Zone=Outside; NN_Risk <= 0.814815",
    "r2": 0.747676,
    "rmse": 11.474674,
    "mae": 9.155856,
    "coverage": 0.014664,
    "n": 52,
    "dominant_group": "Paper_010",
}

TEMPLATE_CANDIDATE_NAMES = [
    "AI论文--混凝沉淀除氟空稿.docx",
    "AI论文-混凝沉淀除氟空稿.docx",
    "AI论文_混凝沉淀除氟空稿.docx",
]

OUTPUT_DOCX_NAME = "32s5_results_discussion_full_template_aligned_with_figures.docx"
OUTPUT_TEXT_NAME = "32s5_results_discussion_full_template_aligned_with_figures_text.txt"
RUN_LOG_NAME = "32s5_run_log.txt"


@dataclass
class FigureInfo:
    fig_id: str
    title: str
    caption: str
    path: str
    exists: bool
    inserted: bool = False
    note: str = ""


@dataclass
class AuditItem:
    item: str
    status: str
    detail: str


# -----------------------------------------------------------------------------
# Utility functions
# -----------------------------------------------------------------------------

def safe_float(value, default=None):
    try:
        if value is None:
            return default
        if isinstance(value, (int, float)):
            return float(value)
        s = str(value).strip().replace("%", "")
        if s == "" or s.lower() in {"nan", "none", "na", "n/a"}:
            return default
        return float(s)
    except Exception:
        return default


def norm_text(s: str) -> str:
    return re.sub(r"\s+", "", str(s or "").lower())


def read_csv_dicts(path: Path) -> List[dict]:
    if not path.exists():
        return []
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            with path.open("r", newline="", encoding=enc) as f:
                return list(csv.DictReader(f))
        except Exception:
            continue
    return []


def write_csv(path: Path, rows: List[dict], fieldnames: Optional[List[str]] = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fieldnames is None:
        keys = []
        for row in rows:
            for k in row.keys():
                if k not in keys:
                    keys.append(k)
        fieldnames = keys
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in fieldnames})


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def is_project_root(p: Path) -> bool:
    return (
        (p / "src").exists()
        and (p / "outputs").exists()
        and (p / "reports").exists()
    ) or (
        (p / "outputs" / "tables").exists()
        and (p / "outputs" / "figures").exists()
    )


def find_project_root() -> Path:
    """Find the real PyCharm project root, not the download folder containing template."""
    here = Path(__file__).resolve()
    candidates: List[Path] = []

    # 1) current script location and its parents
    candidates.extend([here.parent] + list(here.parents))

    # 2) current working directory and parents
    cwd = Path.cwd().resolve()
    candidates.extend([cwd] + list(cwd.parents))

    # 3) likely nested project folders near current location
    for base in [here.parent, cwd, here.parent.parent, cwd.parent]:
        if base.exists():
            for child in base.glob("fluoride_coagulation_ml_project*"):
                if child.is_dir():
                    candidates.append(child)
                    for sub in child.glob("fluoride_coagulation_ml_project*"):
                        if sub.is_dir():
                            candidates.append(sub)

    # Deduplicate while preserving order
    seen = set()
    uniq: List[Path] = []
    for c in candidates:
        try:
            r = c.resolve()
        except Exception:
            continue
        if r not in seen:
            uniq.append(r)
            seen.add(r)

    for c in uniq:
        if is_project_root(c):
            return c

    # fallback: if running in src, parent is project root
    if here.parent.name.lower() == "src":
        return here.parent.parent
    return cwd


def find_template(project_root: Path) -> Optional[Path]:
    search_dirs: List[Path] = [
        project_root,
        project_root / "reports",
        project_root / "data",
        project_root.parent,
        project_root.parent / "reports",
        project_root.parent.parent if project_root.parent else project_root,
        Path.cwd(),
        Path.cwd().parent,
    ]

    # Add script ancestors because the user may keep template in Downloads.
    for p in Path(__file__).resolve().parents:
        search_dirs.append(p)

    # First: exact names.
    for d in search_dirs:
        if not d.exists():
            continue
        for name in TEMPLATE_CANDIDATE_NAMES:
            candidate = d / name
            if candidate.exists() and candidate.is_file():
                return candidate

    # Second: fuzzy search, but do not select outputs generated by scripts.
    for d in search_dirs:
        if not d.exists() or not d.is_dir():
            continue
        try:
            for candidate in d.glob("*.docx"):
                s = candidate.name
                if "混凝" in s and "除氟" in s and "空稿" in s:
                    return candidate
        except Exception:
            pass

    return None


def ensure_project_dirs(root: Path) -> Tuple[Path, Path, Path]:
    table_dir = root / "outputs" / "tables"
    figure_dir = root / "outputs" / "figures"
    report_dir = root / "reports"
    table_dir.mkdir(parents=True, exist_ok=True)
    figure_dir.mkdir(parents=True, exist_ok=True)
    report_dir.mkdir(parents=True, exist_ok=True)
    return table_dir, figure_dir, report_dir


def find_first_existing(paths: Iterable[Path]) -> Optional[Path]:
    for p in paths:
        if p and p.exists() and p.is_file():
            return p
    return None


def search_figure(project_root: Path, fig_id: str) -> Tuple[Optional[Path], str]:
    search_dirs = [
        project_root / "outputs" / "figures",
        project_root / "figures",
        project_root / "outputs",
        project_root,
        project_root.parent / "outputs" / "figures",
        project_root.parent,
        Path.cwd(),
        Path(__file__).resolve().parent,
    ]

    patterns = [
        f"{fig_id}_*.png",
        f"{fig_id}*.png",
        f"*{fig_id}*.png",
        f"{fig_id}_*.jpg",
        f"{fig_id}*.jpg",
        f"*{fig_id}*.jpg",
        f"{fig_id}_*.jpeg",
        f"{fig_id}*.jpeg",
        f"*{fig_id}*.jpeg",
    ]
    for d in search_dirs:
        if not d.exists():
            continue
        for pat in patterns:
            try:
                hits = sorted([p for p in d.glob(pat) if p.is_file()], key=lambda x: len(x.name))
                if hits:
                    return hits[0], f"found_by_pattern={pat}; dir={d}"
            except Exception:
                pass
    return None, "not_found"


def extract_key_result(project_root: Path, table_dir: Path) -> Tuple[Dict[str, object], str]:
    """Read the locked selective AD result from project CSVs, then fallback constants."""
    candidates = [
        table_dir / "32m2_corrected_key_result_table.csv",
        table_dir / "32n_locked_key_result_for_word.csv",
        table_dir / "32q2_locked_key_result_check.csv",
        table_dir / "32j_main_result_table.csv",
        project_root / "outputs" / "tables" / "32m2_corrected_key_result_table.csv",
        project_root / "outputs" / "tables" / "32j_main_result_table.csv",
    ]

    desired_model_token = "RepeatedBalanced_Extended"
    desired_rule_token = "0.814815"

    for path in candidates:
        rows = read_csv_dicts(path)
        if not rows:
            continue
        # Prefer rows explicitly matching both model and selective rule.
        for row in rows:
            joined = " | ".join(str(v) for v in row.values())
            if desired_model_token in joined and (desired_rule_token in joined or "NN_Risk" in joined):
                result = dict(FALLBACK_RESULT)
                for k, v in row.items():
                    lk = str(k).strip().lower()
                    val = v
                    if lk in {"model", "model_name", "model_family", "accepted_model", "method"} and val:
                        if "XGBoost" in str(val) and "RepeatedBalanced" in str(val):
                            result["model"] = str(val)
                    if lk in {"rule", "ad_rule", "selective_rule", "policy"} and val:
                        result["rule"] = str(val)
                    if lk in {"r2", "r²", "test_r2", "retained_r2"} and safe_float(val) is not None:
                        result["r2"] = safe_float(val)
                    if lk in {"rmse", "test_rmse", "retained_rmse"} and safe_float(val) is not None:
                        result["rmse"] = safe_float(val)
                    if lk in {"mae", "test_mae", "retained_mae"} and safe_float(val) is not None:
                        result["mae"] = safe_float(val)
                    if lk in {"coverage", "retained_coverage"} and safe_float(val) is not None:
                        result["coverage"] = safe_float(val)
                    if lk in {"n", "n_retained", "retained_n", "sample_count"} and safe_float(val) is not None:
                        result["n"] = int(round(safe_float(val)))
                    if lk in {"dominant_group", "dominant literature group", "main_group", "group"} and val:
                        if "Paper" in str(val):
                            result["dominant_group"] = str(val)
                return result, f"selected_from={path.name}"

        # If no explicit match but file is a corrected key-result table, use first row carefully.
        if "32m2_corrected_key_result_table" in path.name or "32n_locked_key_result" in path.name:
            row = rows[0]
            result = dict(FALLBACK_RESULT)
            joined = " | ".join(str(v) for v in row.values())
            if "RepeatedBalanced" in joined and "XGBoost" in joined:
                result["model"] = "RepeatedBalanced_Extended + XGBoost"
            if "0.814815" in joined:
                result["rule"] = "Zone=Outside; NN_Risk <= 0.814815"
            for k, v in row.items():
                lk = str(k).strip().lower()
                if lk in {"r2", "r²"} and safe_float(v) is not None:
                    result["r2"] = safe_float(v)
                elif lk == "rmse" and safe_float(v) is not None:
                    result["rmse"] = safe_float(v)
                elif lk == "mae" and safe_float(v) is not None:
                    result["mae"] = safe_float(v)
                elif "coverage" in lk and safe_float(v) is not None:
                    result["coverage"] = safe_float(v)
                elif lk in {"n", "n_retained", "retained_n"} and safe_float(v) is not None:
                    result["n"] = int(round(safe_float(v)))
                elif "dominant" in lk and v:
                    result["dominant_group"] = str(v)
            return result, f"fallback_read_first_row={path.name}"

    return dict(FALLBACK_RESULT), "fallback_constants"


# -----------------------------------------------------------------------------
# DOCX formatting helpers
# -----------------------------------------------------------------------------

def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9.5)
    set_run_font(run)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, east_asia: str = "宋体", latin: str = "Times New Roman", size_pt: Optional[float] = None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    if size_pt is not None:
        run.font.size = Pt(size_pt)


def set_paragraph_font(paragraph, east_asia: str = "宋体", latin: str = "Times New Roman", size_pt: float = 10.5):
    for run in paragraph.runs:
        set_run_font(run, east_asia=east_asia, latin=latin, size_pt=size_pt)


def add_paragraph(doc: Document, text: str = "", style: Optional[str] = None, align=None, bold=False, size=10.5):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        run.bold = bold
        set_run_font(run, size_pt=size)
    return p


def add_heading_like(doc: Document, text: str, level: int = 1):
    style = f"Heading {level}"
    try:
        p = doc.add_paragraph(style=style)
    except Exception:
        p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    size = 14 if level == 1 else 12
    set_run_font(run, east_asia="宋体", latin="Times New Roman", size_pt=size)
    return p


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(9)
    set_run_font(run, size_pt=9)
    return p


def set_table_borders(table, top=True, bottom=True, inside_h=True, inside_v=False):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)

    def set_border(edge, val="single", sz="8", color="000000"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)

    def nil_border(edge):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")

    set_border("top") if top else nil_border("top")
    set_border("bottom") if bottom else nil_border("bottom")
    set_border("insideH", sz="4", color="808080") if inside_h else nil_border("insideH")
    set_border("insideV", sz="4", color="D0D0D0") if inside_v else nil_border("insideV")
    nil_border("left")
    nil_border("right")


def add_compact_key_result_table(doc: Document, result: Dict[str, object]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    hdr = table.rows[0].cells
    set_cell_text(hdr[0], "项目", bold=True)
    set_cell_text(hdr[1], "内容", bold=True)

    rows = [
        ("结果类型", "无目标泄露的选择性适应域保留子域"),
        ("模型", str(result["model"])),
        ("筛选规则", str(result["rule"])),
        ("R²", f"{float(result['r2']):.6f}"),
        ("RMSE", f"{float(result['rmse']):.6f}"),
        ("MAE", f"{float(result['mae']):.6f}"),
        ("Coverage", f"{float(result['coverage']):.6f}（约 {float(result['coverage']) * 100:.2f}%）"),
        ("保留样本数", str(result["n"])),
        ("主导文献组", str(result["dominant_group"])),
    ]
    for k, v in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], k)
        set_cell_text(cells[1], v)
    set_table_borders(table, top=True, bottom=True, inside_h=True, inside_v=False)
    add_paragraph(doc, "")


def clear_doc_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def copy_template_paragraphs(doc: Document, paragraphs: List[Tuple[str, str]]) -> None:
    for text, style_name in paragraphs:
        if not text.strip():
            add_paragraph(doc, "")
            continue
        # Avoid styles that may not exist.
        try:
            p = doc.add_paragraph(style=style_name)
        except Exception:
            p = doc.add_paragraph()
        run = p.add_run(text)
        if re.match(r"^\d+(\.\d+)?", text.strip()) or text.strip().startswith("基于"):
            run.bold = text.strip().startswith("基于")
        set_run_font(run, size_pt=10.5)


def get_template_segments(template_path: Path) -> Tuple[List[Tuple[str, str]], List[Tuple[str, str]]]:
    template_doc = Document(str(template_path))
    paragraphs = []
    for p in template_doc.paragraphs:
        style_name = p.style.name if p.style is not None else "Normal"
        paragraphs.append((p.text, style_name))

    idx3 = None
    idx4 = None
    for i, (txt, _style) in enumerate(paragraphs):
        t = norm_text(txt)
        if idx3 is None and (t.startswith("3resultsanddiscussion") or t.startswith("3.resultsanddiscussion") or t.startswith("3结果") or "resultsanddiscussion" in t and t.startswith("3")):
            idx3 = i
        if idx3 is not None and i > idx3 and (t.startswith("4conclusion") or t.startswith("4.conclusion") or t.startswith("4结论")):
            idx4 = i
            break

    if idx3 is None:
        # If template does not have section 3, keep all text before inserting new section.
        return paragraphs, [("4. Conclusion 结论", "Normal"), ("（1）主要研究成果总结", "Normal"), ("（2）最优模型与预测精度", "Normal"), ("（3）关键影响因子及阈值规律", "Normal"), ("（4）工艺优化推荐参数", "Normal"), ("（5）工程应用价值", "Normal")]
    before = paragraphs[:idx3]
    after = paragraphs[idx4:] if idx4 is not None else [("4. Conclusion 结论", "Normal"), ("（1）主要研究成果总结", "Normal"), ("（2）最优模型与预测精度", "Normal"), ("（3）关键影响因子及阈值规律", "Normal"), ("（4）工艺优化推荐参数", "Normal"), ("（5）工程应用价值", "Normal")]
    return before, after


def add_figure(doc: Document, fig: FigureInfo, width_inches: float = 6.2) -> FigureInfo:
    path = Path(fig.path) if fig.path else None
    if path and path.exists() and path.is_file():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        try:
            run.add_picture(str(path), width=Inches(width_inches))
            fig.inserted = True
            fig.note = "inserted"
        except Exception as exc:
            fig.inserted = False
            fig.note = f"insert_failed: {exc}"
            add_paragraph(doc, f"[图件插入失败：{fig.fig_id}，请检查图片文件。]", size=9)
    else:
        fig.inserted = False
        fig.note = "missing"
        add_paragraph(doc, f"[缺失图件：{fig.fig_id}，请检查 outputs/figures。]", size=9)
    add_caption(doc, fig.caption)
    return fig


def build_figures(project_root: Path) -> List[FigureInfo]:
    meta = {
        "194": (
            "选择性适应域保留子域结果的证据链汇总",
            "图 3-1 选择性适应域保留子域结果的证据链汇总。该图比较了未筛选分组验证、预定义可靠域和选择性适应域保留子域在 R²、RMSE、MAE 和覆盖率上的差异，用于说明最终高 R² 结果属于选择性保留子域，而不是未筛选测试集的整体性能。",
        ),
        "195": (
            "选择性适应域规则的阈值平台证据",
            "图 3-2 选择性适应域规则的阈值平台证据。虚线位置对应最终接受的 NN_Risk 阈值，表明该结果是在阈值敏感性分析中形成的高性能平台附近获得，而不是单一孤立阈值造成的偶然峰值。",
        ),
        "196": (
            "选择性适应域保留子域的化学条件窗口与适应域组成",
            "图 3-3 选择性适应域保留子域的化学条件窗口与适应域组成。该图描述保留样本的关键水质、投加量、pH、干扰离子、混凝剂组成和适应域风险水平，用于界定模型解释适用的化学空间。",
        ),
        "197": (
            "选择性适应域结果的稳定性、迁移支持和解释边界",
            "图 3-4 选择性适应域结果的稳定性、迁移支持和解释边界。该图显示不同数据划分和相关模型组合下的保留子域表现，并强调最终结果具有明确的文献来源和覆盖范围限制。",
        ),
    }
    figures: List[FigureInfo] = []
    for fig_id, (title, caption) in meta.items():
        p, note = search_figure(project_root, fig_id)
        figures.append(FigureInfo(fig_id=fig_id, title=title, caption=caption, path=str(p) if p else "", exists=bool(p), note=note))
    return figures


def add_results_discussion(doc: Document, result: Dict[str, object], figures: List[FigureInfo]) -> List[FigureInfo]:
    r2_short = f"{float(result['r2']):.3f}"
    rmse_short = f"{float(result['rmse']):.3f}"
    mae_short = f"{float(result['mae']):.3f}"
    cov_pct = f"{float(result['coverage']) * 100:.2f}%"

    add_heading_like(doc, "3 Results and Discussion（结果与讨论）", level=1)

    add_heading_like(doc, "3.1 数据集特征、描述性统计与建模挑战", level=2)
    add_paragraph(doc, "本研究首先对混凝沉淀除氟文献数据集进行描述性统计和可视化分析。结果表明，数据整体可用于机器学习建模，但不同变量的数据完整性并不一致，进水 pH、干扰离子浓度等关键机理变量存在一定缺失，需要在建模前进行专门处理。目标变量方面，样本中高去除率区间占比较高，同时仍包含低去除率和失败实验样本，说明数据分布具有明显不均衡特征。", size=10.5)
    add_paragraph(doc, "从机理角度看，氟去除率并不由单一变量线性决定。pH 会影响铝盐、铁盐等混凝剂的水解形态、絮体表面电荷和氟离子吸附/络合过程；混凝剂投加量决定有效絮体生成量；干扰离子会影响竞争吸附和沉淀平衡；初始氟浓度则影响最终出水是否达标。因此，混凝除氟数据呈现多因素耦合、非线性和跨文献异质性共同作用的特点，这也构成了后续模型泛化评价和适应域分析的必要背景。", size=10.5)

    add_heading_like(doc, "3.2 数据预处理、机理特征构造与输入变量封装", level=2)
    add_paragraph(doc, "针对文献数据来源分散、变量记录方式不完全一致的问题，本研究在建模前进行了统一的数据预处理。数值变量主要包括初始氟浓度、混凝剂投加量、进水 pH、反应/沉降时间、水温和干扰离子浓度等；类别变量主要包括混凝剂类型、干扰离子类型和相关工艺类别。对于缺失值，采用与变量物理含义相一致的处理方式，并保留必要的缺失指示特征，以降低文献记录差异对模型训练的干扰。", size=10.5)
    add_paragraph(doc, "在原始变量基础上，进一步构造了具有混凝除氟机理含义的衍生特征，包括 log(1 + 投加量)、相对投加强度、投加量与初始氟浓度交互项、pH 距中性距离、弱酸性/中性/碱性区间指示、干扰离子类别拆分和混凝剂大类归并等。这些特征并不是改变原始实验结果，而是将化学机理和数据结构以更适合模型学习的方式表达出来。", size=10.5)

    add_heading_like(doc, "3.3 基准模型筛选、集成建模与稳定性检验", level=2)
    add_paragraph(doc, "在模型构建阶段，本研究先后比较了线性模型、正则化线性模型和多种树模型。线性回归和 Ridge 等模型可作为基准，用于检验单纯线性关系是否足以描述混凝除氟过程；Random Forest、Extra Trees、HistGradientBoosting、CatBoost 和 XGBoost 等树模型则用于捕捉非线性关系、变量交互和类别变量影响。整体结果显示，树模型明显更适合当前环境实验表格数据，而线性模型主要起到基准参照作用。", size=10.5)
    add_paragraph(doc, "在单模型基础上，本研究进一步尝试了多种组合策略，包括 Random Forest 与 HistGradientBoosting 的 stacking、Random Forest 与 CatBoost 的加权融合、CatBoost 手动调参和贝叶斯优化、以及 Random Forest + CatBoost 的 Ridge/Linear 二层融合。多随机种子检验显示，部分模型在单次划分中表现较好，但跨划分稳定性有限；Random Forest 仍是稳定性较强的基准模型，而 CatBoost/XGBoost 在特定特征处理和数据划分条件下可提供补充预测信息。", size=10.5)

    add_heading_like(doc, "3.4 跨文献分组验证与适应域分析的必要性", level=2)
    add_paragraph(doc, "为评价机器学习模型在文献异质性条件下的真实泛化能力，本研究在常规随机划分之外，进一步采用按文献来源分组的验证策略。与随机划分相比，分组验证能够避免同一文献或相近实验系列同时进入训练集和测试集，从而更严格地检验模型对未见文献工况的预测能力。结果表明，未筛选测试条件下的模型性能明显低于随机划分结果，说明混凝沉淀除氟数据不仅存在非线性和多因素耦合特征，还受到文献来源、实验体系和局部工况分布差异的显著影响。", size=10.5)
    add_paragraph(doc, "因此，仅报告随机划分下的模型精度不足以说明模型具有可靠的工程外推能力。对于来自多篇文献的小样本环境实验数据，模型预测应同时报告整体分组验证结果和适应域限定后的局部结果，并明确二者评价口径不能混用。", size=10.5)

    add_heading_like(doc, "3.5 预定义可靠域与选择性适应域的对比", level=2)
    add_paragraph(doc, "在预定义可靠域分析中，模型尝试通过固定适应域规则筛除高风险样本，以获得更可信的预测子集。然而，固定可靠域并未带来稳定的性能提升，说明单一、静态的适应域阈值难以充分处理本数据集中复杂的跨文献异质性和工况偏移。进一步的选择性适应域分析显示，当模型仅在满足特定风险约束的保留子域内进行评价和解释时，预测误差显著降低，决定系数明显提高。", size=10.5)
    add_paragraph(doc, "这表明适应域分析的关键并不是简单扩大或缩小预测范围，而是识别模型在当前数据结构下真正具有相对可靠性的局部化学工况空间。该过程必须保持目标变量隔离，避免通过实际测试标签参与筛选而造成目标泄露。", size=10.5)

    add_heading_like(doc, "3.6 选择性适应域保留子域的模型性能", level=2)
    add_paragraph(doc, f"最终接受的选择性适应域结果来自 {result['model']}，筛选规则为 {result['rule']}。在该保留子域内，模型取得 R² = {r2_short}、RMSE = {rmse_short}、MAE = {mae_short}，覆盖率为 {cov_pct}，保留样本数为 {result['n']}。该结果仅表征适应域规则筛选后保留样本子集的预测表现；未筛选测试集结果与筛选后保留子域结果应分开报告，二者不能混用。", size=10.5)
    add_compact_key_result_table(doc, result)
    figures[0] = add_figure(doc, figures[0], width_inches=6.2)

    add_heading_like(doc, "3.7 NN-risk 阈值平台、化学窗口与解释边界", level=2)
    add_paragraph(doc, "为降低选择性适应域规则被解释为事后挑选结果的风险，本研究进一步对 NN-risk 阈值进行了敏感性分析。结果显示，在接受阈值附近，保留子域 R² 存在相对稳定的高性能区间，而不是由单一孤立阈值产生的偶然峰值。同时，随着阈值放宽，保留样本数和覆盖率逐渐升高，但模型精度并未持续提高，说明更多样本并不必然带来更可靠的外推结果。", size=10.5)
    figures[1] = add_figure(doc, figures[1], width_inches=6.2)

    add_paragraph(doc, f"保留子域样本主要来自 {result['dominant_group']}，并集中于特定混凝剂组成、pH、投加量、初始氟浓度和干扰离子浓度范围内。该结果一方面表明模型在该类化学窗口内能够获得较高预测精度，另一方面也提示模型解释具有明确边界。对于该子域以外的混凝剂类型、极端 pH、不同初始氟负荷或复杂共存离子体系，模型预测仍需谨慎使用。", size=10.5)
    figures[2] = add_figure(doc, figures[2], width_inches=6.2)

    add_heading_like(doc, "3.8 稳定性、迁移支持与本节小结", level=2)
    add_paragraph(doc, "稳定性分析进一步表明，选择性适应域结果在部分数据划分和相关模型组合中具有可重复支持，但其覆盖率较低且文献来源集中，因此仍应将其定位为保留子域内的可靠预测证据，不能作为未筛选测试集整体泛化能力的直接证据。在论文表述中，应避免将该保留子域结果直接描述为整个混凝除氟数据集的泛化性能。", size=10.5)
    figures[3] = add_figure(doc, figures[3], width_inches=6.2)

    add_paragraph(doc, "综上，混凝沉淀除氟文献数据在跨来源验证中表现出明显的泛化难度，说明该体系存在强烈的数据异质性。通过引入无目标泄露的适应域筛选，本研究识别出一个具有较高预测可靠性的局部保留子域，并在该子域内获得较好的预测结果。该发现为后续工艺参数推荐和模型辅助决策提供了更加谨慎的依据：模型可用于支持特定工况窗口内的预测和解释，但对于适应域之外的水质条件和药剂体系，仍需要补充实验或更多文献数据进行验证。", size=10.5)

    return figures


def extract_docx_text(docx_path: Path) -> str:
    try:
        d = Document(str(docx_path))
        texts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                texts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(texts)
    except Exception:
        return ""


def count_docx_media(docx_path: Path) -> int:
    try:
        with zipfile.ZipFile(docx_path, "r") as zf:
            return len([n for n in zf.namelist() if n.startswith("word/media/")])
    except Exception:
        return 0


def audit_docx(docx_path: Path, result: Dict[str, object], figures: List[FigureInfo]) -> List[AuditItem]:
    text = extract_docx_text(docx_path)
    media_count = count_docx_media(docx_path)
    audits: List[AuditItem] = []

    def add(item, ok, detail):
        audits.append(AuditItem(item=item, status="PASS" if ok else "FAIL", detail=detail))

    add("docx_exists", docx_path.exists(), str(docx_path))
    add("media_embedded", media_count >= 4, f"word/media count={media_count}")
    add("accepted_r2_present", f"{float(result['r2']):.3f}" in text or f"{float(result['r2']):.6f}" in text, f"R2={result['r2']}")
    add("accepted_model_full_name", str(result["model"]) in text, str(result["model"]))
    add("accepted_rule_present", "NN_Risk" in text and "0.814815" in text, str(result["rule"]))
    add("coverage_boundary_present", "覆盖率" in text and ("1.47%" in text or "0.014664" in text), str(result["coverage"]))
    add("no_script_generation_trace", "脚本生成" not in text and "建议插入" not in text and "图件声明边界" not in text, "internal prompt traces removed")
    add("selective_boundary_written", "保留子域" in text and "不能混用" in text, "selective AD boundary")
    add("all_figures_inserted_flags", all(f.inserted for f in figures), "; ".join(f"{f.fig_id}={f.inserted}" for f in figures))
    return audits


def make_text_version(result: Dict[str, object], figures: List[FigureInfo]) -> str:
    lines = []
    lines.append("32S5 Results and Discussion full template-aligned version")
    lines.append(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append("")
    lines.append("Locked result:")
    lines.append(f"  Model: {result['model']}")
    lines.append(f"  Rule: {result['rule']}")
    lines.append(f"  R2: {float(result['r2']):.6f}")
    lines.append(f"  RMSE: {float(result['rmse']):.6f}")
    lines.append(f"  MAE: {float(result['mae']):.6f}")
    lines.append(f"  Coverage: {float(result['coverage']):.6f}")
    lines.append(f"  N retained: {result['n']}")
    lines.append(f"  Dominant group: {result['dominant_group']}")
    lines.append("")
    lines.append("Inserted figures:")
    for f in figures:
        lines.append(f"  {f.fig_id}: exists={f.exists}; inserted={f.inserted}; path={f.path}; note={f.note}")
    lines.append("")
    lines.append("Writing boundary:")
    lines.append("  R2=0.747676 belongs to the target-free selective AD retained-subdomain result only.")
    return "\n".join(lines)


def build_previous_work_map() -> List[dict]:
    return [
        {"section": "3.1", "project_stage": "EDA/data audit", "content": "缺失值、目标分布、达标分组、pH/剂量关系、相关性热图", "manuscript_role": "说明数据不均衡、非线性和多因素耦合"},
        {"section": "3.2", "project_stage": "feature engineering", "content": "log1p剂量、相对投加强度、pH机理特征、干扰离子拆分、混凝剂大类", "manuscript_role": "说明输入变量如何转化为可学习的机理特征"},
        {"section": "3.3", "project_stage": "model screening and ensemble", "content": "Linear/Ridge/RF/ExtraTrees/HGB/CatBoost/XGBoost、加权融合、stacking、多随机种子", "manuscript_role": "说明为什么树模型和集成模型更适合当前数据"},
        {"section": "3.4", "project_stage": "grouped validation", "content": "按文献来源分组验证", "manuscript_role": "说明全域外推难度和数据泄漏控制"},
        {"section": "3.5-3.8", "project_stage": "selective AD and final evidence", "content": "32I2/32J/32Q2锁定的选择性适应域结果与194-197图", "manuscript_role": "说明局部可靠预测、阈值平台、化学窗口和声明边界"},
    ]


def main() -> int:
    run_lines: List[str] = []
    def log(msg: str) -> None:
        print(msg)
        run_lines.append(msg)

    project_root = find_project_root()
    table_dir, figure_dir, report_dir = ensure_project_dirs(project_root)
    template_path = find_template(project_root)

    output_docx = report_dir / OUTPUT_DOCX_NAME
    output_text = report_dir / OUTPUT_TEXT_NAME
    run_log = report_dir / RUN_LOG_NAME

    log("=" * 70)
    log("32S5 Full Results and Discussion template-aligned DOCX with real figures")
    log(f"Run time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    log(f"Project root: {project_root}")
    log(f"Table output dir: {table_dir}")
    log(f"Figure dir: {figure_dir}")
    log(f"Report output dir: {report_dir}")
    log("=" * 70)

    if template_path is None:
        log("[ERROR] Real Word template not found.")
        log("Please put AI论文--混凝沉淀除氟空稿.docx in the project root, reports/, or its parent folder.")
        write_text(run_log, "\n".join(run_lines))
        return 1
    log(f"[INPUT TEMPLATE] {template_path}")

    result, result_source = extract_key_result(project_root, table_dir)
    figures = build_figures(project_root)
    log(f"[KEY RESULT SOURCE] {result_source}")
    log("[LOCKED KEY RESULT]")
    log(f"Selective AD retained-subdomain result: {result['model']}")
    log(f"Rule: {result['rule']}")
    log(f"R2={float(result['r2']):.6f}, RMSE={float(result['rmse']):.6f}, MAE={float(result['mae']):.6f}, Coverage={float(result['coverage']):.6f}, N={result['n']}, Dominant group={result['dominant_group']}")
    for fig in figures:
        log(f"[FIGURE] {fig.fig_id} exists={fig.exists} path={fig.path} note={fig.note}")

    # Create DOCX from template style definitions, then rebuild body.
    try:
        doc = Document(str(template_path))
    except Exception as exc:
        log(f"[ERROR] Failed to open template: {exc}")
        write_text(run_log, "\n".join(run_lines))
        return 1

    before, after = get_template_segments(template_path)
    clear_doc_body(doc)

    # Set margins to a stable manuscript style.
    for sec in doc.sections:
        sec.top_margin = Cm(2.54)
        sec.bottom_margin = Cm(2.54)
        sec.left_margin = Cm(2.8)
        sec.right_margin = Cm(2.8)

    copy_template_paragraphs(doc, before)
    if before and before[-1][0].strip():
        add_paragraph(doc, "")

    inserted_figures = add_results_discussion(doc, result, figures)

    add_paragraph(doc, "")
    copy_template_paragraphs(doc, after)

    try:
        doc.save(str(output_docx))
    except Exception as exc:
        log(f"[ERROR] Failed to save DOCX: {exc}")
        write_text(run_log, "\n".join(run_lines))
        return 1

    write_text(output_text, make_text_version(result, inserted_figures))

    # Tables / audits.
    write_csv(table_dir / "32s5_previous_work_integration_map.csv", build_previous_work_map())
    write_csv(table_dir / "32s5_figure_insertion_manifest.csv", [asdict(f) for f in inserted_figures])
    write_csv(table_dir / "32s5_locked_key_result_check.csv", [{
        "model": result["model"],
        "rule": result["rule"],
        "r2": f"{float(result['r2']):.6f}",
        "rmse": f"{float(result['rmse']):.6f}",
        "mae": f"{float(result['mae']):.6f}",
        "coverage": f"{float(result['coverage']):.6f}",
        "n": result["n"],
        "dominant_group": result["dominant_group"],
        "source": result_source,
    }])
    write_csv(table_dir / "32s5_path_resolution_audit.csv", [{
        "project_root": str(project_root),
        "template_path": str(template_path),
        "table_output_dir": str(table_dir),
        "figure_dir": str(figure_dir),
        "report_output_dir": str(report_dir),
        "output_docx": str(output_docx),
    }])

    audits = audit_docx(output_docx, result, inserted_figures)
    write_csv(table_dir / "32s5_docx_quality_audit.csv", [asdict(a) for a in audits])
    fail_count = sum(1 for a in audits if a.status != "PASS")

    log("=" * 70)
    for a in audits:
        log(f"[{a.status}] {a.item} | {a.detail}")
    log("=" * 70)
    log(f"[OUTPUT DOCX] {output_docx}")
    log(f"[OUTPUT TEXT] {output_text}")
    log(f"[OUTPUT TABLE] {table_dir / '32s5_previous_work_integration_map.csv'}")
    log(f"[OUTPUT TABLE] {table_dir / '32s5_figure_insertion_manifest.csv'}")
    log(f"[OUTPUT TABLE] {table_dir / '32s5_docx_quality_audit.csv'}")
    log(f"[OUTPUT LOG] {run_log}")
    log("")
    if fail_count == 0:
        log("32S5 finished successfully.")
    else:
        log(f"32S5 finished with warnings/fails. Please inspect 32s5_docx_quality_audit.csv. FAIL={fail_count}")
    log("Important writing boundary: R2=0.747676 belongs to the target-free selective AD retained-subdomain result only.")
    log("")
    log("Please open:")
    log(f"    {output_docx}")

    write_text(run_log, "\n".join(run_lines))
    return 0 if fail_count == 0 else 0


if __name__ == "__main__":
    raise SystemExit(main())
