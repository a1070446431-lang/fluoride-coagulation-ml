# -*- coding: utf-8 -*-
"""Create the final JECE Supplementary Information v3 from the completed v2 Word file.

Place this complete script in ``fluoride_coagulation_ml_project/src/`` and run
it in PyCharm. It reads the existing v2 Word file and writes
``outputs/jece_final_package/reports/Supplementary_Information_Paper1_JECE_v3.docx``.
The v2 file is never overwritten. Dependency: python-docx.
"""
from __future__ import annotations

import hashlib
import json
import re
import time
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


INPUT_NAME = "Supplementary_Information_Paper1_JECE_v2.docx"
OUTPUT_NAME = "Supplementary_Information_Paper1_JECE_v3.docx"
RUN_START = time.perf_counter()


TEXT_REPLACEMENTS = {
    "The locked modeling matrix contains 1,011 observations and 26 stored fields, while the source-bearing audit matrix contains the same observations and 27 fields. Stored fields are not synonymous with predictors. Only information available before treatment prediction was eligible as formal input; literature identity, post-treatment measurements, target-derived quantities, raw reconstruction fields, and row identifiers were excluded (Table S1; Figures S1–S2). Descriptive continuous and categorical coverage is summarized in Table S2.":
        "The modeling dataset contains 1,011 observations. The prediction matrix has 26 stored fields, whereas the companion matrix adds literature-source identity for grouped validation. Only variables available before treatment were eligible as predictors. Source identity, post-treatment measurements, target-derived quantities, reconstruction fields, and row identifiers were excluded from model input (Table S1; Figures S1–S2). Table S2 summarizes numerical coverage and the distribution of categorical levels across literature sources.",
    "Estimator settings were recovered from the original analysis record without refitting models (Table S3). Across five prespecified seeds, the weighted four-model ensemble achieved R² = 0.882 ± 0.010, RMSE = 10.249, and MAE = 6.404; the corresponding pooled values were R² = 0.882, RMSE = 10.267, and MAE = 6.404 (Table S4; Figure S3). These results quantify internal learning ability under random partitioning rather than cross-literature transfer.":
        "Table S3 reports the estimator configurations used in the analysis. Across seeds 7, 21, 42, 66, and 88, the weighted four-model ensemble achieved R² = 0.882 ± 0.010, RMSE = 10.249 percentage points, and MAE = 6.404 percentage points. When all test predictions were pooled, R² was 0.882, RMSE was 10.267 percentage points, and MAE was 6.404 percentage points (Table S4; Figure S3). Random splitting measures interpolation within the combined dataset and does not establish transfer to an unseen literature source.",
    "Grouped validation and leave-one-literature-out testing were markedly weaker than random splitting, revealing substantial between-literature heterogeneity (Tables S5–S6; Figures S4–S5). In the repeated balanced audit, the strongest pooled result was obtained by the KRR–CatBoost weighted model (R² = 0.194, RMSE = 27.513, MAE = 22.894). All source-isolated splits had zero source overlap.":
        "Performance declined under grouped validation and leave-one-literature-out testing, indicating substantial variation among literature sources (Tables S5–S6; Figures S4–S5). Under repeated balanced grouped splitting, the KRR–CatBoost weighted model gave the highest pooled R² among the evaluated methods (R² = 0.194, RMSE = 27.513 percentage points, and MAE = 22.894 percentage points). Training, validation, and test source sets did not overlap.",
    "At the exact target-independent P08 operating point, retained predictions yielded R² = 0.748, RMSE = 11.475, and MAE = 9.156, but represented only 52 prediction events and 44 unique observations (Tables S7–S9; Figures S6–S9). All retained events originated from Paper_010, only one of two retained splits achieved R² ≥ 0.5, and transfer across validation settings was inconsistent. The retained region is therefore interpreted as a narrow source-specific candidate window rather than a universal cross-literature applicability domain.":
        "The P08 rule retained 52 prediction events corresponding to 44 unique observations and produced R² = 0.748, RMSE = 11.475 percentage points, and MAE = 9.156 percentage points (Tables S7–S9; Figures S6–S9). All retained events came from Paper_010, and only one of the two retained splits reached R² ≥ 0.5. Applying the same threshold in other validation settings did not yield consistent performance. The retained region therefore represents a narrow, source-specific operating window rather than a general cross-literature applicability domain.",
    "SHAP attribution varied across the full dataset, Paper_010, and the retained domain, with coagulant identity, dosage, and interferent identity remaining prominent (Table S10; Figure S10). The two-dimensional partial-dependence interpretation was evaluated against empirical grid support (Figure S11). SHAP and PDP patterns describe the fitted model and are not interpreted as causal effects or prescriptive chemical optima.":
        "SHAP importance changed between the full dataset, Paper_010, and the retained subset, although coagulant identity, dose, and interferent identity remained influential (Table S10; Figure S10). Figure S11 compares the two-dimensional partial-dependence surfaces with observed-data support in the same feature space. These analyses describe the fitted model; they do not demonstrate causal effects or define process optima.",
}


SHORT_REPLACEMENTS = {
    "Table S1. Variable dictionary, analytical roles, and leakage controls": "Table S1. Variable definitions and predictor eligibility",
    "Table S2. Data support and categorical-source structure": "Table S2. Data coverage and categorical structure",
    "Table S3. Recovered estimator configurations": "Table S3. Model configurations used in the analysis",
    "Table S5. Pooled source-isolated validation performance": "Table S5. Performance under literature-source-isolated validation",
    "Table S6. Source-isolation and source-macro performance": "Table S6. Source separation and LOLO macro performance",
    "Table S7. Exact selective-AD policy and split stability": "Table S7. P08 selective-domain performance and split stability",
    "Table S8. Transfer of the fixed AD rule across validation settings": "Table S8. Transfer of the fixed P08 rule with XGBoost",
    "Table S9. Selective-domain chemistry and representativeness": "Table S9. Representativeness of the retained domain",
    "Table S10. SHAP importance and PDP response range": "Table S10. Context-dependent SHAP importance and PDP response ranges",
    "S6a. Source-isolation audit": "S6a. Source separation summary",
    "S9b. Representativeness": "Representativeness of the retained domain",
    "Note. All fields are presented in English. Source identity and post-treatment or target-derived quantities were never formal predictors.": "Note. Predictor eligibility was determined from information available before treatment. Source identity and post-treatment or target-derived fields were excluded.",
    "Note. Statistics are descriptive and do not redefine the locked prediction target.": "Note. Statistics describe the observations available for model development.",
    "Note. Sparse and single-source categorical levels were retained in the processed evidence.": "Note. Several categorical levels occur in only one literature source, which limits independent evaluation of their effects.",
    "Note. Configurations were recovered read-only from the original analysis scripts; no model was retrained.": "Note. Only explicitly specified settings are shown; software defaults are omitted.",
    "Note. The same five prespecified seeds were used for each method.": "Note. Each method was evaluated with seeds 7, 21, 42, 66, and 88.",
    "Note. Every validation split isolates literature sources. Strategy-specific results should not be conflated with random-split performance.": "Note. Each split separates literature sources between model fitting and evaluation.",
    "Note. Maximum source overlap was required to equal zero.": "Note. A value of zero confirms that no literature source appears on both sides of a split.",
    "Note. The rule uses target-independent AD scores; the retained evidence is source-specific.": "Note. The P08 rule uses applicability-domain scores calculated without the response variable.",
    "Note. For each strategy, XGBoost and the best-R² evaluated method are shown; this compact comparison is descriptive, not a new selection step.": "Note. The same XGBoost model family and P08 threshold are shown for every validation strategy; no strategy-specific best model was selected.",
    "Note. Fractions quantify the limited coverage and source concentration of the retained domain.": "Note. Coverage and source concentration indicate how narrowly the P08 subset represents the complete dataset.",
    "Note. Context labels define diagnostic subsets only; source identity was not a model feature.": "Note. Importance values are reported for three evaluation subsets. Literature-source identity was not a predictor.",
    "Note. Response extrema describe the fitted model over the evaluated grid and are not causal process optima.": "Note. Prediction ranges describe the fitted response surfaces within the evaluated grid and do not establish causal process optima.",
}


REMOVE_PARAGRAPHS = {
    "S9a. Numeric operating window",
    "Note. No records met this display condition.",
    "Note. The retained range is descriptive and is not a prospective universal specification.",
}


FIGURES = {
    "Figure S1": ("Data structure and predictor eligibility", "Completeness and analytical role of the fields in the source-bearing matrix. Panel a reports field completeness; panel b distinguishes pre-treatment predictors from fields reserved for response calculation, validation, or data checks."),
    "Figure S2": ("Continuous-variable coverage across literature sources", "Distributional support and literature-source coverage for continuous predictors and fluoride removal. Panel a shows the within-variable distributions; panel b gives the number of contributing sources and records. Normalization is used only for visualization."),
    "Figure S3": ("Random-split model stability across five seeds", "R², RMSE, and MAE for nine candidate methods across seeds 7, 21, 42, 66, and 88. The weighted four-model ensemble showed the strongest overall accuracy with limited between-seed variation."),
    "Figure S4": ("Source-isolated performance across validation strategies", "Pooled R², RMSE, and MAE under balanced grouped splitting, repeated grouped splitting, GroupKFold, and leave-one-literature-out validation. The decline relative to random splitting indicates substantial between-source heterogeneity."),
    "Figure S5": ("LOLO RMSE across held-out literature sources", "Leave-one-literature-out RMSE for each held-out source and model. Variation across rows identifies sources that are consistently easier or more difficult to predict and illustrates uneven transferability among literature sources."),
    "Figure S6": ("Sensitivity of selective-domain performance to the threshold", "Retained sample size, coverage, predictive performance, and observed-versus-predicted values as the nearest-neighbor risk threshold changes. The vertical marker denotes the P08 threshold of 0.814815."),
    "Figure S7": ("Stability and uncertainty within the retained domain", "Split-level R², bootstrap distributions, source composition, and representativeness for the P08-retained subset. Performance remains uncertain, and the retained observations are concentrated in Paper_010."),
    "Figure S8": ("Chemical conditions represented in the retained domain", "Continuous operating ranges and the most frequent categorical conditions in the P08-retained subset. These ranges describe the chemistry represented by the retained observations and are not universal operating recommendations."),
    "Figure S9": ("Transferability of the fixed P08 rule", "Performance and coverage after applying the P08 rule across validation strategies and model families. The rule does not transfer consistently, supporting a source-specific rather than universal interpretation."),
    "Figure S10": ("Context-dependent SHAP importance", "Feature ranks and grouped SHAP contributions for the full dataset, Paper_010, and the retained subset. Coagulant identity, dose, and interferent identity remain influential, while their ordering changes with evaluation context."),
    "Figure S11": ("Empirical support for two-dimensional PDP surfaces", "Two-dimensional partial-dependence surfaces for coagulant dose with initial fluoride and influent pH, paired with observed-data support. Sparsely occupied grid cells indicate regions where the fitted surfaces require cautious interpretation."),
}


STAGE_LABELS = {"19": "Random-split benchmark", "31g": "Grouped model screening", "27": "Out-of-fold stacking", "31l": "Grouped non-tree screening"}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def locate_input(project_root: Path) -> Path:
    candidates = [project_root / "outputs" / "jece_final_package" / "reports" / INPUT_NAME,
                  project_root / INPUT_NAME, Path(__file__).resolve().parent / INPUT_NAME, Path.cwd() / INPUT_NAME]
    found = [path for path in candidates if path.is_file()]
    if not found:
        raise FileNotFoundError(f"Could not find {INPUT_NAME}. Place it in the project root or outputs/jece_final_package/reports/.")
    return found[0]


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph_text(paragraph, text: str, bold_lead: str | None = None) -> None:
    clear_paragraph(paragraph)
    pieces = [(bold_lead, True), (text[len(bold_lead):], False)] if bold_lead and text.startswith(bold_lead) else [(text, False)]
    for value, bold in pieces:
        run = paragraph.add_run(value)
        run.bold = bold
        run.font.name = "Times New Roman"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")


def remove_paragraph(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def set_cell_text(cell, text: str, size: float = 6.2) -> None:
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
            run.font.size = Pt(size)


def concise_parameters(raw: str) -> str:
    try:
        values = json.loads(raw)
    except Exception:
        return re.sub(r"<expression[^>]*>", "set during fitting", raw).replace("_", " ")
    omit = {"n_jobs", "verbose", "allow_writing_files", "thread_count", "verbosity"}
    parts = []
    for key, value in values.items():
        if key in omit:
            continue
        if isinstance(value, str) and "<expression" in value:
            if key in {"random_state", "random_seed", "seed"}: value = "varied by split"
            elif key == "alphas": value = "log-spaced candidate grid"
            elif key == "n_components": value = "limited by available dimensions"
            else: value = "set during fitting"
        parts.append(f"{key.replace('_', ' ')}={value}")
    return "; ".join(parts) if parts else "Standard estimator specification"


def remove_cell(cell) -> None:
    cell._tc.getparent().remove(cell._tc)


def edit_configuration_table(table) -> None:
    headers = [cell.text.strip().lower() for cell in table.rows[0].cells]
    if headers[:3] != ["stage", "model", "parameters"]:
        raise ValueError("Table S3 was not found in the expected structure.")
    if len(headers) == 4:
        for row in table.rows:
            remove_cell(row.cells[-1])
        grid = table._tbl.tblGrid
        grid.remove(grid.gridCol_lst[-1])
    set_cell_text(table.rows[0].cells[0], "Analysis stage")
    for row in table.rows[1:]:
        stage = row.cells[0].text.strip().lower()
        set_cell_text(row.cells[0], STAGE_LABELS.get(stage, stage))
        set_cell_text(row.cells[1], row.cells[1].text.strip().replace("Regressor", "").strip())
        set_cell_text(row.cells[2], concise_parameters(row.cells[2].text.strip()))


def edit_transfer_table(table) -> None:
    headers = [cell.text.strip() for cell in table.rows[0].cells]
    if "Validation strategy" not in headers or "Method" not in headers:
        raise ValueError("Table S8 was not found in the expected structure.")
    method_index = headers.index("Method")
    for row in list(table.rows[1:]):
        if row.cells[method_index].text.strip() != "XGBoost":
            row._tr.getparent().remove(row._tr)


def update_paragraphs(doc: Document) -> None:
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if text in REMOVE_PARAGRAPHS:
            remove_paragraph(paragraph)
        elif text in TEXT_REPLACEMENTS:
            set_paragraph_text(paragraph, TEXT_REPLACEMENTS[text])
        elif text in SHORT_REPLACEMENTS:
            replacement = SHORT_REPLACEMENTS[text]
            set_paragraph_text(paragraph, replacement, "Note. " if replacement.startswith("Note. ") else None)
        else:
            for number, (title, caption) in FIGURES.items():
                if text.startswith(number + "."):
                    lead = f"{number}. {title}. "
                    set_paragraph_text(paragraph, lead + caption, lead)
                    break


def validate(doc: Document, input_hash: str, input_path: Path) -> None:
    visible = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    forbidden = ["locked evidence", "recovered read-only", "source file", "<expression>", ".py", "prespecified P08"]
    present = [term for term in forbidden if term.lower() in visible.lower()]
    if present: raise AssertionError("Internal workflow wording remains in v3: " + ", ".join(present))
    if re.search(r"[\u4e00-\u9fff]", visible): raise AssertionError("Chinese text remains in the English SI file.")
    if len(doc.inline_shapes) != 11: raise AssertionError(f"Expected 11 figures, found {len(doc.inline_shapes)}.")
    captions = [p.text for p in doc.paragraphs if re.match(r"Figure S\d+\.", p.text.strip())]
    if len(captions) != 11: raise AssertionError(f"Expected 11 figure captions, found {len(captions)}.")
    config = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Analysis stage")
    if len(config.columns) != 3: raise AssertionError("Table S3 still contains internal source-file information.")
    transfer = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Validation strategy" and len(t.columns) == 10)
    if any(row.cells[1].text.strip() != "XGBoost" for row in transfer.rows[1:]):
        raise AssertionError("Table S8 contains a strategy-specific alternative model.")
    if sha256(input_path) != input_hash: raise AssertionError("The v2 input changed during generation.")


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    reports = project_root / "outputs" / "jece_final_package" / "reports"
    reports.mkdir(parents=True, exist_ok=True)
    input_path = locate_input(project_root)
    output_path = reports / OUTPUT_NAME
    input_hash = sha256(input_path)
    doc = Document(input_path)
    if len(doc.inline_shapes) != 11 or len(doc.tables) < 15:
        raise ValueError("The selected v2 file does not match the completed 11-figure SI structure.")
    update_paragraphs(doc)
    edit_configuration_table(doc.tables[3])
    edit_transfer_table(doc.tables[11])
    core = doc.core_properties
    core.author = ""
    core.last_modified_by = ""
    core.comments = ""
    core.subject = "Supplementary Information for the Journal of Environmental Chemical Engineering"
    validate(doc, input_hash, input_path)
    doc.save(output_path)
    if not output_path.is_file() or output_path.stat().st_size == 0: raise IOError("The v3 Word file was not created.")
    print("PASS: JECE Supplementary Information v3 created.")
    print(f"Input preserved: {input_path}")
    print(f"Output: {output_path}")
    print("Supplementary Text: 5; Tables: 10; Figures: 11.")
    print("Table S3 internal file references removed; Table S8 restricted to fixed XGBoost comparisons.")
    print(f"Total runtime: {time.perf_counter() - RUN_START:.2f} s")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        print(f"Total runtime: {time.perf_counter() - RUN_START:.2f} s")
        raise
