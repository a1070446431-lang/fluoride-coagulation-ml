# -*- coding: utf-8 -*-
"""
32W EST-style main figure suite + supplementary figure package
Project: fluoride coagulation defluoridation ML project

What this script does
---------------------
1) Rebuilds the manuscript figure system into EST-style composite figures.
2) Each main figure is limited to exactly 6 panels: A-F, 2 x 3 layout.
3) Removes redundant sentence-style labels below each subfigure. Detailed explanation is left to the Word caption.
4) Creates a supplementary-material Word document containing the remaining project figures in project order.
5) Keeps output locations unchanged:
   - outputs/figures/
   - outputs/tables/
   - reports/

Important manuscript boundary
-----------------------------
R2 = 0.747676 belongs only to the target-free selective AD retained-subdomain result:
RepeatedBalanced_Extended + XGBoost; Zone=Outside; NN_Risk <= 0.814815.
It must not be described as all-domain grouped-validation performance.
"""

from __future__ import annotations

import csv
import math
import re
import sys
import time
from pathlib import Path
from typing import Dict, List, Optional

try:
    from PIL import Image, ImageDraw, ImageFont, ImageOps, ImageChops
except Exception as e:
    print("[ERROR] Pillow is required. Install it by: pip install pillow")
    print(f"Original error: {e}")
    sys.exit(1)

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except Exception as e:
    print("[ERROR] python-docx is required. Install it by: pip install python-docx")
    print(f"Original error: {e}")
    sys.exit(1)


# --------------------------
# Basic utilities
# --------------------------

def now_str() -> str:
    return time.strftime("%Y-%m-%d %H:%M:%S")


def natural_key(text: str):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", text)]


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def write_csv(path: Path, rows: List[Dict]) -> None:
    ensure_dir(path.parent)
    if not rows:
        rows = [{"empty": ""}]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def prefix_number(path: Optional[Path]) -> Optional[int]:
    if path is None:
        return None
    m = re.match(r"^(\d+)", path.name)
    return int(m.group(1)) if m else None


def find_project_root(start: Path) -> Path:
    start = start.resolve()
    candidates = [start] + list(start.parents)
    for base in candidates:
        if (base / "outputs" / "figures").exists() and (base / "outputs" / "tables").exists():
            return base
    for base in candidates[:3]:
        try:
            for p in base.rglob("outputs"):
                if (p / "figures").exists() and (p / "tables").exists():
                    return p.parent
        except Exception:
            pass
    raise FileNotFoundError("Cannot find project root containing outputs/figures and outputs/tables.")


def resolve_template(project_root: Path) -> Optional[Path]:
    names = [
        "AI论文--混凝沉淀除氟初稿.docx",
        "AI论文--混凝沉淀除氟空稿.docx",
        "32t2_full_manuscript_intro_results_random_grouped_ad.docx",
        "32v2_est_style_expanded_storyline_figure_suite_inserted.docx",
    ]
    search_dirs = [
        project_root,
        project_root / "reports",
        project_root.parent,
        project_root.parent / "reports",
        Path.cwd(),
    ]
    for d in search_dirs:
        if not d.exists():
            continue
        for name in names:
            p = d / name
            if p.exists():
                return p
    return None


def list_pngs(fig_dir: Path) -> List[Path]:
    if not fig_dir.exists():
        return []
    return sorted([p for p in fig_dir.iterdir() if p.suffix.lower() == ".png"], key=lambda p: natural_key(p.name))


def pick_figure(fig_dir: Path, prefixes: List[int], keywords: Optional[List[str]] = None) -> Optional[Path]:
    keywords = keywords or []
    pngs = list_pngs(fig_dir)
    for num in prefixes:
        matches = [p for p in pngs if prefix_number(p) == num]
        if matches:
            return matches[0]
    low_keywords = [k.lower() for k in keywords]
    for p in pngs:
        name = p.name.lower()
        if low_keywords and all(k in name for k in low_keywords):
            return p
    for p in pngs:
        name = p.name.lower()
        if low_keywords and any(k in name for k in low_keywords):
            return p
    return None


def get_font(size: int, bold: bool = False):
    candidates = []
    if bold:
        candidates += [
            "arialbd.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
        ]
    candidates += [
        "arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for c in candidates:
        try:
            return ImageFont.truetype(c, size=size)
        except Exception:
            continue
    return ImageFont.load_default()


def trim_white(img: Image.Image, border: int = 12) -> Image.Image:
    rgb = img.convert("RGB")
    bg = Image.new("RGB", rgb.size, "white")
    diff = ImageChops.difference(rgb, bg)
    bbox = diff.getbbox()
    if not bbox:
        return rgb
    l, t, r, b = bbox
    l = max(0, l - border)
    t = max(0, t - border)
    r = min(rgb.width, r + border)
    b = min(rgb.height, b + border)
    return rgb.crop((l, t, r, b))


def safe_image(path: Optional[Path], label: str) -> Image.Image:
    if path and path.exists():
        try:
            img = Image.open(path)
            img = ImageOps.exif_transpose(img).convert("RGB")
            return trim_white(img)
        except Exception:
            pass
    img = Image.new("RGB", (1200, 900), "white")
    d = ImageDraw.Draw(img)
    d.rectangle((20, 20, 1180, 880), outline=(150, 150, 150), width=4)
    d.text((70, 120), f"Missing source panel: {label}", fill=(120, 120, 120), font=get_font(44, True))
    return img


# --------------------------
# Figure construction
# --------------------------

def build_composite(title: str, panels: List[Dict], out_png: Path, out_pdf: Path) -> None:
    """Build a clean 2 x 3 EST-style composite figure."""
    ensure_dir(out_png.parent)
    W, H = 4200, 3000
    margin = 110
    top_title_h = 150
    gap_x, gap_y = 80, 90
    ncols, nrows = 3, 2
    header_h = 70

    canvas = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(canvas)

    title_font = get_font(58, True)
    label_font = get_font(52, True)
    header_font = get_font(30, True)

    tb = draw.textbbox((0, 0), title, font=title_font)
    draw.text(((W - (tb[2] - tb[0])) / 2, 35), title, fill="black", font=title_font)

    panel_w = int((W - 2 * margin - (ncols - 1) * gap_x) / ncols)
    panel_h = int((H - margin - top_title_h - margin - (nrows - 1) * gap_y) / nrows)

    for i, p in enumerate(panels[:6]):
        row, col = divmod(i, ncols)
        x0 = margin + col * (panel_w + gap_x)
        y0 = top_title_h + row * (panel_h + gap_y)
        x1 = x0 + panel_w
        y1 = y0 + panel_h

        draw.rectangle((x0, y0, x1, y1), outline=(95, 95, 95), width=2)
        draw.rectangle((x0, y0, x1, y0 + header_h), fill=(246, 246, 246), outline=(95, 95, 95), width=1)
        draw.text((x0 + 16, y0 + 9), p["label"], fill="black", font=label_font)
        draw.text((x0 + 85, y0 + 20), p["header"], fill="black", font=header_font)

        content_x0 = x0 + 22
        content_y0 = y0 + header_h + 18
        content_w = panel_w - 44
        content_h = panel_h - header_h - 36

        src = safe_image(p.get("path"), p["label"])
        src.thumbnail((content_w, content_h), Image.Resampling.LANCZOS)
        px = content_x0 + (content_w - src.width) // 2
        py = content_y0 + (content_h - src.height) // 2
        canvas.paste(src, (px, py))

    canvas.save(out_png, dpi=(300, 300))
    canvas.save(out_pdf, resolution=300.0)


# --------------------------
# Word helpers
# --------------------------

def set_doc_style(doc: Document) -> None:
    try:
        st = doc.styles["Normal"]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        st.font.size = Pt(11)
    except Exception:
        pass


def add_centered_picture(doc: Document, fig_path: Path, caption: str, width_inches: float = 6.4) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(fig_path), width=Inches(width_inches))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(10.5)


def create_main_doc(template: Optional[Path], figures: List[Dict], out_docx: Path) -> None:
    doc = Document(str(template)) if template and template.exists() else Document()
    set_doc_style(doc)
    doc.add_page_break()
    doc.add_heading("EST-style Main Figure Suite（正文主图套图）", level=1)
    note = doc.add_paragraph(
        "以下主图采用 2×3 复合图结构，每张图仅保留 6 个关键子图，子图使用 A–F 面板标识。"
        "其余过程图按项目顺序整理至补充材料文档。"
    )
    note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for fig in figures:
        add_centered_picture(doc, fig["path"], fig["caption"], width_inches=6.5)
    ensure_dir(out_docx.parent)
    doc.save(str(out_docx))


def create_supplement_doc(project_root: Path, used_prefixes: List[int], out_docx: Path, manifest_rows: List[Dict]) -> None:
    fig_dir = project_root / "outputs" / "figures"
    all_pngs = list_pngs(fig_dir)
    used = set(used_prefixes)
    supp = []
    for p in all_pngs:
        n = prefix_number(p)
        if n is not None and n not in used and n < 300:
            supp.append(p)

    doc = Document()
    set_doc_style(doc)
    doc.add_heading("Supplementary Figures（补充材料图件）", level=1)
    doc.add_paragraph("本文件按项目图件编号顺序保存未进入正文主图的过程性图件，用于支撑随机划分、模型调参、分组验证、适应域和稳定性审计等完整证据链。")

    for i, p in enumerate(supp, 1):
        cap = f"Figure S{i}. Project figure {p.name}."
        doc.add_paragraph(cap).alignment = WD_ALIGN_PARAGRAPH.CENTER
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            para.add_run().add_picture(str(p), width=Inches(6.2))
            status = "inserted"
        except Exception as e:
            para.add_run(f"[Figure insertion failed: {p.name}; {e}]")
            status = "failed"
        manifest_rows.append({
            "supplement_index": i,
            "source_file": p.name,
            "source_prefix": prefix_number(p),
            "status": status,
            "caption": cap,
        })
    ensure_dir(out_docx.parent)
    doc.save(str(out_docx))


# --------------------------
# Main
# --------------------------

def main():
    t0 = time.time()
    start = Path.cwd()
    project_root = find_project_root(start)
    fig_dir = project_root / "outputs" / "figures"
    table_dir = project_root / "outputs" / "tables"
    report_dir = project_root / "reports"
    ensure_dir(fig_dir)
    ensure_dir(table_dir)
    ensure_dir(report_dir)

    template = resolve_template(project_root)

    print("=" * 72)
    print("32W EST-style main and supplementary figure suite")
    print(f"Run time: {now_str()}")
    print(f"Project root: {project_root}")
    print(f"Figure dir: {fig_dir}")
    print(f"Table dir: {table_dir}")
    print(f"Report dir: {report_dir}")
    print(f"Template/draft: {template if template else 'None'}")
    print("=" * 72)

    configs = [
        {
            "num": 301,
            "title": "Data landscape and modeling challenge",
            "caption": "图 3-1 数据景观与建模挑战。该图以 6 个关键面板概括文献来源、目标分布、变量结构、机理特征和后续建模问题。",
            "panels": [
                ("A", "Literature groups", [69], ["paper", "group"]),
                ("B", "Target distribution", [70], ["target", "distribution"]),
                ("C", "Variable landscape", [21], ["distribution", "figure1"]),
                ("D", "Mechanism features", [63], ["feature"]),
                ("E", "Correlation/nonlinearity", [22, 23], ["correlation", "heatmap"]),
                ("F", "Workflow anchor", [201, 205, 211], ["storyline", "framework"]),
            ],
        },
        {
            "num": 302,
            "title": "Random-split modeling to two-model fusion",
            "caption": "图 3-2 随机划分建模与双模型形成过程。该图展示从基准模型筛选、CatBoost 调参、模型融合、权重稳定性到 CatBoost-HGB 加权双模型的形成证据；随机划分最高 R² 约为 0.825–0.829。",
            "panels": [
                ("A", "R² milestones", [206, 212], ["random", "twomodel"]),
                ("B", "Baseline models", [64], ["model", "comparison"]),
                ("C", "Fusion metrics", [65], ["fusion", "metrics"]),
                ("D", "Weight stability", [99, 46, 33], ["weight", "stability"]),
                ("E", "Prediction accuracy", [111, 47, 52], ["predicted", "actual"]),
                ("F", "Ablation support", [67, 68], ["ablation"]),
            ],
        },
        {
            "num": 303,
            "title": "Grouped validation and transfer boundary",
            "caption": "图 3-3 分组验证与跨文献泛化边界。该图同时给出探索性全域分组最优结果约 R² = 0.25 与最终审计全域基准 R² = 0.194，并展示训练-测试分布偏移、分组预测偏差和文献层级误差。",
            "panels": [
                ("A", "Metric hierarchy", [207, 213], ["grouped", "boundary"]),
                ("B", "Train-test shift", [72], ["shift"]),
                ("C", "All-domain metrics", [80], ["grouped", "metrics"]),
                ("D", "Grouped prediction", [81], ["predicted", "observed"]),
                ("E", "Literature errors", [129], ["leave", "error"]),
                ("F", "Transfer boundary", [82, 83], ["boundary"]),
            ],
        },
        {
            "num": 304,
            "title": "Applicability domain and retained-subdomain evidence",
            "caption": "图 3-4 适应域发展与最终保留子域证据。该图展示适应域风险诊断、预定义可靠域不足、选择性适应域证据链、阈值平台、化学窗口和稳定性边界。R² = 0.748 仅代表选择性适应域保留子域。",
            "panels": [
                ("A", "AD risk vs error", [139, 214], ["risk", "error"]),
                ("B", "Reliable-domain tradeoff", [153], ["reliable", "domain"]),
                ("C", "Evidence chain", [194], ["evidence", "summary"]),
                ("D", "Threshold plateau", [195], ["threshold", "plateau"]),
                ("E", "Chemistry window", [196], ["chemistry", "window"]),
                ("F", "Stability boundary", [197], ["stability", "boundary"]),
            ],
        },
    ]

    main_rows = []
    figures = []
    used_prefixes = []

    for cfg in configs:
        panels = []
        for label, header, prefixes, keywords in cfg["panels"]:
            src = pick_figure(fig_dir, prefixes, keywords)
            n = prefix_number(src)
            if n is not None:
                used_prefixes.append(n)
            panels.append({"label": label, "header": header, "path": src})
            main_rows.append({
                "main_figure": cfg["num"],
                "panel": label,
                "header": header,
                "source_file": src.name if src else "NOT_FOUND",
                "source_prefix": n if n is not None else "",
                "found": bool(src),
            })
        out_png = fig_dir / f"{cfg['num']}_32w_est_main_{cfg['num']}.png"
        out_pdf = fig_dir / f"{cfg['num']}_32w_est_main_{cfg['num']}.pdf"
        build_composite(cfg["title"], panels, out_png, out_pdf)
        figures.append({"num": cfg["num"], "path": out_png, "caption": cfg["caption"]})
        print(f"[OUTPUT FIGURE] {out_png}")

    main_doc = report_dir / "32w_est_style_main_figures_inserted.docx"
    create_main_doc(template, figures, main_doc)

    supp_rows: List[Dict] = []
    supp_doc = report_dir / "32w_supplementary_figures.docx"
    create_supplement_doc(project_root, sorted(set(used_prefixes)), supp_doc, supp_rows)

    write_csv(table_dir / "32w_main_figure_panel_manifest.csv", main_rows)
    write_csv(table_dir / "32w_supplementary_figure_manifest.csv", supp_rows)
    write_csv(table_dir / "32w_metric_claim_boundary.csv", [
        {
            "claim": "R2=0.747676",
            "correct_scope": "target-free selective AD retained-subdomain only",
            "model_rule": "RepeatedBalanced_Extended + XGBoost; Zone=Outside; NN_Risk <= 0.814815",
            "not_allowed": "all-domain grouped validation performance",
        }
    ])

    log = [
        "32W EST-style main and supplementary figure suite",
        f"Run time: {now_str()}",
        f"Project root: {project_root}",
        f"Figure dir: {fig_dir}",
        f"Template/draft: {template if template else 'None'}",
        f"Main docx: {main_doc}",
        f"Supplement docx: {supp_doc}",
        f"Main figure count: {len(figures)}",
        f"Supplement figure count: {len(supp_rows)}",
        "Rule: each main figure has exactly six panels; other figures are moved to supplementary material.",
        f"Total runtime: {time.time() - t0:.2f} s",
    ]
    (report_dir / "32w_run_log.txt").write_text("\n".join(log), encoding="utf-8-sig")

    print("=" * 72)
    print("32W finished successfully.")
    print(f"[OUTPUT MAIN DOCX] {main_doc}")
    print(f"[OUTPUT SUPPLEMENT DOCX] {supp_doc}")
    print(f"[OUTPUT TABLE] {table_dir / '32w_main_figure_panel_manifest.csv'}")
    print(f"[OUTPUT TABLE] {table_dir / '32w_supplementary_figure_manifest.csv'}")
    print(f"[OUTPUT LOG] {report_dir / '32w_run_log.txt'}")
    print(f"Total runtime: {time.time() - t0:.2f} s")


if __name__ == "__main__":
    main()
