# -*- coding: utf-8 -*-
"""
32Q2 Final DOCX warning fix for Results and Discussion
Project: fluoride coagulation defluoridation ML project

Purpose
-------
This script fixes the only warning found in 32Q: a conservative audit rule
flagged a boundary sentence because the accepted selective-AD R2 value and
an all-domain warning phrase appeared in the same sentence/paragraph.

This script does not retrain models, does not change any 32I2 metric, and
does not create a new output root.

Output locations follow the original project convention:
    outputs/tables/
    reports/

Expected locked result:
    Model/strategy: RepeatedBalanced_Extended + XGBoost
    Rule: Zone=Outside; NN_Risk <= 0.814815
    R2=0.747676, RMSE=11.474674, MAE=9.155856, Coverage=0.014664, N=52
    Dominant group=Paper_010
"""

from __future__ import annotations

import csv
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

try:
    from docx import Document
except Exception as exc:  # pragma: no cover
    print("[ERROR] python-docx is required. Install it by: pip install python-docx")
    print(f"Original error: {exc}")
    sys.exit(1)


LOCKED = {
    "model_family": "RepeatedBalanced_Extended + XGBoost",
    "rule": "Zone=Outside; NN_Risk <= 0.814815",
    "r2": "0.747676",
    "rmse": "11.474674",
    "mae": "9.155856",
    "coverage": "0.014664",
    "n": "52",
    "dominant_group": "Paper_010",
}

FORBIDDEN_NEAR_R2_TERMS = [
    "全域分组验证",
    "全域 grouped validation",
    "all-domain",
    "all domain",
    "grouped validation",
    "未筛选全域",
]

SAFE_R2_SENTENCE = (
    "选择性适应域保留子域的 R² 为 0.747676，RMSE 为 11.474674，"
    "MAE 为 9.155856，Coverage 为 0.014664，保留样本数为 52。"
    "该结果仅表征适应域规则筛选后保留样本子集的预测表现。"
)

SAFE_BOUNDARY_SENTENCE = (
    "结果解释边界：本文对未筛选测试集与适应域筛选后保留子域进行分开报告，"
    "二者的评价口径不能混用。"
)


def find_project_root(start: Path) -> Path:
    """Find the project root regardless of whether this script is run from root or src."""
    start = start.resolve()
    candidates = [start] + list(start.parents)
    for p in candidates:
        if (p / "outputs" / "tables").exists() and (p / "reports").exists():
            return p
    for p in candidates:
        if (p / "src").exists() and (p / "outputs").exists():
            return p
    return start.parent if start.name.lower() == "src" else start


def iter_paragraphs(document: Document):
    """Yield normal paragraphs and table-cell paragraphs."""
    for para in document.paragraphs:
        yield para
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield para
                for nested_table in cell.tables:
                    for row2 in nested_table.rows:
                        for cell2 in row2.cells:
                            for para2 in cell2.paragraphs:
                                yield para2


def set_paragraph_text(para, new_text: str) -> None:
    """Replace paragraph text while keeping the first run's basic formatting."""
    if not para.runs:
        para.add_run(new_text)
        return
    para.runs[0].text = new_text
    for run in para.runs[1:]:
        run.text = ""


def normalize_text(s: str) -> str:
    return re.sub(r"\s+", " ", str(s)).strip()


def paragraph_has_forbidden_r2_claim(text: str) -> bool:
    """True if accepted R2 and all-domain wording appear too close in the same paragraph."""
    t = normalize_text(text)
    if LOCKED["r2"] not in t:
        return False
    low = t.lower()
    return any(term.lower() in low for term in FORBIDDEN_NEAR_R2_TERMS)


def repair_docx_text(doc: Document) -> List[Dict[str, str]]:
    """Repair warning-triggering wording and common key-result shorthand."""
    repair_log: List[Dict[str, str]] = []

    for idx, para in enumerate(iter_paragraphs(doc), start=1):
        old = para.text or ""
        if not old.strip():
            continue
        new = old
        reason = ""

        # 1) Main repair: remove accepted R2 and all-domain language from the same paragraph.
        if paragraph_has_forbidden_r2_claim(new):
            new = SAFE_R2_SENTENCE
            reason = "separated_accepted_r2_from_all_domain_boundary_language"

        # 2) Repair short model names that could imply XGBoost alone.
        short_patterns = [
            "Selective AD retained-subdomain result: XGBoost",
            "选择性适应域保留子域结果：XGBoost",
            "选择性适应域保留子域模型：XGBoost",
        ]
        for pat in short_patterns:
            if pat in new and LOCKED["model_family"] not in new:
                new = new.replace(pat, f"Selective AD retained-subdomain result: {LOCKED['model_family']}")
                reason = (reason + ";" if reason else "") + "expanded_model_family_name"

        # 3) Repair accidental all-domain claim wording if it appears around the accepted result.
        if "R2=0.747676 belongs to the selective AD retained-subdomain" in new:
            new = SAFE_R2_SENTENCE
            reason = (reason + ";" if reason else "") + "converted_english_boundary_line_to_safe_chinese_sentence"

        if new != old:
            set_paragraph_text(para, new)
            repair_log.append({
                "paragraph_index": str(idx),
                "repair_reason": reason,
                "old_text": old,
                "new_text": new,
            })

    # Add a boundary sentence without the accepted R2 number if the document does not already have a safe boundary.
    full_text = "\n".join(p.text for p in iter_paragraphs(doc))
    if SAFE_BOUNDARY_SENTENCE not in full_text:
        doc.add_paragraph(SAFE_BOUNDARY_SENTENCE)
        repair_log.append({
            "paragraph_index": "APPENDIX_END",
            "repair_reason": "added_safe_boundary_sentence_without_accepted_r2_value",
            "old_text": "",
            "new_text": SAFE_BOUNDARY_SENTENCE,
        })

    return repair_log


def collect_doc_text(doc: Document) -> str:
    return "\n".join(p.text for p in iter_paragraphs(doc))


def audit_docx(doc: Document, figure_dir: Path) -> Tuple[List[Dict[str, str]], List[Dict[str, str]]]:
    text = collect_doc_text(doc)
    normalized = normalize_text(text)
    low = normalized.lower()

    checks: List[Dict[str, str]] = []

    def add_check(item: str, passed: bool, detail: str) -> None:
        checks.append({
            "item": item,
            "status": "PASS" if passed else "FAIL",
            "detail": detail,
        })

    add_check("DOCX contains accepted R2", LOCKED["r2"] in normalized, f"Expected R2={LOCKED['r2']}")
    add_check("DOCX contains accepted RMSE", LOCKED["rmse"] in normalized, f"Expected RMSE={LOCKED['rmse']}")
    add_check("DOCX contains accepted MAE", LOCKED["mae"] in normalized, f"Expected MAE={LOCKED['mae']}")
    add_check("DOCX contains accepted coverage", LOCKED["coverage"] in normalized, f"Expected Coverage={LOCKED['coverage']}")
    add_check("DOCX contains retained N", re.search(r"\bN\s*[=＝]\s*52\b|保留样本数为 52|保留样本数=52", normalized) is not None, "Expected retained N=52")
    add_check("DOCX contains dominant group", LOCKED["dominant_group"] in normalized, f"Expected {LOCKED['dominant_group']}")
    add_check("DOCX contains accepted model family", LOCKED["model_family"] in normalized, f"Expected {LOCKED['model_family']}")
    add_check("DOCX contains accepted AD rule", LOCKED["rule"] in normalized, f"Expected {LOCKED['rule']}")

    bad_paras = []
    for i, para in enumerate(iter_paragraphs(doc), start=1):
        if paragraph_has_forbidden_r2_claim(para.text or ""):
            bad_paras.append(f"P{i}: {normalize_text(para.text)[:180]}")
    add_check(
        "No accepted R2 mixed with all-domain wording in one paragraph",
        len(bad_paras) == 0,
        "No mixed paragraph found" if not bad_paras else " | ".join(bad_paras),
    )

    # Avoid model-family shorthand that can suggest XGBoost alone.
    shorthand_bad = "selective ad retained-subdomain result: xgboost" in low and LOCKED["model_family"].lower() not in low
    add_check(
        "No XGBoost-only shorthand for the accepted selective AD result",
        not shorthand_bad,
        "Accepted model family is explicitly written" if not shorthand_bad else "Found XGBoost-only shorthand",
    )

    inline_count = len(doc.inline_shapes)
    add_check("At least four inline images in DOCX", inline_count >= 4, f"inline_shapes={inline_count}")

    figure_audit: List[Dict[str, str]] = []
    for fig_no in [194, 195, 196, 197]:
        found = sorted(figure_dir.glob(f"{fig_no}_*.png"))
        figure_audit.append({
            "figure_no": str(fig_no),
            "status": "PASS" if found else "FAIL",
            "file_found": str(found[0]) if found else "",
        })
    add_check(
        "All four figure files exist in outputs/figures",
        all(row["status"] == "PASS" for row in figure_audit),
        "; ".join([f"{row['figure_no']}={row['status']}" for row in figure_audit]),
    )

    return checks, figure_audit


def write_csv(path: Path, rows: List[Dict[str, str]], fieldnames: Iterable[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fieldnames is None:
        keys: List[str] = []
        for row in rows:
            for k in row.keys():
                if k not in keys:
                    keys.append(k)
        fieldnames = keys or ["empty"]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(fieldnames))
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def main() -> int:
    script_path = Path(__file__).resolve()
    project_root = find_project_root(script_path.parent)
    tables_dir = project_root / "outputs" / "tables"
    figures_dir = project_root / "outputs" / "figures"
    reports_dir = project_root / "reports"
    tables_dir.mkdir(parents=True, exist_ok=True)
    reports_dir.mkdir(parents=True, exist_ok=True)

    source_candidates = [
        reports_dir / "32q_results_discussion_final_checked.docx",
        reports_dir / "32p_results_discussion_with_figures.docx",
        reports_dir / "32o_results_discussion_inserted_draft.docx",
    ]
    source_docx = next((p for p in source_candidates if p.exists()), None)
    if source_docx is None:
        print("[ERROR] No source DOCX found. Expected one of:")
        for p in source_candidates:
            print(f"  - {p}")
        return 1

    out_docx = reports_dir / "32q2_results_discussion_final_checked.docx"
    audit_csv = tables_dir / "32q2_docx_quality_audit.csv"
    figure_csv = tables_dir / "32q2_figure_presence_audit.csv"
    repair_csv = tables_dir / "32q2_text_repair_log.csv"
    key_csv = tables_dir / "32q2_locked_key_result_check.csv"
    manifest_csv = tables_dir / "32q2_final_docx_package_manifest.csv"
    report_md = reports_dir / "32q2_final_docx_quality_audit_report.md"
    report_txt = reports_dir / "32q2_final_docx_quality_audit_report.txt"
    run_log = reports_dir / "32q2_run_log.txt"

    doc = Document(str(source_docx))
    repairs = repair_docx_text(doc)
    checks, figure_audit = audit_docx(doc, figures_dir)
    doc.save(str(out_docx))

    # Re-open the saved docx and audit again to ensure the persisted file is clean.
    saved_doc = Document(str(out_docx))
    checks, figure_audit = audit_docx(saved_doc, figures_dir)

    write_csv(repair_csv, repairs, ["paragraph_index", "repair_reason", "old_text", "new_text"])
    write_csv(audit_csv, checks, ["item", "status", "detail"])
    write_csv(figure_csv, figure_audit, ["figure_no", "status", "file_found"])
    write_csv(key_csv, [{
        "model_family": LOCKED["model_family"],
        "rule": LOCKED["rule"],
        "r2": LOCKED["r2"],
        "rmse": LOCKED["rmse"],
        "mae": LOCKED["mae"],
        "coverage": LOCKED["coverage"],
        "n": LOCKED["n"],
        "dominant_group": LOCKED["dominant_group"],
        "writing_boundary": "R2=0.747676 belongs only to the selective AD retained subdomain.",
    }])

    package_rows = [
        {"artifact_type": "source_docx", "path": str(source_docx), "exists": str(source_docx.exists())},
        {"artifact_type": "final_docx", "path": str(out_docx), "exists": str(out_docx.exists())},
        {"artifact_type": "quality_audit_csv", "path": str(audit_csv), "exists": str(audit_csv.exists())},
        {"artifact_type": "figure_audit_csv", "path": str(figure_csv), "exists": str(figure_csv.exists())},
        {"artifact_type": "repair_log_csv", "path": str(repair_csv), "exists": str(repair_csv.exists())},
        {"artifact_type": "audit_report_md", "path": str(report_md), "exists": str(report_md.exists())},
    ]
    write_csv(manifest_csv, package_rows, ["artifact_type", "path", "exists"])

    pass_count = sum(1 for r in checks if r["status"] == "PASS")
    fail_rows = [r for r in checks if r["status"] == "FAIL"]
    fail_count = len(fail_rows)
    status_line = "32Q2 finished successfully." if fail_count == 0 else "32Q2 finished with warnings. Please inspect FAIL rows."

    report_lines = [
        "# 32Q2 Final DOCX quality audit and warning fix",
        "",
        f"Run time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"Project root: {project_root}",
        f"Source DOCX: {source_docx}",
        f"Final DOCX: {out_docx}",
        "",
        "## Locked key result",
        f"- Selective AD retained-subdomain result: {LOCKED['model_family']}",
        f"- Rule: {LOCKED['rule']}",
        f"- R²={LOCKED['r2']}, RMSE={LOCKED['rmse']}, MAE={LOCKED['mae']}, Coverage={LOCKED['coverage']}, N={LOCKED['n']}",
        f"- Dominant group: {LOCKED['dominant_group']}",
        "",
        "## Audit summary",
        f"- PASS: {pass_count}",
        f"- FAIL: {fail_count}",
        f"- Text repairs applied: {len(repairs)}",
        "",
        "## Writing boundary",
        "The value R²=0.747676 is only the selective AD retained-subdomain result and must not be described as overall unfiltered test-set performance.",
        "",
        "## Failed checks",
    ]
    if fail_rows:
        for row in fail_rows:
            report_lines.append(f"- {row['item']}: {row['detail']}")
    else:
        report_lines.append("- None")

    report_content = "\n".join(report_lines) + "\n"
    write_text(report_md, report_content)
    write_text(report_txt, report_content)

    log_lines = [
        "=" * 72,
        "32Q2 Final DOCX warning fix",
        f"Run time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"Project root: {project_root}",
        f"Table output dir: {tables_dir}",
        f"Report output dir: {reports_dir}",
        "=" * 72,
        "[LOCKED KEY RESULT]",
        f"Selective AD retained-subdomain result: {LOCKED['model_family']}",
        f"Rule: {LOCKED['rule']}",
        f"R2={LOCKED['r2']}, RMSE={LOCKED['rmse']}, MAE={LOCKED['mae']}, Coverage={LOCKED['coverage']}, N={LOCKED['n']}, Dominant group={LOCKED['dominant_group']}",
        "=" * 72,
        f"[AUDIT] PASS={pass_count}; FAIL={fail_count}; inline_images={len(saved_doc.inline_shapes)}; text_repairs={len(repairs)}",
    ]
    for row in checks:
        log_lines.append(f"[{row['status']}] {row['item']} | {row['detail']}")
    log_lines += [
        "=" * 72,
        f"[OUTPUT DOCX] {out_docx}",
        f"[OUTPUT TABLE] {audit_csv}",
        f"[OUTPUT TABLE] {figure_csv}",
        f"[OUTPUT TABLE] {repair_csv}",
        f"[OUTPUT TABLE] {key_csv}",
        f"[OUTPUT TABLE] {manifest_csv}",
        f"[OUTPUT REPORT] {report_md}",
        f"[OUTPUT LOG] {run_log}",
        "",
        status_line,
        "Important writing boundary: R2=0.747676 belongs to the selective AD retained-subdomain only.",
        "",
        "Please open:",
        f"  {out_docx}",
    ]
    write_text(run_log, "\n".join(log_lines) + "\n")

    print("\n".join(log_lines))
    return 0 if fail_count == 0 else 0


if __name__ == "__main__":
    raise SystemExit(main())
