# -*- coding: utf-8 -*-
"""Build a concise, submission-ready English Supplementary Information Word file.

Place this complete script in ``fluoride_coagulation_ml_project/src/`` and run
it from PyCharm. It reads the locked SI evidence already stored under
``outputs/jece_final_package`` and the latest main-manuscript Word file found
in the project root. It does not train models, alter evidence, or overwrite v1.

Primary output
--------------
outputs/jece_final_package/reports/Supplementary_Information_Paper1_JECE_v2.docx

Dependencies: Python >= 3.9, pandas, python-docx, Pillow.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

import pandas as pd
from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


STEM = "SI_build_submission_ready_word_v2"
OUTPUT_DOCX = "Supplementary_Information_Paper1_JECE_v2.docx"
MANUSCRIPT_OVERRIDE: str | None = None
RUN_START = time.perf_counter()

BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
BLACK = RGBColor(0, 0, 0)
A4_W, A4_H = Cm(21.0), Cm(29.7)
MARGIN_P, MARGIN_L = Cm(1.65), Cm(1.35)


@dataclass(frozen=True)
class Paths:
    root: Path
    config: Path
    figures: Path
    tables: Path
    reports: Path
    logs: Path
    manuscript: Path
    output_docx: Path
    audit: Path
    log: Path


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def require(path: Path, purpose: str) -> Path:
    if not path.is_file():
        raise FileNotFoundError(f"Missing {purpose}: {path}")
    return path


def extract_english_title(manuscript: Path) -> str:
    doc = Document(manuscript)
    candidates: list[tuple[float, str]] = []
    keywords = ("fluor", "coagul", "predict", "applicability", "data-driven", "model")
    for i, para in enumerate(doc.paragraphs[:120]):
        text = " ".join(para.text.split()).strip(" -–—")
        words = re.findall(r"[A-Za-z][A-Za-z\-']*", text)
        if not (7 <= len(words) <= 35) or len(text) > 300:
            continue
        ascii_ratio = sum(ord(c) < 128 for c in text) / max(len(text), 1)
        if ascii_ratio < 0.92:
            continue
        score = 5 * sum(k in text.lower() for k in keywords) - 0.03 * i
        if para.style and "title" in para.style.name.lower():
            score += 8
        candidates.append((score, text))
    if not candidates or max(candidates)[0] < 10:
        raise ValueError(f"Could not identify the English article title in {manuscript.name}.")
    return max(candidates, key=lambda x: x[0])[1]


def find_latest_manuscript(root: Path) -> Path:
    if MANUSCRIPT_OVERRIDE:
        candidate = Path(MANUSCRIPT_OVERRIDE)
        if not candidate.is_absolute():
            candidate = root / candidate
        require(candidate, "main manuscript Word file")
        extract_english_title(candidate)
        return candidate
    ranked: list[tuple[int, float, Path]] = []
    for folder in (root, root.parent):
        for p in folder.glob("*.docx"):
            low = p.name.lower()
            if "supplement" in low or "si_" in low or p.name.startswith("~$"):
                continue
            score = 0
            if "图4文字更新" in p.name:
                score += 100
            if "总版3" in p.name or "总版(3)" in p.name:
                score += 50
            if "总版" in p.name:
                score += 20
            try:
                extract_english_title(p)
            except Exception:
                continue
            ranked.append((score, p.stat().st_mtime, p))
    if not ranked:
        raise FileNotFoundError(
            "No main-manuscript Word file containing a recognizable English title was found. "
            "Copy the latest manuscript to the project root or set MANUSCRIPT_OVERRIDE."
        )
    return max(ranked, key=lambda x: (x[0], x[1]))[2]


def project_paths() -> Paths:
    root = Path(__file__).resolve().parents[1]
    config = require(root / "src" / "config.py", "project configuration")
    base = root / "outputs" / "jece_final_package"
    figures, tables, reports, logs = [base / x for x in ("figures", "tables", "reports", "logs")]
    for folder in (figures, tables, reports, logs):
        folder.mkdir(parents=True, exist_ok=True)
    manuscript = find_latest_manuscript(root)
    return Paths(root, config, figures, tables, reports, logs, manuscript,
                 reports / OUTPUT_DOCX, reports / f"{STEM}_audit.json", logs / f"{STEM}.log")


def read_csv(tables: Path, name: str) -> pd.DataFrame:
    return pd.read_csv(require(tables / name, f"SI evidence '{name}'"))


def set_run_font(run, size: float | None = None, bold: bool | None = None,
                 color: RGBColor | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.06
    for name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 11.5), ("Heading 3", 10)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True


def apply_section(section, landscape: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = A4_H, A4_W
        margin = MARGIN_L
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = A4_W, A4_H
        margin = MARGIN_P
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)


def add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def add_headers_footers(doc: Document) -> None:
    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    hp.text = "Supplementary Information"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(hp.runs[0], size=8, color=RGBColor(90, 90, 90))
    fp = section.footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page ")
    set_run_font(r, size=8, color=RGBColor(90, 90, 90))
    add_field(fp, "PAGE")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9", size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


HEADER_MAP = {
    "variable": "Stored field / variable", "display_name": "English name",
    "data_type": "Type", "modeling_role": "Analytical role",
    "formal_prediction_input": "Formal predictor", "leakage_or_provenance_control": "Leakage / provenance control",
    "n_records": "Records", "n_sources": "Sources", "n_levels": "Levels",
    "n_single_source_levels": "Single-source levels", "n_records_in_single_source_levels": "Records in single-source levels",
    "stage": "Stage", "model": "Model", "parameters_json": "Parameters",
    "Label": "Method", "Seed": "Seed", "N_seeds": "Seeds",
    "N_prediction_events": "Prediction events", "Unique_experimental_rows": "Unique rows",
    "R2": "R²", "R2_mean": "Mean R²", "R2_std": "SD R²",
    "RMSE": "RMSE", "RMSE_mean": "Mean RMSE", "RMSE_std": "SD RMSE",
    "MAE": "MAE", "MAE_mean": "Mean MAE", "MAE_std": "SD MAE",
    "Strategy": "Validation strategy", "Method": "Method", "Bias": "Bias",
    "N_events": "Events", "N_unique": "Unique rows", "N_sources": "Sources", "N_splits": "Splits",
    "Maximum_source_overlap": "Maximum source overlap", "Total_sources": "Total sources",
    "Score_Column": "AD score", "Threshold": "Threshold", "N_Retained": "Retained events",
    "Unique_Row_N": "Unique rows", "Coverage_within_strategy_method": "Coverage",
    "Coverage_within_split": "Split coverage", "N_Literature_Groups": "Literature groups",
    "Dominant_Group": "Dominant source", "Dominant_Group_Fraction": "Dominant-source fraction",
    "Pass_R2_0p5": "R² ≥ 0.5", "Context": "Context", "Variable_Key": "Variable",
    "rank_all_domain": "Full-data rank", "share_pct_all_domain": "Full-data share (%)",
    "rank_paper_010": "Paper 010 rank", "share_pct_paper_010": "Paper 010 share (%)",
    "rank_retained_domain": "Retained rank", "share_pct_retained_domain": "Retained share (%)",
    "feature_label": "Feature", "prediction_min": "Minimum prediction",
    "prediction_max": "Maximum prediction", "prediction_range": "Prediction range",
    "support_class": "Support class", "Grid_cells": "Grid cells",
}

ENGLISH_VARIABLE_MAP = {
    "氟形态": "Fluoride form",
    "游离氟占比": "Free-fluoride fraction",
    "初始氟_mg_L": "Influent fluoride (mg L-1)",
    "干扰离子": "Interferent label",
    "干扰离子浓度_mmol_L": "Interferent concentration (mmol L-1)",
    "水温_C": "Water temperature (deg C)",
    "混凝剂": "Coagulant",
    "剂量_mM": "Coagulant dose (mM)",
    "进水pH": "Influent pH",
    "停留时间_min": "Contact time (min)",
    "出水pH": "Effluent pH",
    "氟去除率_pct": "Fluoride removal (%)",
    "出水氟_mg_L": "Effluent fluoride (mg L-1)",
    "原始行号": "Original row ID",
    "游离氟占比_原始": "Raw free-fluoride fraction",
    "初始氟_mg_L_原始": "Raw influent fluoride (mg L-1)",
    "干扰离子浓度_mmol_L_原始": "Raw interferent concentration (mmol L-1)",
    "水温_C_原始": "Raw water temperature (deg C)",
    "剂量_mM_原始": "Raw coagulant dose (mM)",
    "进水pH_原始": "Raw influent pH",
    "停留时间_min_原始": "Raw contact time (min)",
    "出水pH_原始": "Raw effluent pH",
    "氟去除率_pct_原始": "Raw fluoride removal (%)",
    "出水氟_mg_L_原始": "Raw effluent fluoride (mg L-1)",
    "由浓度反算去除率_pct": "Removal reconstructed from concentrations (%)",
    "去除率差值_pct": "Removal-rate discrepancy (%)",
    "文献编号": "Source ID",
}


def translate_variable_values(df: pd.DataFrame, columns: Sequence[str]) -> pd.DataFrame:
    out = df.copy()
    for column in columns:
        if column in out.columns:
            out[column] = out[column].map(lambda value: ENGLISH_VARIABLE_MAP.get(str(value), value))
    return out


def friendly_header(column: str) -> str:
    return HEADER_MAP.get(column, column.replace("_", " "))


def fmt_value(value, column: str = "") -> str:
    if pd.isna(value):
        return "—"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, int):
        return f"{value:,}"
    if isinstance(value, float):
        if any(k in column.lower() for k in ("r2", "fraction", "coverage", "share")):
            return f"{value:.3f}"
        if abs(value) >= 1000:
            return f"{value:,.1f}"
        if value != 0 and abs(value) < 0.001:
            return f"{value:.2e}"
        return f"{value:.3f}"
    text = str(value)
    return text.replace("_", " ") if column in {"Strategy", "Method", "Context"} else text


def clean_frame(df: pd.DataFrame, columns: Sequence[str]) -> pd.DataFrame:
    missing = [c for c in columns if c not in df.columns]
    if missing:
        raise KeyError(f"Missing expected columns: {missing}")
    return df[list(columns)].copy()


def add_caption(doc: Document, number: str, title: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{number}. ")
    set_run_font(r, size=10, bold=True)
    r = p.add_run(title)
    set_run_font(r, size=10)


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Note. ")
    set_run_font(r, size=8.3, italic=True)
    r = p.add_run(text)
    set_run_font(r, size=8.3, italic=True)


def add_dataframe(doc: Document, df: pd.DataFrame, font_size: float = 7.3) -> None:
    if df.empty:
        add_note(doc, "No records met this display condition.")
        return
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for j, col in enumerate(df.columns):
        cell = table.rows[0].cells[j]
        cell.text = friendly_header(str(col))
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                set_run_font(r, size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for i, (_, row) in enumerate(df.iterrows()):
        cells = table.add_row().cells
        if i % 2:
            for cell in cells:
                set_cell_shading(cell, PALE_BLUE)
        for j, col in enumerate(df.columns):
            cells[j].text = fmt_value(row[col], str(col))
            set_cell_border(cells[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[j].paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    set_run_font(r, size=font_size)


TEXTS = [
    ("Supplementary Text S1", "Data structure and leakage control",
     "The locked modeling matrix contains 1,011 observations and 26 stored fields, while the source-bearing audit matrix contains the same observations and 27 fields. Stored fields are not synonymous with predictors. Only information available before treatment prediction was eligible as formal input; literature identity, post-treatment measurements, target-derived quantities, raw reconstruction fields, and row identifiers were excluded (Table S1; Figures S1–S2). Descriptive continuous and categorical coverage is summarized in Table S2."),
    ("Supplementary Text S2", "Model specification and random-split performance",
     "Estimator settings were recovered from the original analysis record without refitting models (Table S3). Across five prespecified seeds, the weighted four-model ensemble achieved R² = 0.882 ± 0.010, RMSE = 10.249, and MAE = 6.404; the corresponding pooled values were R² = 0.882, RMSE = 10.267, and MAE = 6.404 (Table S4; Figure S3). These results quantify internal learning ability under random partitioning rather than cross-literature transfer."),
    ("Supplementary Text S3", "Source-isolated generalization",
     "Grouped validation and leave-one-literature-out testing were markedly weaker than random splitting, revealing substantial between-literature heterogeneity (Tables S5–S6; Figures S4–S5). In the repeated balanced audit, the strongest pooled result was obtained by the KRR–CatBoost weighted model (R² = 0.194, RMSE = 27.513, MAE = 22.894). All source-isolated splits had zero source overlap."),
    ("Supplementary Text S4", "Selective applicability-domain evidence",
     "At the exact target-independent P08 operating point, retained predictions yielded R² = 0.748, RMSE = 11.475, and MAE = 9.156, but represented only 52 prediction events and 44 unique observations (Tables S7–S9; Figures S6–S9). All retained events originated from Paper_010, only one of two retained splits achieved R² ≥ 0.5, and transfer across validation settings was inconsistent. The retained region is therefore interpreted as a narrow source-specific candidate window rather than a universal cross-literature applicability domain."),
    ("Supplementary Text S5", "Extended model interpretation",
     "SHAP attribution varied across the full dataset, Paper_010, and the retained domain, with coagulant identity, dosage, and interferent identity remaining prominent (Table S10; Figure S10). The two-dimensional partial-dependence interpretation was evaluated against empirical grid support (Figure S11). SHAP and PDP patterns describe the fitted model and are not interpreted as causal effects or prescriptive chemical optima."),
]


TABLE_TITLES = {
    "Table S1": "Variable dictionary, analytical roles, and leakage controls",
    "Table S2": "Data support and categorical-source structure",
    "Table S3": "Recovered estimator configurations",
    "Table S4": "Random-split performance across five seeds",
    "Table S5": "Pooled source-isolated validation performance",
    "Table S6": "Source-isolation and source-macro performance",
    "Table S7": "Exact selective-AD policy and split stability",
    "Table S8": "Transfer of the fixed AD rule across validation settings",
    "Table S9": "Selective-domain chemistry and representativeness",
    "Table S10": "SHAP importance and PDP response range",
}


def build_table_frames(tables: Path):
    r = lambda name: read_csv(tables, name)
    frames = {}
    dictionary = clean_frame(
        r("SI_final_numbering_caption_manifest_v1_Table_S1_variable_dictionary.csv"),
        ["display_name", "data_type", "modeling_role", "formal_prediction_input", "leakage_or_provenance_control"])
    dictionary = dictionary.rename(columns={"display_name": "Variable"})
    frames["Table S1"] = [("", dictionary,
        "All fields are presented in English. Source identity and post-treatment or target-derived quantities were never formal predictors.")]

    cont = r("SI_final_numbering_caption_manifest_v1_Table_S3_continuous_statistics.csv")
    cont_cols = [c for c in ["display_name", "n_records", "n_sources", "mean", "std", "min", "median", "max"] if c in cont.columns]
    cont = cont.rename(columns={"display_name": "Variable"})
    cont_cols = ["Variable" if c == "display_name" else c for c in cont_cols]
    cat = r("SI_categorical_composition_source_coverage_v1_variable_summary.csv")
    cat_cols = [c for c in ["variable", "n_records", "n_sources", "n_levels", "n_single_source_levels", "n_records_in_single_source_levels"] if c in cat.columns]
    frames["Table S2"] = [("S2a. Continuous-variable support", cont[cont_cols].copy(),
        "Statistics are descriptive and do not redefine the locked prediction target."),
        ("S2b. Categorical-variable source structure", cat[cat_cols].copy(),
        "Sparse and single-source categorical levels were retained in the processed evidence.")]

    cfg = clean_frame(r("SI_model_parameters_and_search_spaces_v1_explicit_model_parameters.csv"),
                      ["stage", "model", "parameters_json", "source_file"])
    frames["Table S3"] = [("", cfg, "Configurations were recovered read-only from the original analysis scripts; no model was retrained.")]

    seed = clean_frame(r("SI_random_split_complete_results_v1_seed_summary.csv"),
                       ["Label", "N_seeds", "R2_mean", "R2_std", "RMSE_mean", "RMSE_std", "MAE_mean", "MAE_std"])
    pooled = clean_frame(r("SI_random_split_complete_results_v1_pooled_metrics.csv"),
                         ["Label", "N_prediction_events", "Unique_experimental_rows", "R2", "RMSE", "MAE"])
    frames["Table S4"] = [("S4a. Across-seed summary", seed,
        "The same five prespecified seeds were used for each method."),
        ("S4b. Five-seed pooled performance", pooled,
        "Pooled rows are repeated prediction events across seeds and are not independent experimental samples.")]

    grouped = clean_frame(r("SI_grouped_validation_lolo_complete_results_v1_recomputed_pooled_metrics.csv"),
                          ["Strategy", "Method", "R2", "RMSE", "MAE", "Bias", "N_unique", "N_sources", "N_splits"])
    frames["Table S5"] = [("", grouped,
        "Every validation split isolates literature sources. Strategy-specific results should not be conflated with random-split performance.")]

    iso = r("SI_grouped_validation_lolo_complete_results_v1_source_isolation_audit.csv")
    iso_sum = iso.groupby("Strategy", as_index=False).agg(
        Splits=("Split_ID", "nunique"), Maximum_source_overlap=("Source_overlap", "max"), Total_sources=("Total_sources", "max"))
    macro = r("SI_grouped_validation_lolo_complete_results_v1_original_paper_macro_summary.csv")
    macro = macro[macro["Strategy"].eq("LOLO_LeaveOneLiteratureOut")].copy()
    macro_cols = ["Method", "PaperMacro_RMSE_mean", "PaperMacro_RMSE_median", "PaperMacro_MAE_mean", "PaperMacro_MAE_median", "n_papers"]
    frames["Table S6"] = [("S6a. Source-isolation audit", iso_sum,
        "Maximum source overlap was required to equal zero."),
        ("S6b. LOLO source-macro performance", clean_frame(macro, macro_cols),
        "Macro statistics assign equal weight to each represented literature source.")]

    exact_cols = ["Strategy", "Method", "Zone", "Score_Column", "Threshold", "R2", "RMSE", "MAE",
                  "N_Retained", "Unique_Row_N", "Coverage_within_strategy_method", "N_Literature_Groups",
                  "Dominant_Group", "Dominant_Group_Fraction"]
    exact = clean_frame(r("SI_ad_threshold_coverage_sensitivity_audit_v1_exact_policy_metrics.csv"), exact_cols)
    stable_cols = ["Split_ID", "R2", "RMSE", "MAE", "N_Retained", "Coverage_within_split", "Dominant_Group", "Pass_R2_0p5"]
    stable = clean_frame(r("SI_ad_threshold_coverage_sensitivity_audit_v1_split_stability.csv"), stable_cols)
    frames["Table S7"] = [("S7a. Exact P08 policy", exact,
        "The rule uses target-independent AD scores; the retained evidence is source-specific."),
        ("S7b. Split-level stability", stable,
        "Only one of two retained splits achieved R² ≥ 0.5.")]

    transfer = r("SI_ad_threshold_coverage_sensitivity_audit_v1_strategy_method_transfer.csv")
    keys = ["Strategy", "Method"]
    keep = []
    for _, group in transfer.groupby("Strategy"):
        chosen = pd.concat([group[group["Method"].eq("XGBoost")], group.nlargest(1, "R2")]).drop_duplicates(keys)
        keep.append(chosen)
    transfer = pd.concat(keep, ignore_index=True)
    transfer_cols = ["Strategy", "Method", "R2", "RMSE", "MAE", "N_Retained",
                     "Coverage_within_strategy_method", "N_Literature_Groups", "Dominant_Group", "Dominant_Group_Fraction"]
    frames["Table S8"] = [("", clean_frame(transfer, transfer_cols),
        "For each strategy, XGBoost and the best-R² evaluated method are shown; this compact comparison is descriptive, not a new selection step.")]

    window = translate_variable_values(
        r("SI_ad_threshold_coverage_sensitivity_audit_v1_numeric_operating_window.csv"), ["Variable_Key"])
    window = window[window["Context"].isin(["Retained_events", "Paper_010", "All_data"])].copy()
    window_cols = [c for c in ["Context", "Variable_Key", "N", "Min", "Q25", "Median", "Q75", "Max"] if c in window.columns]
    rep = r("SI_ad_threshold_coverage_sensitivity_audit_v1_representativeness.csv")
    frames["Table S9"] = [("S9a. Numeric operating window", window[window_cols].copy(),
        "The retained range is descriptive and is not a prospective universal specification."),
        ("S9b. Representativeness", rep,
        "Fractions quantify the limited coverage and source concentration of the retained domain.")]

    shap = clean_frame(r("SI_shap_pdp_extended_interpretation_v1_feature_context_metrics.csv"),
                       ["feature_label", "rank_all_domain", "share_pct_all_domain", "rank_paper_010",
                        "share_pct_paper_010", "rank_retained_domain", "share_pct_retained_domain"])
    pdp = translate_variable_values(
        r("SI_shap_pdp_extended_interpretation_v1_pdp_surface_summary.csv"), ["feature_x", "feature_y"])
    pdp_cols = ["feature_x", "feature_y", "x_min", "x_max", "y_min", "y_max",
                "prediction_min", "prediction_max", "prediction_range"]
    frames["Table S10"] = [("S10a. Feature importance across contexts", shap,
        "Context labels define diagnostic subsets only; source identity was not a model feature."),
        ("S10b. PDP surface range", clean_frame(pdp, pdp_cols),
        "Response extrema describe the fitted model over the evaluated grid and are not causal process optima.")]
    return frames


FIGURE_SELECTION = [
    ("Figure S1", "Supplementary data architecture and leakage-control audit"),
    ("Figure S2", "Continuous-variable support and source coverage"),
    ("Figure S7", "Seed-level stability of random-split models"),
    ("Figure S9", "Grouped-validation performance across strategies and models"),
    ("Figure S10", "LOLO error by held-out literature source"),
    ("Figure S12", "Selective-AD threshold evidence"),
    ("Figure S13", "Selective-AD stability, uncertainty, and representativeness"),
    ("Figure S14", "Selective-AD chemistry window"),
    ("Figure S15", "Selective-AD transfer across strategies and methods"),
    ("Figure S16", "SHAP importance across interpretation contexts"),
    ("Figure S17", "Empirical support for two-dimensional PDP interpretation"),
]


def selected_figure_manifest(full: pd.DataFrame) -> pd.DataFrame:
    lookup = full.set_index("number")
    rows = []
    for i, (old_number, expected_title) in enumerate(FIGURE_SELECTION, 1):
        if old_number not in lookup.index:
            raise KeyError(f"Missing figure-manifest entry: {old_number}")
        row = lookup.loc[old_number].copy()
        if str(row["title"]) != expected_title:
            raise AssertionError(f"Unexpected title for {old_number}: {row['title']}")
        row["source_number"] = old_number
        row["number"] = f"Figure S{i}"
        rows.append(row)
    return pd.DataFrame(rows)


def new_page_section(doc: Document, landscape: bool) -> None:
    apply_section(doc.add_section(WD_SECTION.NEW_PAGE), landscape)


def add_one_figure(doc: Document, row: pd.Series, img: Path, width_cm: float) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(img), width=Cm(width_cm))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = cp.add_run(f"{row['number']}. {row['title']}. ")
    set_run_font(r, size=8.8, bold=True)
    r = cp.add_run(str(row["caption"]))
    set_run_font(r, size=8.8)


def build_document(paths: Paths, title: str, table_frames, figures: pd.DataFrame) -> Document:
    doc = Document()
    configure_styles(doc)
    apply_section(doc.sections[0], False)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run("Supplementary Information")
    set_run_font(r, size=20, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run(title)
    set_run_font(r, size=15, bold=True)
    doc.add_paragraph("\n\n")  # Author information intentionally left blank.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Journal of Environmental Chemical Engineering")
    set_run_font(r, size=11, italic=True)
    doc.add_page_break()

    doc.add_heading("Contents", level=1)
    for number, heading, _ in TEXTS:
        p = doc.add_paragraph(f"{number}. {heading}")
        p.paragraph_format.left_indent = Cm(0.35)
    p = doc.add_paragraph("Supplementary Tables S1–S10")
    p.paragraph_format.left_indent = Cm(0.35)
    p = doc.add_paragraph("Supplementary Figures S1–S11")
    p.paragraph_format.left_indent = Cm(0.35)

    doc.add_heading("Supplementary Text", level=1)
    for number, heading, text in TEXTS:
        doc.add_heading(f"{number}. {heading}", level=2)
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for i in range(1, 11):
        number = f"Table S{i}"
        parts = table_frames[number]
        max_cols = max(len(frame.columns) for _, frame, _ in parts)
        new_page_section(doc, max_cols >= 8)
        if i == 1:
            doc.add_heading("Supplementary Tables", level=1)
        add_caption(doc, number, TABLE_TITLES[number])
        for j, (subtitle, frame, note) in enumerate(parts):
            if subtitle:
                doc.add_heading(subtitle, level=3)
            font = 6.1 if len(frame.columns) >= 10 else 6.7 if len(frame.columns) >= 8 else 7.3
            add_dataframe(doc, frame, font)
            add_note(doc, note)
            if j < len(parts) - 1:
                doc.add_paragraph()

    new_page_section(doc, False)
    doc.add_heading("Supplementary Figures", level=1)
    pending = []

    def flush_pending(first: bool = False) -> None:
        nonlocal pending
        if not pending:
            return
        if not first:
            doc.add_page_break()
        if len(pending) == 1:
            row, img = pending[0]
            add_one_figure(doc, row, img, 15.4)
        else:
            table = doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, (row, img) in zip(table.rows[0].cells, pending):
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(img), width=Cm(7.7))
                cp = cell.add_paragraph()
                r = cp.add_run(f"{row['number']}. {row['title']}. ")
                set_run_font(r, size=7.8, bold=True)
                r = cp.add_run(str(row["caption"]))
                set_run_font(r, size=7.8)
        pending = []

    figure_page_started = False
    for _, row in figures.iterrows():
        img = require(paths.figures / f"{row['file_stem']}.png", f"{row['source_number']} PNG")
        with Image.open(img) as im:
            aspect = im.width / im.height
            pixels = im.width * im.height
        small = pixels <= 14_000_000 and 0.78 <= aspect <= 1.55
        if small:
            pending.append((row, img))
            if len(pending) == 2:
                flush_pending(first=not figure_page_started)
                figure_page_started = True
        else:
            if pending:
                flush_pending(first=not figure_page_started)
                figure_page_started = True
            if figure_page_started:
                doc.add_page_break()
            add_one_figure(doc, row, img, 15.4)
            figure_page_started = True
    if pending:
        flush_pending(first=not figure_page_started)

    add_headers_footers(doc)
    core = doc.core_properties
    core.author = ""
    core.last_modified_by = ""
    core.title = f"Supplementary Information — {title}"
    core.subject = "Concise Supplementary Information for the Journal of Environmental Chemical Engineering"
    core.keywords = "coagulation; defluoridation; machine learning; applicability domain; grouped validation"
    core.comments = "Generated from locked evidence; author information intentionally left blank."
    return doc


def main() -> None:
    start = time.perf_counter()
    paths = project_paths()
    sys.path.insert(0, str(paths.config.parent))
    __import__("config")

    full_figures = read_csv(paths.tables, "SI_final_numbering_caption_manifest_v1_figure_manifest.csv")
    figures = selected_figure_manifest(full_figures)
    for _, row in figures.iterrows():
        require(paths.figures / f"{row['file_stem']}.png", f"selected figure {row['source_number']}")

    protected = sorted({paths.config, paths.manuscript, *paths.figures.glob("*.png"),
                        *paths.figures.glob("*.pdf"), *paths.tables.glob("*.csv")})
    before = {str(p.relative_to(paths.root)): sha256(p) for p in protected if p.is_file()}

    title = extract_english_title(paths.manuscript)
    table_frames = build_table_frames(paths.tables)
    doc = build_document(paths, title, table_frames, figures)
    doc.save(paths.output_docx)
    require(paths.output_docx, "generated concise SI Word file")

    after = {str(p.relative_to(paths.root)): sha256(p) for p in protected if p.is_file()}
    changed = [name for name in before if before[name] != after.get(name)]
    if changed:
        raise AssertionError("Protected evidence changed during Word assembly: " + ", ".join(changed))

    runtime = time.perf_counter() - start
    audit = {
        "script": STEM, "article_title": title,
        "manuscript_source": str(paths.manuscript.relative_to(paths.root)) if paths.manuscript.is_relative_to(paths.root) else str(paths.manuscript),
        "output_docx": str(paths.output_docx.relative_to(paths.root)),
        "supplementary_text_count": len(TEXTS), "table_count": 10, "figure_count": len(figures),
        "source_figure_numbers": list(figures["source_number"]),
        "author_information": "intentionally blank",
        "protected_file_count": len(before), "protected_files_changed": changed,
        "output_sha256": sha256(paths.output_docx), "runtime_seconds": round(runtime, 3),
    }
    paths.audit.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    lines = [
        "PASS: concise submission-ready SI Word assembled from locked evidence.",
        f"Title: {title}",
        f"Supplementary Text: {len(TEXTS)}; Tables: 10; Figures: {len(figures)}.",
        f"Output: {paths.output_docx}",
        f"Protected evidence files unchanged: {not changed}.",
        f"Total runtime: {runtime:.2f} s",
    ]
    paths.log.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    try:
        main()
    except Exception:
        print(f"Total runtime: {time.perf_counter() - RUN_START:.2f} s")
        raise
