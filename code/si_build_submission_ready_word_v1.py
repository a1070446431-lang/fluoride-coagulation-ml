# -*- coding: utf-8 -*-
"""Build the submission-ready English Supplementary Information Word file.

Place this complete script in ``fluoride_coagulation_ml_project/src/`` and run
it from PyCharm. It reads the locked SI tables/figures already stored under
``outputs/jece_final_package`` and the latest main-manuscript Word file found
in the project root. It does not train models or modify existing evidence.

Outputs
-------
outputs/jece_final_package/reports/Supplementary_Information_Paper1_JECE_v1.docx
outputs/jece_final_package/tables/SI_submission_machine_readable_files_v1.csv
outputs/jece_final_package/reports/SI_build_submission_ready_word_v1_audit.json
outputs/jece_final_package/logs/SI_build_submission_ready_word_v1.log

Dependencies: Python >= 3.9, pandas, python-docx, Pillow.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable, Sequence

import pandas as pd
from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


STEM = "SI_build_submission_ready_word_v1"
OUTPUT_DOCX = "Supplementary_Information_Paper1_JECE_v1.docx"
MANUSCRIPT_OVERRIDE: str | None = None
RUN_START = time.perf_counter()

BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
LIGHT_GREY = "D9E2F3"
MID_GREY = "B7C9DC"
WHITE = "FFFFFF"
BLACK = RGBColor(0, 0, 0)

A4_W, A4_H = Cm(21.0), Cm(29.7)
MARGIN_P = Cm(1.8)
MARGIN_L = Cm(1.4)


@dataclass(frozen=True)
class Paths:
    root: Path
    config: Path
    figures: Path
    tables: Path
    reports: Path
    logs: Path
    manuscript: Path
    output_docx: Path
    attachment_manifest: Path
    audit: Path
    log: Path


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def require(path: Path, purpose: str) -> Path:
    if not path.is_file():
        raise FileNotFoundError(f"Missing {purpose}: {path}")
    return path


def find_latest_manuscript(root: Path) -> Path:
    if MANUSCRIPT_OVERRIDE:
        candidate = Path(MANUSCRIPT_OVERRIDE)
        if not candidate.is_absolute():
            candidate = root / candidate
        return require(candidate, "main manuscript Word file")
    candidates = []
    for folder in (root, root.parent):
        for p in folder.glob("*.docx"):
            low = p.name.lower()
            if "supplement" in low or "si_" in low or p.name == OUTPUT_DOCX:
                continue
            score = 0
            if "图4文字更新" in p.name:
                score += 100
            if "总版3" in p.name or "总版(3)" in p.name:
                score += 50
            if "总版" in p.name:
                score += 20
            candidates.append((score, p.stat().st_mtime, p))
    if not candidates:
        raise FileNotFoundError(
            "No main-manuscript .docx was found. Copy the latest manuscript Word "
            "file to the project root or set MANUSCRIPT_OVERRIDE near the top of this script."
        )
    return max(candidates, key=lambda x: (x[0], x[1]))[2]


def project_paths() -> Paths:
    root = Path(__file__).resolve().parents[1]
    config = require(root / "src" / "config.py", "project configuration")
    base = root / "outputs" / "jece_final_package"
    figures, tables, reports, logs = [base / x for x in ("figures", "tables", "reports", "logs")]
    for d in (figures, tables, reports, logs):
        d.mkdir(parents=True, exist_ok=True)
    manuscript = find_latest_manuscript(root)
    return Paths(
        root, config, figures, tables, reports, logs, manuscript,
        reports / OUTPUT_DOCX,
        tables / "SI_submission_machine_readable_files_v1.csv",
        reports / f"{STEM}_audit.json",
        logs / f"{STEM}.log",
    )


def read_csv(tables: Path, name: str) -> pd.DataFrame:
    p = require(tables / name, f"SI table evidence '{name}'")
    return pd.read_csv(p)


def extract_english_title(manuscript: Path) -> str:
    doc = Document(manuscript)
    candidates: list[tuple[float, str]] = []
    keywords = ("fluor", "coagul", "predict", "applicability", "data-driven", "model")
    for i, para in enumerate(doc.paragraphs[:120]):
        text = " ".join(para.text.split()).strip(" -–—")
        words = re.findall(r"[A-Za-z][A-Za-z\-']*", text)
        if not (7 <= len(words) <= 35) or len(text) > 300:
            continue
        ascii_ratio = sum(ord(c) < 128 for c in text) / max(len(text), 1)
        if ascii_ratio < 0.92:
            continue
        key_hits = sum(k in text.lower() for k in keywords)
        score = 5 * key_hits - 0.03 * i
        if para.style and "title" in para.style.name.lower():
            score += 8
        candidates.append((score, text))
    if not candidates or max(candidates)[0] < 10:
        raise ValueError(
            f"Could not identify the English article title in {manuscript.name}. "
            "Ensure the latest manuscript contains an English title near its beginning."
        )
    return max(candidates, key=lambda x: x[0])[1]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "B8C4CE", size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def set_repeatable_heading(paragraph) -> None:
    set_keep_with_next(paragraph)
    paragraph.paragraph_format.keep_together = True


def set_run_font(run, name: str = "Times New Roman", size: float | None = None,
                 bold: bool | None = None, color: RGBColor | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = name
    # English follows JECE-style Times New Roman; the original Chinese labels
    # retained at first occurrence use a standard Windows CJK font in Word.
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, bold in (("Title", 18, True), ("Heading 1", 14, True),
                                   ("Heading 2", 12, True), ("Heading 3", 10.5, True)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def apply_section(section, landscape: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = A4_H, A4_W
        margin = MARGIN_L
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = A4_W, A4_H
        margin = MARGIN_P
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)


def add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def add_headers_footers(doc: Document) -> None:
    # Later sections remain linked to the first. Editing every linked footer
    # would append the same PAGE field repeatedly to one shared footer part.
    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    hp.text = "Supplementary Information"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(hp.runs[0], size=8, color=RGBColor(90, 90, 90))
    fp = section.footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page ")
    set_run_font(r, size=8, color=RGBColor(90, 90, 90))
    add_field(fp, "PAGE")


def fmt_value(v, column: str = "") -> str:
    if pd.isna(v):
        return "—"
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, (int,)):
        return f"{v:,}"
    if isinstance(v, float):
        if not pd.notna(v):
            return "—"
        ac = abs(v)
        if ac >= 1000:
            return f"{v:,.1f}"
        if any(x in column.lower() for x in ("r2", "fraction", "correlation", "coverage", "share")):
            return f"{v:.3f}"
        if ac != 0 and ac < 0.001:
            return f"{v:.2e}"
        return f"{v:.3f}"
    s = str(v)
    return s.replace("_", " ") if column in {"Strategy", "Method", "Context"} else s


HEADER_MAP = {
    "variable": "Stored field / variable", "display_name": "English display name",
    "data_type": "Type", "modeling_role": "Analytical role",
    "formal_prediction_input": "Formal predictor", "leakage_or_provenance_control": "Control",
    "n_records": "Records", "n_sources": "Sources", "n_levels": "Levels",
    "n_single_source_levels": "Single-source levels", "n_records_in_single_source_levels": "Records in single-source levels",
    "n_empty_cells": "Empty cells", "n_literal_none": "Literal ‘None’",
    "category_display": "English category", "category_original": "Original category",
    "pct_all_records": "% records", "pct_all_sources": "% sources",
    "dominant_source_pct": "Dominant-source %", "source_id": "Literature source",
    "audit_item": "Audit item", "expected_or_policy": "Expected value / policy",
    "pass": "Passed", "parameters_json": "Parameters", "declared_parameters_json": "Declared parameters",
    "explicit_effective_parameters_json": "Effective parameters", "range_or_expression": "Range / expression",
    "listed_values": "Listed values", "n_exact_configs": "Exact configurations",
    "N_prediction_events": "Prediction events", "Unique_experimental_rows": "Unique experimental rows",
    "R2": "R²", "R2_mean": "Mean R²", "R2_std": "SD R²", "Pooled_R2": "Pooled R²",
    "RMSE": "RMSE", "RMSE_mean": "Mean RMSE", "RMSE_std": "SD RMSE", "Pooled_RMSE": "Pooled RMSE",
    "MAE": "MAE", "MAE_mean": "Mean MAE", "MAE_std": "SD MAE", "Pooled_MAE": "Pooled MAE",
    "Mean_bias": "Mean bias", "Bias": "Bias", "N_events": "Events", "N_unique": "Unique rows",
    "N_sources": "Sources", "N_splits": "Splits", "文献编号": "Literature source",
    "Paper_R2": "Source R²", "Paper_RMSE": "Source RMSE", "Paper_MAE": "Source MAE",
    "Score_Column": "AD score", "Threshold": "Threshold", "N_Retained": "Retained events",
    "Unique_Row_N": "Unique rows", "N_Literature_Groups": "Literature groups",
    "Dominant_Group": "Dominant source", "Dominant_Group_Fraction": "Dominant-source fraction",
    "feature_label": "Feature", "spearman_rank_correlation": "Spearman ρ",
    "contribution_pct": "Contribution (%)", "sum_mean_abs_shap": "Σ mean |SHAP|",
    "profile_axis": "Profile axis", "grid_value": "Grid value", "pdp_mean": "Mean PDP",
    "pdp_min": "Minimum PDP", "pdp_max": "Maximum PDP",
}


def friendly_header(column: str) -> str:
    return HEADER_MAP.get(column, column.replace("_", " "))


def add_caption(doc: Document, number: str, title: str) -> None:
    p = doc.add_paragraph()
    set_repeatable_heading(p)
    r = p.add_run(f"{number}. ")
    set_run_font(r, size=10, bold=True)
    r = p.add_run(title)
    set_run_font(r, size=10, bold=False)
    p.paragraph_format.space_after = Pt(4)


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("Note. ")
    set_run_font(r, size=8.5, italic=True)
    r = p.add_run(text)
    set_run_font(r, size=8.5, italic=True)


def clean_frame(df: pd.DataFrame, columns: Sequence[str] | None = None,
                rename: dict[str, str] | None = None) -> pd.DataFrame:
    out = df.copy()
    if columns is not None:
        missing = [c for c in columns if c not in out.columns]
        if missing:
            raise KeyError(f"Missing expected columns: {missing}")
        out = out[list(columns)]
    if rename:
        out = out.rename(columns=rename)
    return out


def add_dataframe(doc: Document, df: pd.DataFrame, font_size: float = 7.5,
                  max_rows: int | None = None) -> None:
    if max_rows is not None and len(df) > max_rows:
        raise ValueError(f"Refusing silent truncation: {len(df)} rows exceed explicit limit {max_rows}")
    if df.empty:
        add_note(doc, "No records met this display condition.")
        return
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for j, col in enumerate(df.columns):
        cell = table.rows[0].cells[j]
        cell.text = friendly_header(str(col))
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                set_run_font(r, size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for i, (_, row) in enumerate(df.iterrows()):
        cells = table.add_row().cells
        if i % 2 == 1:
            for c in cells:
                set_cell_shading(c, PALE_BLUE)
        for j, col in enumerate(df.columns):
            cells[j].text = fmt_value(row[col], str(col))
            set_cell_border(cells[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[j].paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    set_run_font(r, size=font_size)


def new_table_section(doc: Document, landscape: bool) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_section(section, landscape)


def weighted_summary(df: pd.DataFrame, keys: list[str], value: str) -> pd.DataFrame:
    return (df.groupby(keys, dropna=False)[value]
            .agg(["count", "mean", "min", "max"]).reset_index()
            .rename(columns={"count": "N", "mean": f"{value}_mean", "min": f"{value}_min", "max": f"{value}_max"}))


def build_table_frames(t: Path) -> dict[str, list[tuple[str, pd.DataFrame, str]]]:
    R = lambda n: read_csv(t, n)
    frames: dict[str, list[tuple[str, pd.DataFrame, str]]] = {}
    frames["Table S1"] = [("", R("SI_final_numbering_caption_manifest_v1_Table_S1_variable_dictionary.csv"),
        "The stored-field column preserves the original processed-matrix labels. Source identity and post-treatment or target-derived fields are never formal prediction inputs.")]
    frames["Table S2"] = [("", R("SI_final_numbering_caption_manifest_v1_Table_S2_quality_control_audit.csv"),
        "All checks are evaluated against the locked 1011-row matrices; a pass value of Yes indicates exact agreement with the stated policy or expected value.")]
    frames["Table S3"] = [("", R("SI_final_numbering_caption_manifest_v1_Table_S3_continuous_statistics.csv"),
        "Statistics are descriptive. Source coverage is counted from the source-bearing matrix and does not make source identity a predictor.")]

    varsum = R("SI_categorical_composition_source_coverage_v1_variable_summary.csv")
    cats = clean_frame(R("SI_categorical_composition_source_coverage_v1_category_summary.csv"),
        ["variable", "category_original", "category_display", "n_records", "pct_all_records", "n_sources", "pct_all_sources", "single_source", "dominant_source_pct"])
    frames["Table S4"] = [("S4a. Variable-level composition", varsum,
        "Original Chinese labels are retained at first occurrence; the English display labels are used throughout the remaining SI."),
        ("S4b. Level-level composition", cats,
        "Sparse and single-source levels were retained. Percentages are descriptive and may not sum exactly to 100 because of rounding.")]
    sizes = R("SI_categorical_composition_source_coverage_v1_source_sizes.csv")
    cov = (cats.groupby("variable", as_index=False)
           .agg(levels=("category_display", "size"), single_source_levels=("single_source", "sum"),
                median_sources_per_level=("n_sources", "median"), min_sources_per_level=("n_sources", "min"),
                max_sources_per_level=("n_sources", "max")))
    frames["Table S5"] = [("S5a. Literature-source sizes", sizes,
        "Source identifiers are anonymized provenance labels used only for validation, coverage diagnosis, and applicability-domain auditing."),
        ("S5b. Category-level source coverage summary", cov,
        "Complete category-by-source counts remain in the locked evidence files and are visualized in Figures S3–S6.")]

    explicit = clean_frame(R("SI_model_parameters_and_search_spaces_v1_explicit_model_parameters.csv"),
        ["stage", "model", "parameters_json", "source_file", "line"])
    configs = R("SI_model_parameters_and_search_spaces_v1_grouped_exact_configurations.csv")
    cfgsum = (configs.groupby(["stage", "model", "block"], as_index=False)
              .agg(n_exact_configurations=("config_id", "nunique"), first_source=("source_file", "first")))
    non_tree = clean_frame(R("SI_model_parameters_and_search_spaces_v1_non_tree_exact_configurations.csv"),
        ["Stage", "Model_Family", "Config_Name", "Params", "Allowed_K", "Weighting_Mode"])
    frames["Table S6"] = [("S6a. Explicit estimator parameters", explicit,
        "Parameters were recovered from the original analysis scripts; no model was retrained for this SI."),
        ("S6b. Grouped-validation configuration blocks", cfgsum,
        "Counts summarize the exact declared configuration records."),
        ("S6c. Exact non-tree configurations", non_tree,
        "Configuration names and parameter strings reproduce the analysis record.")]
    hist = R("SI_model_parameters_and_search_spaces_v1_historical_search_spaces.csv")
    ranges = R("SI_model_parameters_and_search_spaces_v1_grouped_listed_value_ranges.csv")
    frames["Table S7"] = [("S7a. Historical random-split search spaces", hist,
        "These are the searched ranges, not necessarily the finally selected settings."),
        ("S7b. Listed grouped-validation values", ranges,
        "Listed values summarize the discrete grouped-validation configuration record.")]
    rw = R("SI_model_parameters_and_search_spaces_v1_random_selected_weights.csv")
    rw_sum = weighted_summary(rw, ["Method", "Model"], "Weight")
    gs = R("SI_model_parameters_and_search_spaces_v1_grouped_selected_settings.csv")
    gs["Selected_K_numeric"] = pd.to_numeric(gs["Selected_K"], errors="coerce")
    gs["Selected_K_includes_all"] = gs["Selected_K"].astype(str).str.lower().eq("all")
    gs_sum = (gs.groupby(["Stage", "Method"], as_index=False)
              .agg(N=("Seed", "size"), selected_k_numeric_min=("Selected_K_numeric", "min"),
                   selected_k_numeric_max=("Selected_K_numeric", "max"),
                   selected_k_includes_all=("Selected_K_includes_all", "max"),
                   validation_RMSE_mean=("Validation_RMSE", "mean"), test_RMSE_mean=("Test_RMSE", "mean")))
    cross = R("SI_model_parameters_and_search_spaces_v1_selection_crosscheck.csv")
    cross_sum = cross.groupby(["stage", "status"], as_index=False).size().rename(columns={"size": "N"})
    frames["Table S8"] = [("S8a. Random-split ensemble-weight summary", rw_sum,
        "Weights are summarized across the five prespecified random seeds."),
        ("S8b. Grouped-validation selected-setting summary", gs_sum,
        "Values summarize the locked selected configurations and test records."),
        ("S8c. Configuration cross-check", cross_sum,
        "Every selected configuration was checked against the recovered configuration record.")]

    seed = clean_frame(R("SI_random_split_complete_results_v1_seed_metrics.csv"),
        ["Label", "Seed", "N_prediction_events", "Selected_K", "Validation_RMSE", "R2", "RMSE", "MAE"])
    summ = clean_frame(R("SI_random_split_complete_results_v1_seed_summary.csv"),
        ["Label", "N_seeds", "R2_mean", "R2_std", "RMSE_mean", "RMSE_std", "MAE_mean", "MAE_std"])
    frames["Table S9"] = [("S9a. Seed-level metrics", seed,
        "Each method was evaluated using the same five prespecified seeds (7, 21, 42, 66, and 88)."),
        ("S9b. Across-seed summary", summ,
        "Standard deviations describe between-seed variability.")]
    pooled = clean_frame(R("SI_random_split_complete_results_v1_pooled_metrics.csv"),
        ["Label", "N_prediction_events", "N_seeds", "Unique_experimental_rows", "R2", "RMSE", "MAE", "Correlation", "Std_ratio"])
    resid = clean_frame(R("SI_random_split_complete_results_v1_residual_statistics.csv"),
        ["Method", "Mean_bias", "Residual_Q05", "Residual_Q50", "Residual_Q95", "Absolute_error_Q50", "Absolute_error_Q90", "N_abs_error_gt_50"])
    frames["Table S10"] = [("S10a. Pooled performance", pooled,
        "Pooled rows are repeated prediction events across seeds, not independent experimental samples."),
        ("S10b. Residual statistics", resid,
        "Residual is defined as observed minus predicted fluoride removal.")]
    cand = R("SI_random_split_complete_results_v1_complete_validation_candidates.csv")
    cand_sum = (cand.groupby(["Seed", "Kind"], as_index=False)
                .agg(N_candidates=("Validation_RMSE", "size"), minimum_validation_RMSE=("Validation_RMSE", "min"),
                     maximum_validation_R2=("Validation_R2", "max")))
    sel = R("SI_random_split_complete_results_v1_validation_selection_audit.csv")
    checks = R("SI_random_split_complete_results_v1_metric_consistency_checks.csv")
    check_sum = checks.groupby("Passed", as_index=False).size().rename(columns={"size": "N"})
    frames["Table S11"] = [("S11a. Candidate-screening summary", cand_sum,
        "Candidate counts and extrema are reconstructed from the complete locked screening table."),
        ("S11b. Selection audit", sel,
        "Selection was based on validation RMSE; test performance was not used to select candidates."),
        ("S11c. Numerical consistency checks", check_sum,
        "The complete seed results and recorded prediction events are designated as separate machine-readable files.")]

    gp = clean_frame(R("SI_grouped_validation_lolo_complete_results_v1_recomputed_pooled_metrics.csv"),
        ["Strategy", "Method", "R2", "RMSE", "MAE", "Bias", "N_events", "N_unique", "N_sources", "N_splits"])
    frames["Table S12"] = [("", gp,
        "All splits isolate literature sources. Strategy-specific values must not be conflated with the historical unified audit baseline reported in the main text.")]
    pm = R("SI_grouped_validation_lolo_complete_results_v1_recomputed_paper_metrics.csv")
    lolo = pm[pm["Strategy"] == "LOLO_LeaveOneLiteratureOut"].copy()
    rows = []
    for source, g in lolo.groupby("Source"):
        best = g.loc[g["RMSE"].idxmin()]
        xgb = g[g["Method"] == "XGBoost"]
        x = xgb.iloc[0] if len(xgb) else best
        rows.append({"Literature_source": source, "N_unique": int(best["N_unique"]),
                     "Best_RMSE_method": best["Method"], "Best_RMSE": best["RMSE"], "Best_MAE": best["MAE"],
                     "XGBoost_R2": x["R2"], "XGBoost_RMSE": x["RMSE"], "XGBoost_MAE": x["MAE"]})
    source_summary = pd.DataFrame(rows).sort_values("Literature_source")
    macro = R("SI_grouped_validation_lolo_complete_results_v1_original_paper_macro_summary.csv")
    frames["Table S13"] = [("S13a. LOLO source-level summary", source_summary,
        "The best-RMSE column is descriptive across evaluated methods; XGBoost columns provide a fixed-method comparator for every held-out source."),
        ("S13b. Source-macro performance", macro,
        "Macro statistics weight each represented literature source equally.")]
    pred = R("SI_grouped_validation_lolo_complete_results_v1_original_predictions.csv")
    pred_sum = (pred.groupby(["Strategy", "Method"], as_index=False)
                .agg(Prediction_events=("Observed", "size"), Sources=("文献编号", "nunique"),
                     Mean_residual=("Residual", "mean"), Median_absolute_residual=("Residual", lambda x: x.abs().median())))
    rc = R("SI_grouped_validation_lolo_complete_results_v1_original_residual_correlation_by_strategy.csv")
    frames["Table S14"] = [("S14a. Prediction-record summary", pred_sum,
        "Complete grouped-validation and LOLO predictions are designated as a separate machine-readable file."),
        ("S14b. Residual correlation by strategy", rc,
        "Residual correlation is diagnostic and does not itself establish transferable predictive skill.")]
    iso = R("SI_grouped_validation_lolo_complete_results_v1_source_isolation_audit.csv")
    iso_sum = (iso.groupby("Strategy", as_index=False)
               .agg(Splits=("Split_ID", "nunique"), Maximum_source_overlap=("Source_overlap", "max"),
                    Total_sources=("Total_sources", "max")))
    fw = R("SI_grouped_validation_lolo_complete_results_v1_original_fusion_weights.csv")
    fw_sum = weighted_summary(fw, ["Strategy", "Method", "Base_Model"], "Weight")
    nc = R("SI_grouped_validation_lolo_complete_results_v1_numeric_checks.csv")
    nc_sum = nc.groupby("Passed", as_index=False).size().rename(columns={"size": "N"})
    frames["Table S15"] = [("S15a. Source-isolation audit", iso_sum,
        "Maximum source overlap must be zero for source-isolated validation."),
        ("S15b. Fusion-weight summary", fw_sum,
        "Weights are summarized over the locked strategy/split records."),
        ("S15c. Numerical checks", nc_sum,
        "No models were retrained during SI reconstruction.")]

    exact = R("SI_ad_threshold_coverage_sensitivity_audit_v1_exact_policy_metrics.csv")
    exact_cols = ["Strategy", "Method", "Zone", "Score_Column", "Threshold", "R2", "RMSE", "MAE",
                  "N_Retained", "Unique_Row_N", "Coverage_within_strategy_method", "N_Literature_Groups",
                  "Dominant_Group", "Dominant_Group_Fraction"]
    comp = R("SI_ad_threshold_coverage_sensitivity_audit_v1_ad_component_summary.csv")
    frames["Table S16"] = [("S16a. Exact P08 policy", clean_frame(exact, exact_cols),
        "The P08 rule uses only target-independent AD scores. Its retained predictions are source-specific evidence, not a universal cross-literature domain."),
        ("S16b. AD-component distribution summary", comp,
        "Complete retained prediction events are designated as a separate machine-readable file.")]
    th = clean_frame(R("SI_ad_threshold_coverage_sensitivity_audit_v1_threshold_sensitivity.csv"),
        ["Threshold", "R2", "RMSE", "MAE", "N_Retained", "Unique_Row_N", "Coverage_within_strategy_method",
         "N_Literature_Groups", "Dominant_Group", "Dominant_Group_Fraction", "Pass_R2_0p5", "Is_Exact_Base_Threshold"])
    stable = clean_frame(R("SI_ad_threshold_coverage_sensitivity_audit_v1_split_stability.csv"),
        ["Split_ID", "Threshold", "R2", "RMSE", "MAE", "N_Retained", "Coverage_within_split", "Dominant_Group", "Pass_R2_0p5"])
    transfer = clean_frame(R("SI_ad_threshold_coverage_sensitivity_audit_v1_strategy_method_transfer.csv"),
        ["Strategy", "Method", "R2", "RMSE", "MAE", "N_Retained", "Coverage_within_strategy_method", "N_Literature_Groups", "Dominant_Group"])
    boot = R("SI_ad_threshold_coverage_sensitivity_audit_v1_event_level_bootstrap_summary.csv")
    frames["Table S17"] = [("S17a. Threshold sensitivity", th,
        "Thresholds were audited around the fixed P08 operating point; test errors were not used to choose the threshold."),
        ("S17b. Split stability", stable,
        "Only one of the two retained splits achieved R² ≥ 0.5, underscoring limited stability."),
        ("S17c. Strategy/method transfer", transfer,
        "Transfer results distinguish a narrow candidate window from generalizable reliability."),
        ("S17d. Event-level bootstrap summary", boot,
        "Bootstrap intervals quantify uncertainty conditional on the retained events and do not correct source concentration.")]
    win = clean_frame(R("SI_ad_threshold_coverage_sensitivity_audit_v1_numeric_operating_window.csv"),
        ["Context", "Variable_Key", "N", "Min", "Q25", "Median", "Q75", "Max"])
    cat = R("SI_ad_threshold_coverage_sensitivity_audit_v1_categorical_composition.csv")
    topcat = (cat.sort_values(["Context", "Variable_Key", "Count"], ascending=[True, True, False])
              .groupby(["Context", "Variable_Key"], as_index=False, group_keys=False).head(5))
    topcat = clean_frame(topcat, ["Context", "Variable_Key", "Category", "Count", "Fraction"])
    rep = R("SI_ad_threshold_coverage_sensitivity_audit_v1_representativeness.csv")
    frames["Table S18"] = [("S18a. Numeric operating window", win,
        "The retained window is compared with Paper_010 and the full data support; it is not a prospective universal specification."),
        ("S18b. Five most frequent categories per context", topcat,
        "Only the five most frequent categories are displayed for compactness; all original categories remain in the evidence file."),
        ("S18c. Representativeness", rep,
        "Fractions quantify the narrow coverage of the selectively retained domain.")]

    fcm = R("SI_shap_pdp_extended_interpretation_v1_feature_context_metrics.csv")
    fgc = R("SI_shap_pdp_extended_interpretation_v1_feature_group_contributions.csv")
    crc = R("SI_shap_pdp_extended_interpretation_v1_context_rank_correlations.csv")
    frames["Table S19"] = [("S19a. Feature importance across contexts", fcm,
        "Context labels are diagnostic filters. Literature-source identity is not a model feature."),
        ("S19b. Feature-group contributions", fgc,
        "Contributions are based on mean absolute SHAP magnitude."),
        ("S19c. Rank correlations", crc,
        "SHAP importance is model-descriptive and does not establish causality.")]
    surf = R("SI_shap_pdp_extended_interpretation_v1_pdp_surface_summary.csv")
    prof = R("SI_shap_pdp_extended_interpretation_v1_pdp_marginal_profiles.csv")
    prof_sum = (prof.groupby(["feature_x", "feature_y", "profile_axis"], as_index=False)
                .agg(Grid_points=("grid_value", "size"), Grid_min=("grid_value", "min"), Grid_max=("grid_value", "max"),
                     PDP_min=("pdp_mean", "min"), PDP_max=("pdp_mean", "max")))
    support = R("SI_shap_pdp_extended_interpretation_v1_pdp_paired_support_grid.csv")
    support_sum = support.groupby(["feature_x", "feature_y", "support_class"], as_index=False).size().rename(columns={"size": "Grid_cells"})
    frames["Table S20"] = [("S20a. PDP surface summary", surf,
        "Prediction extrema describe the fitted model over the evaluated grid, not causal process optima."),
        ("S20b. Marginal-profile summary", prof_sum,
        "Complete marginal profiles are designated as a separate machine-readable file."),
        ("S20c. Empirical support of paired grids", support_sum,
        "Sparse-support grid regions should not be assigned mechanistic meaning.")]
    return frames


TABLE_TITLES = {
    "Table S1": "Variable dictionary, modeling roles, and leakage controls",
    "Table S2": "Processed-matrix quality-control audit",
    "Table S3": "Descriptive statistics and source coverage for continuous variables",
    "Table S4": "Categorical-variable composition",
    "Table S5": "Categorical levels across literature sources",
    "Table S6": "Exact model configurations recovered from the analysis record",
    "Table S7": "Historical hyperparameter search spaces and listed ranges",
    "Table S8": "Selected settings, ensemble weights, and configuration evidence",
    "Table S9": "Complete seed-level random-split metrics",
    "Table S10": "Pooled random-split performance and residual statistics",
    "Table S11": "Random-split candidates, prediction events, and selection audit",
    "Table S12": "Grouped-validation split and pooled performance",
    "Table S13": "Source-level grouped-validation and LOLO metrics",
    "Table S14": "LOLO predictions and residual diagnostics",
    "Table S15": "Grouped-validation configurations, weights, and source-isolation checks",
    "Table S16": "Exact selective-AD policy and retained predictions",
    "Table S17": "AD threshold sensitivity, stability, transfer, and bootstrap uncertainty",
    "Table S18": "Selective-domain chemistry and representativeness",
    "Table S19": "SHAP importance across interpretation contexts",
    "Table S20": "PDP response profiles and empirical support",
}


TEXTS = [
    ("Supplementary Text S1", "Data structure, variable roles, and quality control",
     "The locked modeling matrix contains 1,011 observations and 26 stored fields, whereas the source-bearing audit matrix contains the same observations and 27 fields. Stored fields are not synonymous with predictors: only information available before treatment prediction was eligible as formal model input. Literature identity, post-treatment fluoride and pH, target-derived removal quantities, reconstructed values, raw fields, and row identifiers were excluded from predictive inputs (Table S1; Figure S1).",
     "Quality-control checks verified matrix dimensions, target support, source coverage, formal input roles, and leakage controls (Table S2). Descriptive support and literature-source availability for the continuous variables are reported in Table S3 and Figure S2. These summaries characterize the evidence base and were not used to redefine the locked prediction target or models."),
    ("Supplementary Text S2", "Categorical composition and literature-source coverage",
     "The processed data contain three fluoride-form levels, 25 interferent levels, and 53 coagulant levels. Nineteen interferent levels and 46 coagulant levels occur in only one literature source, accounting for 307 and 535 records, respectively (Table S4). This long-tailed categorical structure was retained rather than collapsed after model evaluation.",
     "Source-level coverage is summarized in Table S5 and Figures S3–S6. The anonymized Paper_### labels represent provenance groups only; they were used for grouped validation, LOLO testing, and applicability-domain auditing, never as chemical predictors."),
    ("Supplementary Text S3", "Model configurations, hyperparameter spaces, and selection evidence",
     "Tables S6–S8 document the exact estimator settings, historical search spaces, selected configurations, and ensemble weights recovered from the original analysis record. The reconstruction is read-only: no estimator was refitted, no search was rerun, and no performance result was rewritten.",
     "Random-split and source-isolated validation settings are reported separately because they answer different questions. Selection records identify validation RMSE as the decision criterion, while test outcomes remain evaluation evidence rather than selection inputs."),
    ("Supplementary Text S4", "Complete random-split results and stability",
     "Across five prespecified seeds, the weighted four-model ensemble achieved mean R² = 0.882 ± 0.010, RMSE = 10.249, and MAE = 6.404 (Table S9; Figure S7). Its five-seed pooled performance was R² = 0.882, RMSE = 10.267, and MAE = 6.404 (Table S10; Figure S8). These results quantify internal learning ability under random partitioning.",
     "Candidate screening and numerical consistency checks are summarized in Table S11. Pooled records represent repeated prediction events across seed-specific test sets and therefore are not treated as independent experimental observations. Complete seed metrics and prediction events are provided as machine-readable files."),
    ("Supplementary Text S5", "Grouped validation and leave-one-literature-out generalization",
     "Source-isolated grouped validation and LOLO testing produced markedly weaker performance than random splitting (Tables S12–S14; Figures S9–S11), revealing substantial between-literature heterogeneity. For the pooled grouped audit, XGBoost achieved R² = 0.183, RMSE = 26.631, and MAE = 21.031 across 1,011 unique observations; this strategy-specific value is distinct from the unified audit baseline reported elsewhere in the manuscript.",
     "Table S15 verifies zero source overlap and summarizes recovered ensemble weights and numerical checks. Complete grouped-validation and LOLO prediction records are supplied separately so that source-specific errors and pooled metrics can be independently reproduced."),
    ("Supplementary Text S6", "Selective applicability-domain evidence and sensitivity",
     "At the exact P08 operating point, the retained prediction events yielded R² = 0.748, RMSE = 11.475, and MAE = 9.156, but coverage was only 52 events representing 44 unique observations (Table S16; Figures S12–S14). All retained events originated from Paper_010, and only one of two retained splits achieved R² ≥ 0.5.",
     "Threshold, split, bootstrap, transfer, chemistry-window, and representativeness audits are summarized in Tables S17–S18 and Figure S15. Because the rule is target-independent but the retained support is narrow and source-concentrated, it is interpreted as a candidate source-specific reliability window rather than a universal cross-literature applicability domain."),
    ("Supplementary Text S7", "Extended SHAP and partial-dependence interpretation",
     "SHAP importance was compared across the full dataset, Paper_010, and the retained selective domain to diagnose how model attribution changes with context (Table S19; Figure S16). Source identity defines audit subsets only and is not supplied to the prediction model.",
     "Partial-dependence surfaces and marginal profiles are accompanied by empirical-support diagnostics (Table S20; Figures S17–S18). Low-support regions are explicitly flagged, and all SHAP and PDP patterns are interpreted as descriptions of the fitted model rather than causal effects or prescriptive chemical optima."),
]


MACHINE_FILES = [
    ("Random-split complete seed results", "SI_random_split_complete_results_v1_seed_metrics.csv", "Table S9"),
    ("Random-split pooled prediction events", "SI_random_split_complete_results_v1_recorded_prediction_events.csv", "Table S11"),
    ("Grouped-validation and LOLO complete predictions", "SI_grouped_validation_lolo_complete_results_v1_original_predictions.csv", "Tables S12–S15"),
    ("AD threshold sensitivity", "SI_ad_threshold_coverage_sensitivity_audit_v1_threshold_sensitivity.csv", "Table S17"),
    ("P08 retained prediction events", "SI_ad_threshold_coverage_sensitivity_audit_v1_retained_prediction_events.csv", "Table S16"),
    ("SHAP importance across contexts", "SI_shap_pdp_extended_interpretation_v1_feature_context_metrics.csv", "Table S19"),
    ("PDP marginal response curves", "SI_shap_pdp_extended_interpretation_v1_pdp_marginal_profiles.csv", "Table S20"),
]


def make_machine_manifest(paths: Paths) -> pd.DataFrame:
    rows = []
    for description, filename, crossref in MACHINE_FILES:
        p = require(paths.tables / filename, f"selected machine-readable file '{filename}'")
        df = pd.read_csv(p)
        rows.append({"description": description, "filename": filename, "SI_cross_reference": crossref,
                     "rows": len(df), "columns": len(df.columns), "sha256": sha256(p),
                     "submission_role": "Separate machine-readable supplementary data"})
    out = pd.DataFrame(rows)
    out.to_csv(paths.attachment_manifest, index=False, encoding="utf-8-sig")
    return out


def build_document(paths: Paths, article_title: str, figure_manifest: pd.DataFrame,
                   table_frames: dict[str, list[tuple[str, pd.DataFrame, str]]],
                   machine_manifest: pd.DataFrame) -> Document:
    doc = Document()
    configure_styles(doc)
    apply_section(doc.sections[0], False)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("Supplementary Information")
    set_run_font(r, size=20, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run(article_title)
    set_run_font(r, size=15, bold=True)
    # Deliberately blank author area for later completion.
    doc.add_paragraph("\n\n")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Journal of Environmental Chemical Engineering")
    set_run_font(r, size=11, italic=True)
    doc.add_page_break()

    h = doc.add_heading("Contents", level=1)
    set_repeatable_heading(h)
    for num, title, _, _ in TEXTS:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Cm(0.4)
        p.add_run(f"{num}. {title}")
    p = doc.add_paragraph("Supplementary Tables S1–S20", style="Normal")
    p.paragraph_format.left_indent = Cm(0.4)
    p = doc.add_paragraph("Supplementary Figures S1–S18", style="Normal")
    p.paragraph_format.left_indent = Cm(0.4)
    p = doc.add_paragraph("Machine-readable supplementary data files", style="Normal")
    p.paragraph_format.left_indent = Cm(0.4)
    doc.add_page_break()

    doc.add_heading("Supplementary Text", level=1)
    for num, title, p1, p2 in TEXTS:
        h = doc.add_heading(f"{num}. {title}", level=2)
        set_repeatable_heading(h)
        for text in (p1, p2):
            p = doc.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_heading("Machine-readable supplementary data files", level=1)
    p = doc.add_paragraph(
        "The seven files below are selected for separate submission because their complete row-level or grid-level records are not efficiently represented in a formatted Word table. The files are read-only outputs; their original fields, labels, values, and filenames are unchanged."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    display = clean_frame(machine_manifest, ["description", "filename", "SI_cross_reference", "rows", "columns"])
    add_dataframe(doc, display, font_size=8)

    for i in range(1, 21):
        number = f"Table S{i}"
        parts = table_frames[number]
        max_cols = max(len(frame.columns) for _, frame, _ in parts)
        landscape = max_cols >= 9
        new_table_section(doc, landscape)
        if i == 1:
            doc.add_heading("Supplementary Tables", level=1)
        add_caption(doc, number, TABLE_TITLES[number])
        for j, (subtitle, frame, note) in enumerate(parts):
            if subtitle:
                h = doc.add_heading(subtitle, level=3)
                set_repeatable_heading(h)
            size = 6.2 if len(frame.columns) >= 12 else 6.8 if len(frame.columns) >= 9 else 7.5
            add_dataframe(doc, frame, font_size=size)
            add_note(doc, note)
            if j < len(parts) - 1:
                doc.add_paragraph()

    # Return to portrait for all figures.
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_section(section, False)

    pending_small: list[tuple[pd.Series, Path, float]] = []
    figures_started = False

    def flush_small() -> None:
        nonlocal pending_small, figures_started
        if not pending_small:
            return
        if not figures_started:
            doc.add_heading("Supplementary Figures", level=1)
            figures_started = True
        else:
            doc.add_page_break()
        if len(pending_small) == 1:
            row, img, _ = pending_small[0]
            add_one_figure(doc, row, img, 15.8)
        else:
            table = doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, (row, img, _) in zip(table.rows[0].cells, pending_small):
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(img), width=Cm(7.65))
                cp = cell.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = cp.add_run(f"{row['number']}. {row['title']}. ")
                set_run_font(r, size=8, bold=True)
                r = cp.add_run(str(row["caption"]))
                set_run_font(r, size=8)
        pending_small = []

    def add_one_figure(document: Document, row: pd.Series, img: Path, width_cm: float) -> None:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(img), width=Cm(width_cm))
        cp = document.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = cp.add_run(f"{row['number']}. {row['title']}. ")
        set_run_font(r, size=9, bold=True)
        r = cp.add_run(str(row["caption"]))
        set_run_font(r, size=9)

    for _, row in figure_manifest.iterrows():
        img = require(paths.figures / f"{row['file_stem']}.png", f"{row['number']} PNG")
        with Image.open(img) as im:
            aspect = im.width / im.height
            pixels = im.width * im.height
        # A figure is paired only when it remains legible at half-page width.
        # Wide or high-density figures are always assigned a dedicated page.
        small = pixels <= 8_000_000 and 0.8 <= aspect <= 1.6
        if small:
            pending_small.append((row, img, aspect))
            if len(pending_small) == 2:
                flush_small()
        else:
            flush_small()
            if not figures_started:
                doc.add_heading("Supplementary Figures", level=1)
                figures_started = True
            else:
                doc.add_page_break()
            add_one_figure(doc, row, img, 15.8)
    flush_small()

    add_headers_footers(doc)
    core = doc.core_properties
    core.author = ""
    core.last_modified_by = ""
    core.title = f"Supplementary Information — {article_title}"
    core.subject = "Supplementary Information for submission to the Journal of Environmental Chemical Engineering"
    core.keywords = "coagulation; defluoridation; machine learning; applicability domain; grouped validation"
    core.comments = "Generated from locked SI evidence files; author information intentionally left blank."
    return doc


def main() -> None:
    start = time.perf_counter()
    paths = project_paths()
    # Import the project configuration as an additional structure check.
    sys.path.insert(0, str(paths.config.parent))
    __import__("config")

    figure_manifest = read_csv(paths.tables, "SI_final_numbering_caption_manifest_v1_figure_manifest.csv")
    table_manifest = read_csv(paths.tables, "SI_final_numbering_caption_manifest_v1_table_manifest.csv")
    text_manifest = read_csv(paths.tables, "SI_final_numbering_caption_manifest_v1_text_manifest.csv")
    if list(text_manifest["number"]) != [f"Supplementary Text S{i}" for i in range(1, 8)]:
        raise AssertionError("Supplementary Text numbering is not continuous from S1 to S7.")
    if list(table_manifest["number"]) != [f"Table S{i}" for i in range(1, 21)]:
        raise AssertionError("Table numbering is not continuous from S1 to S20.")
    if list(figure_manifest["number"]) != [f"Figure S{i}" for i in range(1, 19)]:
        raise AssertionError("Figure numbering is not continuous from S1 to S18.")
    if not bool(table_manifest["all_files_present"].all()):
        raise AssertionError("At least one numbered-table evidence component is missing.")
    if not bool(figure_manifest[["png_present", "pdf_present"]].all().all()):
        raise AssertionError("At least one numbered figure is missing a PNG or PDF source.")

    protected = sorted({paths.config, paths.manuscript, *paths.figures.glob("*.png"),
                        *paths.figures.glob("*.pdf"), *paths.tables.glob("*.csv")})
    protected = [p for p in protected if p not in {paths.attachment_manifest}]
    before = {str(p.relative_to(paths.root)): sha256(p) for p in protected if p.is_file()}

    article_title = extract_english_title(paths.manuscript)
    machine_manifest = make_machine_manifest(paths)
    table_frames = build_table_frames(paths.tables)
    doc = build_document(paths, article_title, figure_manifest, table_frames, machine_manifest)
    doc.save(paths.output_docx)
    require(paths.output_docx, "generated SI Word file")

    after = {str(p.relative_to(paths.root)): sha256(p) for p in protected if p.is_file()}
    changed = [k for k in before if before[k] != after.get(k)]
    if changed:
        raise AssertionError("Protected evidence changed during Word assembly: " + ", ".join(changed))

    audit = {
        "script": STEM,
        "article_title": article_title,
        "manuscript_source": str(paths.manuscript.relative_to(paths.root)) if paths.manuscript.is_relative_to(paths.root) else str(paths.manuscript),
        "output_docx": str(paths.output_docx.relative_to(paths.root)),
        "supplementary_text_count": 7,
        "table_count": 20,
        "figure_count": 18,
        "machine_readable_file_count": 7,
        "author_information": "intentionally blank",
        "layout": "A4 portrait; automatic A4 landscape sections for wide tables; dedicated pages for wide/high-density figures; small figures may be paired",
        "protected_file_count": len(before),
        "protected_files_changed": changed,
        "output_sha256": sha256(paths.output_docx),
        "runtime_seconds": round(time.perf_counter() - start, 3),
    }
    paths.audit.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    log_lines = [
        "PASS: submission-ready SI Word assembled from locked evidence.",
        f"Title: {article_title}",
        "Supplementary Text: 7; Tables: 20; Figures: 18; machine-readable files: 7.",
        f"Output: {paths.output_docx}",
        f"Protected evidence files unchanged: {len(changed) == 0}.",
        f"Total runtime: {audit['runtime_seconds']:.2f} s",
    ]
    paths.log.write_text("\n".join(log_lines) + "\n", encoding="utf-8")
    print("\n".join(log_lines))


if __name__ == "__main__":
    try:
        main()
    except Exception:
        elapsed = time.perf_counter() - RUN_START
        print(f"Total runtime: {elapsed:.2f} s")
        raise
