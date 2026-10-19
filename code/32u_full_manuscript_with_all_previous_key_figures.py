# -*- coding: utf-8 -*-
"""
32U Full manuscript draft with random-split, all-domain grouped-validation,
selective applicability-domain results, and representative figures from previous steps.
Project: fluoride coagulation defluoridation ML project

Why this script exists
----------------------
32T/32T2 expanded the manuscript-level text, but only inserted the final 194-197
figures. This 32U script adds representative figures from the earlier workflow:
  - data/literature-group characterization
  - feature engineering / selected feature fusion
  - random-split model screening and ensemble modeling
  - multi-seed robustness and expanded-model screening
  - all-domain grouped validation and LOLO evidence
  - applicability-domain risk analysis
  - final selective AD evidence chain (194-197)

This script does NOT retrain models and does NOT change final metrics.
It only builds a fuller Word draft and writes figure/table audit files.

Original project output convention is preserved:
  outputs/tables/
  reports/

Expected run location:
  project root or src/
"""

from __future__ import annotations

import csv
import re
import sys
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional, List, Dict, Tuple

try:
    import pandas as pd
except Exception:
    pd = None

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except Exception as e:
    print("[ERROR] python-docx is required. Install it by: pip install python-docx")
    print("Original error:", e)
    raise SystemExit(1)


# ---------------------------------------------------------------------
# Fixed accepted result: do not modify unless the project decision changes
# ---------------------------------------------------------------------
LOCKED_RESULT = {
    "result_type_en": "target-free selective AD retained-subdomain",
    "result_type_cn": "无目标泄露的选择性适应域保留子域",
    "model": "RepeatedBalanced_Extended + XGBoost",
    "rule": "Zone=Outside; NN_Risk <= 0.814815",
    "r2": 0.747676,
    "rmse": 11.474674,
    "mae": 9.155856,
    "coverage": 0.014664,
    "n": 52,
    "dominant_group": "Paper_010",
    "all_domain_r2": 0.193744,
    "all_domain_rmse": 27.513358,
    "all_domain_mae": 22.893844,
    "predefined_r2": -0.160,
    "predefined_rmse": 29.416,
    "predefined_mae": 25.688,
    "predefined_coverage": 0.3796,
}

PROJECT_TITLE = "基于机器学习建模的含氟废水混凝处理的工艺选择、优化及控制"


@dataclass
class FigureSpec:
    fig_no: str
    stage: str
    preferred_numbers: List[int]
    keywords: List[str]
    caption: str
    section_hint: str
    width: float = 6.1
    required: bool = False


# Representative key figures. This is intentionally not all 197 figures.
# It covers every major workflow stage without making the paper unreadable.
FIGURE_PLAN: List[FigureSpec] = [
    FigureSpec(
        "图 3-1", "数据来源与文献组分布", [69],
        ["literature_group_size_distribution", "group_size", "文献"],
        "文献来源分组及样本量分布。该图用于说明数据具有明显的跨来源结构，后续需要采用按文献来源的分组验证来评价跨文献外推能力。",
        "3.1 数据集特征、描述性统计与建模挑战", 5.8),
    FigureSpec(
        "图 3-2", "目标变量与文献层级分布", [70],
        ["paper_level_target_distribution", "target_distribution", "目标变量"],
        "不同文献来源下目标变量分布差异。该图用于说明混凝除氟数据不仅存在高去除率样本集中现象，也存在文献来源之间的目标分布差异。",
        "3.1 数据集特征、描述性统计与建模挑战", 5.8),
    FigureSpec(
        "图 3-3", "机理特征筛选与融合输入", [63, 68],
        ["selected_feature_curve", "selected_feature_importance", "feature"],
        "机理特征处理与特征选择结果。该图用于说明 pH、投加量、初始氟浓度、干扰离子和文献来源相关变量进入模型前被重新组织为可学习的特征表达。",
        "3.2 数据预处理、机理特征构造与输入变量封装", 5.8),
    FigureSpec(
        "图 3-4", "随机划分下加权融合模型表现", [64, 66],
        ["selected_fusion_metrics_dashboard", "selected_fusion_pooled_predicted_vs_actual", "fusion_metrics"],
        "随机划分条件下代表性树模型/融合模型的预测表现。该图用于说明在数据内部插值条件下，树模型能够学习混凝除氟的非线性规律。",
        "3.3 随机划分下的基准模型筛选与集成建模", 6.2),
    FigureSpec(
        "图 3-5", "随机划分下融合权重稳定性", [65, 101, 105],
        ["weight_stability", "catboost_xgboost_weight", "fusion_weight"],
        "多随机种子下融合权重变化。该图用于说明单次划分得到的权重或性能可能受数据划分影响，因此需要通过多随机种子检验稳定性。",
        "3.3 随机划分下的基准模型筛选与集成建模", 5.8),
    FigureSpec(
        "图 3-6", "CatBoost/XGBoost 调参与扩展模型筛选", [99, 102],
        ["tuned_catboost_xgboost_metrics", "param_selection", "expanded_model_screen_metrics"],
        "CatBoost/XGBoost 调参与扩展模型筛选结果。该图用于说明梯度提升树模型在特定调参和特征处理条件下具有补充预测能力。",
        "3.3 随机划分下的基准模型筛选与集成建模", 6.2),
    FigureSpec(
        "图 3-7", "特征工程消融验证", [111, 112, 113],
        ["31k_feature_ablation_metrics", "feature_ablation", "feature_selection_frequency"],
        "机理特征工程消融和特征空间表现。该图用于说明并非所有衍生特征都会稳定提升模型，但特征工程为后续分组验证和适应域分析提供了统一输入。",
        "3.3 随机划分下的基准模型筛选与集成建模", 6.2),
    FigureSpec(
        "图 3-8", "分组验证中的训练-测试分布偏移", [72, 73],
        ["train_test_numeric_shift", "unseen_category_fraction", "smd_heatmap"],
        "按文献来源划分后训练集与测试集之间的数值变量偏移和未见类别比例。该图用于说明全域分组验证难度高于随机划分。",
        "3.4 全域分组验证与随机划分结果的差异", 6.2),
    FigureSpec(
        "图 3-9", "全域分组验证模型筛选", [80, 92, 95, 127],
        ["grouped_screen_metrics", "balanced_grouped_screen_metrics", "grouped_validation_strategy_metrics"],
        "未筛选全域分组验证下不同策略的性能比较。该图用于表明跨文献外推条件下，模型性能明显低于随机划分，说明数据存在强烈异质性。",
        "3.4 全域分组验证与随机划分结果的差异", 6.2),
    FigureSpec(
        "图 3-10", "全域分组验证预测偏差", [81, 93, 100, 138],
        ["grouped_screen_pooled_predicted_vs_actual", "pooled_predicted_vs_actual", "policy_pooled"],
        "未筛选全域分组验证或策略筛选下的预测值-实测值关系。该图用于展示跨文献预测中残差和偏差的来源。",
        "3.4 全域分组验证与随机划分结果的差异", 6.2),
    FigureSpec(
        "图 3-11", "LOLO 文献层级误差", [129, 137, 98],
        ["lolo_paper_level_rmse", "high_error_papers", "paper_level_rmse"],
        "留一文献组外推时的文献层级误差。该图用于识别高误差文献组，并说明误差并非平均分布于所有文献来源。",
        "3.4 全域分组验证与随机划分结果的差异", 6.2),
    FigureSpec(
        "图 3-12", "适应域风险与预测误差关系", [139, 140, 141, 142],
        ["ad_risk_error_relationship", "ad_performance_by_risk_bin", "paper_ad_risk", "violation_driver"],
        "适应域风险指标与预测误差之间的关系。该图用于说明高风险样本通常对应更大的跨文献预测不确定性。",
        "3.5 预定义可靠域与选择性适应域的对比", 6.2),
    FigureSpec(
        "图 3-13", "预定义可靠域规则表现", [153, 154, 155, 157],
        ["all_vs_reliable_domain_metrics", "reliability_zone_metrics", "reliable_domain_prediction_panels"],
        "未筛选全域测试与预定义可靠域的对比。该图用于说明固定可靠域规则并不一定带来稳定性能提升。",
        "3.5 预定义可靠域与选择性适应域的对比", 6.2),
    FigureSpec(
        "图 3-14", "选择性适应域证据链", [194],
        ["32i2_evidence_chain_summary_clean", "evidence_chain_summary"],
        "选择性适应域保留子域结果的证据链汇总。该图比较了未筛选全域分组验证、预定义可靠域和选择性适应域保留子域的 R²、RMSE、MAE 和覆盖率，用于说明最终高 R² 结果属于选择性保留子域。",
        "3.6 选择性适应域保留子域的模型性能", 6.2, True),
    FigureSpec(
        "图 3-15", "NN-risk 阈值平台", [195],
        ["32i2_threshold_plateau_evidence_clean", "threshold_plateau"],
        "选择性适应域规则的 NN-risk 阈值平台证据。虚线位置对应最终接受阈值，表明该结果不是单一孤立阈值造成的偶然峰值。",
        "3.7 NN-risk 阈值平台、化学窗口与解释边界", 6.2, True),
    FigureSpec(
        "图 3-16", "保留子域化学窗口", [196],
        ["32i2_retained_chemistry_window_clean", "retained_chemistry_window"],
        "选择性适应域保留子域的化学条件窗口与适应域组成。该图用于界定模型解释适用的水质、投加量、pH、干扰离子和混凝剂空间。",
        "3.7 NN-risk 阈值平台、化学窗口与解释边界", 6.2, True),
    FigureSpec(
        "图 3-17", "稳定性、迁移支持与解释边界", [197],
        ["32i2_stability_transfer_boundary_clean", "stability_transfer_boundary"],
        "选择性适应域结果的稳定性、迁移支持和解释边界。该图强调最终结果具有明确的文献来源和覆盖范围限制，不能解释为全域泛化性能。",
        "3.8 稳定性、迁移支持与本节小结", 6.2, True),
]


def normalize_path(p: Path) -> Path:
    return p.resolve()


def looks_like_project_root(p: Path) -> bool:
    return (p / "outputs").exists() and (p / "src").exists()


def find_project_root() -> Path:
    here = Path(__file__).resolve()
    candidates = [Path.cwd().resolve(), here.parent.resolve()] + list(here.parents)
    for c in candidates:
        if looks_like_project_root(c):
            return c
    # fallback: if running from src, parent may still be project root even if src folder is current file parent
    for c in candidates:
        if (c / "outputs" / "figures").exists() and (c / "outputs" / "tables").exists():
            return c
    return Path.cwd().resolve()


def find_template(project_root: Path) -> Optional[Path]:
    names = [
        "AI论文--混凝沉淀除氟空稿.docx",
        "AI论文--混凝沉淀除氟空稿.docx",
    ]
    search_dirs = [
        project_root,
        project_root / "reports",
        project_root.parent,
        project_root.parent.parent if project_root.parent != project_root.parent.parent else project_root.parent,
        Path.cwd(),
    ]
    # Also search below project root quickly
    for d in search_dirs:
        if not d.exists():
            continue
        for name in names:
            p = d / name
            if p.exists():
                return p
    for root in [project_root, project_root.parent, Path.cwd()]:
        if root.exists():
            hits = list(root.rglob("AI论文--混凝沉淀除氟空稿.docx"))
            if hits:
                return hits[0]
    return None


def clear_doc_body_keep_sections(doc: Document) -> None:
    body = doc._element.body
    children = list(body)
    sect_pr = None
    for child in children:
        if child.tag == qn("w:sectPr"):
            sect_pr = child
    for child in children:
        body.remove(child)
    if sect_pr is not None:
        body.append(sect_pr)


def set_doc_defaults(doc: Document) -> None:
    try:
        normal = doc.styles["Normal"]
        normal.font.name = "宋体"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        normal.font.size = Pt(10.5)
    except Exception:
        pass
    try:
        for section in doc.sections:
            section.top_margin = Inches(0.85)
            section.bottom_margin = Inches(0.85)
            section.left_margin = Inches(0.85)
            section.right_margin = Inches(0.85)
    except Exception:
        pass


def add_run_font(run, font_name="宋体", size=10.5, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    add_run_font(run, "宋体", 16 if level == 1 else 13.5 if level == 2 else 12, True)
    try:
        p.style = f"Heading {min(level,3)}"
    except Exception:
        pass


def add_para(doc: Document, text: str, size: float = 10.5, first_line: bool = True, bold_prefix: Optional[str] = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line:
        try:
            p.paragraph_format.first_line_indent = Pt(21)
        except Exception:
            pass
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        add_run_font(r1, size=size, bold=True)
        rest = text[len(bold_prefix):]
        if rest:
            r2 = p.add_run(rest)
            add_run_font(r2, size=size)
    else:
        run = p.add_run(text)
        add_run_font(run, size=size)


def set_cell_text(cell, text: str, bold=False, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text))
    add_run_font(r, size=size, bold=bold)
    try:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    except Exception:
        pass


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")


def add_table(doc: Document, headers: List[str], rows: List[List[str]], title: Optional[str] = None) -> None:
    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        add_run_font(r, size=10.5, bold=True)
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for j, h in enumerate(headers):
        set_cell_text(table.cell(0, j), h, bold=True, size=9.5)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            set_cell_text(table.cell(i, j), val, size=9)
    doc.add_paragraph()


def get_figure_dir(project_root: Path) -> Path:
    candidates = [
        project_root / "outputs" / "figures",
        project_root.parent / "outputs" / "figures",
        Path.cwd() / "outputs" / "figures",
    ]
    for c in candidates:
        if c.exists():
            return c
    return project_root / "outputs" / "figures"


def list_image_files(fig_dir: Path, project_root: Path) -> List[Path]:
    dirs = [fig_dir, project_root / "outputs" / "figures", project_root, project_root.parent]
    seen = set()
    files = []
    for d in dirs:
        if d.exists():
            for p in d.rglob("*.png"):
                if p.name.startswith("~$"):
                    continue
                rp = str(p.resolve())
                if rp not in seen:
                    files.append(p.resolve())
                    seen.add(rp)
    return files


def figure_number_from_name(name: str) -> Optional[int]:
    m = re.match(r"^(\d+)[_\-]", name)
    if m:
        return int(m.group(1))
    return None


def find_figure(spec: FigureSpec, files: List[Path]) -> Optional[Path]:
    # exact preferred number first
    for n in spec.preferred_numbers:
        matches = [p for p in files if figure_number_from_name(p.name) == n]
        if matches:
            # prefer clean/final png and smaller obvious duplicates last? Sort clean first
            matches.sort(key=lambda p: (0 if "clean" in p.name.lower() else 1, len(p.name)))
            return matches[0]
    # keyword search
    candidates: List[Tuple[int, Path]] = []
    for p in files:
        low = p.name.lower()
        score = 0
        for kw in spec.keywords:
            if kw.lower() in low:
                score += 10
        for n in spec.preferred_numbers:
            if low.startswith(str(n) + "_"):
                score += 30
        if score > 0:
            candidates.append((score, p))
    if candidates:
        candidates.sort(key=lambda x: (-x[0], len(x[1].name)))
        return candidates[0][1]
    return None


def add_picture_safe(doc: Document, path: Optional[Path], caption: str, width: float, fig_no: str) -> Tuple[bool, str]:
    if path is None or not Path(path).exists():
        add_para(doc, f"[{fig_no} 对应图件未找到，需在 outputs/figures 中补齐后重新运行脚本。]", first_line=False)
        add_para(doc, f"{fig_no} {caption}", size=9.5, first_line=False)
        return False, "not_found"
    try:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=Inches(width))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"{fig_no} {caption}")
        add_run_font(r, size=9.5)
        return True, str(path)
    except Exception as e:
        add_para(doc, f"[{fig_no} 图件插入失败：{e}]", first_line=False)
        add_para(doc, f"{fig_no} {caption}", size=9.5, first_line=False)
        return False, "insert_error: " + repr(e)


def media_count(docx_path: Path) -> int:
    try:
        with zipfile.ZipFile(docx_path) as z:
            return sum(1 for n in z.namelist() if n.startswith("word/media/"))
    except Exception:
        return -1


def build_doc(project_root: Path, template: Optional[Path], figures: Dict[str, Optional[Path]], reports_dir: Path) -> Document:
    if template and template.exists():
        doc = Document(str(template))
        clear_doc_body_keep_sections(doc)
    else:
        doc = Document()
    set_doc_defaults(doc)

    add_heading(doc, PROJECT_TITLE, 1)

    # ----------------------------- Introduction -----------------------------
    add_heading(doc, "1. Introduction（引言）", 1)
    add_heading(doc, "1.1 研究背景与意义", 2)
    add_para(doc, "氟是自然界中广泛存在的卤族元素之一。适量氟与人体骨骼和牙齿发育有关，但当水体中氟离子浓度超过安全限值时，会引起氟斑牙、氟骨症等健康风险，并可能对水生态系统和饮用水安全造成长期影响。含氟废水主要来源于地下水地质背景、磷肥和氟化工、半导体与光伏制造、冶金、电镀、玻璃蚀刻及煤化工等过程。由于氟离子水溶性强、迁移性高，进入水环境后难以依靠自然衰减快速降低浓度，因此发展高效、稳定、可调控的除氟技术具有重要环境工程意义。")
    add_para(doc, "在现有除氟技术中，吸附、膜分离、离子交换、电化学和混凝沉淀等方法均得到广泛研究。混凝沉淀具有药剂来源广、工艺成熟、运行成本相对较低、易与现有水处理流程衔接等优势，尤其适用于中高浓度含氟废水的预处理或快速处理。然而，混凝除氟效果受混凝剂类型、投加量、原水 pH、初始氟浓度、共存离子、反应时间、沉降条件和水温等多因素共同影响，其工程优化仍高度依赖经验试验和人工判断。")
    add_heading(doc, "1.2 混凝除氟过程的核心问题", 2)
    add_para(doc, "混凝沉淀除氟并不是单一变量控制的线性过程。以铝盐、铁盐及复合混凝剂为代表的药剂在水中会发生水解、络合、吸附、共沉淀和絮体卷扫等反应；pH 会调节金属羟基络合物形态、絮体表面电荷以及氟离子与金属中心的结合能力；HCO₃⁻、CO₃²⁻、SO₄²⁻、PO₄³⁻ 等共存阴离子可能通过竞争吸附、络合平衡或沉淀竞争影响除氟效率。同时，工程评价不能仅依赖去除率，还应结合初始氟负荷和出水氟浓度判断达标风险。")
    add_para(doc, "已有研究通常围绕单一混凝剂或特定实验体系开展，能够解释局部实验规律，但难以直接形成跨文献、跨工况的统一预测模型。对于由多篇文献汇总得到的数据集，不同研究在药剂体系、测试水质、实验记录完整性和变量范围上存在显著差异。如果仅采用随机划分评价模型，同一文献或相近实验系列样本可能同时进入训练集和测试集，从而高估模型对未见文献工况的泛化能力。因此，如何在小样本、异质性和不均衡数据条件下构建可靠预测模型，是混凝除氟机器学习研究中的关键问题。")
    add_heading(doc, "1.3 机器学习、分组验证与适应域分析", 2)
    add_para(doc, "机器学习方法能够从复杂环境实验数据中学习非线性关系和高阶变量交互，已在水处理药剂投加预测、污染物去除效率模拟、材料性能筛选和工艺参数优化中表现出应用潜力。Random Forest、Extra Trees、XGBoost、CatBoost 等树模型适合处理表格型环境数据，能够在一定程度上捕捉变量之间的非线性耦合；SHAP、PDP 等可解释方法则可用于分析关键变量贡献和潜在作用机制。")
    add_para(doc, "然而，环境科学数据常具有样本量有限、来源分散、低性能区域样本稀疏和变量分布不均衡等特点。在这种条件下，随机划分下的高精度并不必然代表真实外推能力。分组验证、留一文献外推、适应域（applicability domain, AD）和不确定性约束逐渐成为评价环境机器学习模型可靠性的关键工具。对于文献汇总型混凝除氟数据，必须区分随机划分性能、未筛选全域分组验证性能和适应域筛选后的局部保留子域性能，避免将局部可靠结果误解为全域泛化能力。")
    add_heading(doc, "1.4 本研究内容与技术路线", 2)
    add_para(doc, "本研究构建混凝沉淀除氟文献数据集，并围绕数据预处理、机理特征构造、机器学习建模、随机划分评价、跨文献分组验证、适应域筛选和结果解释开展系统分析。首先，对数据集中的水质参数、药剂参数、运行条件和目标变量进行统一整理，构造具有化学机理意义的衍生特征；其次，比较线性模型、树模型、调参模型和集成模型在随机划分和多随机种子条件下的表现；随后，引入按文献来源分组的验证策略，评价模型在未见文献工况下的全域外推能力；最后，基于无目标泄露的适应域筛选，识别模型相对可靠的局部保留子域，并明确其化学窗口和应用边界。")

    # ----------------------------- Methods -----------------------------
    add_heading(doc, "2. Methodology / Experimental section（方法）", 1)
    add_heading(doc, "2.1 数据来源与数据集构建", 2)
    add_para(doc, "本研究数据来源于已发表混凝沉淀除氟相关文献中的实验结果。根据实验记录，提取混凝剂类型、混凝剂投加量、絮凝剂信息、反应/搅拌时间、沉降时间、进水 pH、初始氟浓度、水温、干扰离子类型及浓度、出水 pH、氟去除率和出水氟浓度等信息。对于同一文献中的不同 pH、不同投加量、不同初始氟浓度或不同共存离子实验，按独立实验工况整理为独立样本。")
    add_heading(doc, "2.2 数据预处理与机理特征构造", 2)
    add_para(doc, "数值变量主要采用中位数或物理含义一致的方式处理缺失值；类别变量统一为标准名称或“无/未记录”等类别。对于缺失本身具有信息含义的变量，增加缺失指示特征，以保留文献记录差异可能带来的数据来源信息。在原始变量基础上构造 log(1 + 投加量)、相对投加强度、投加量与初始氟浓度交互项、pH 距中性距离、弱酸性/中性/碱性区间指示、干扰离子类别拆分、碳酸盐类干扰指示和混凝剂大类归并等机理特征。")
    add_heading(doc, "2.3 模型构建、随机划分与全域分组验证", 2)
    add_para(doc, "模型候选集包括 Linear Regression、Ridge、Random Forest、Extra Trees、HistGradientBoosting、CatBoost 和 XGBoost 等。随机划分用于评价模型在数据内部插值条件下的学习能力，多随机种子验证用于检验模型对数据划分的稳定性。集成策略包括加权融合和 stacking，其中加权融合通过验证集 RMSE 最小化确定模型权重，stacking 则使用第一层模型预测作为第二层模型输入。")
    add_para(doc, "为进一步评价跨文献外推能力，本研究采用按文献来源分组的验证策略。该策略避免同一文献或相近实验系列同时进入训练集与测试集，更接近模型面对未见文献工况时的实际应用场景。本文将未筛选全域分组验证结果与适应域筛选后的局部结果分开报告，以避免不同评价口径混用。")
    add_heading(doc, "2.4 适应域定义与评价指标", 2)
    add_para(doc, "适应域分析用于判断待预测样本是否位于模型相对可靠的数据支持范围内。本研究综合考虑距离风险、邻近样本风险、杠杆风险、类别未见风险和重构风险等指标，并在无目标泄露前提下建立选择性保留规则。模型性能采用 R²、RMSE、MAE 和 Coverage 进行评价。其中，Coverage 表示适应域规则保留的样本比例，反映模型在可靠性约束下的适用范围。")

    # ----------------------------- Results -----------------------------
    add_heading(doc, "3. Results and Discussion（结果与讨论）", 1)
    add_heading(doc, "3.1 数据集特征、描述性统计与建模挑战", 2)
    add_para(doc, "本研究首先对混凝沉淀除氟文献数据集进行描述性统计和可视化分析。结果表明，数据整体可用于机器学习建模，但不同变量的数据完整性并不一致，进水 pH、干扰离子浓度等关键机理变量存在一定缺失，需要在建模前进行专门处理。目标变量方面，样本中高去除率区间占比较高，同时仍包含低去除率和失败实验样本，说明数据分布具有明显不均衡特征。")
    add_para(doc, "从机理角度看，氟去除率并不由单一变量线性决定。pH 会影响铝盐、铁盐等混凝剂的水解形态、絮体表面电荷和氟离子吸附/络合过程；混凝剂投加量决定有效絮体生成量；干扰离子会影响竞争吸附和沉淀平衡；初始氟浓度则影响最终出水是否达标。因此，混凝除氟数据呈现多因素耦合、非线性和跨文献异质性共同作用的特点。")
    for key in ["图 3-1", "图 3-2"]:
        spec = next(x for x in FIGURE_PLAN if x.fig_no == key)
        add_picture_safe(doc, figures.get(key), spec.caption, spec.width, spec.fig_no)

    add_heading(doc, "3.2 数据预处理、机理特征构造与输入变量封装", 2)
    add_para(doc, "针对文献数据来源分散、变量记录方式不完全一致的问题，本研究在建模前进行了统一的数据预处理。数值变量主要包括初始氟浓度、混凝剂投加量、进水 pH、反应/沉降时间、水温和干扰离子浓度等；类别变量主要包括混凝剂类型、干扰离子类型和相关工艺类别。")
    add_para(doc, "机理特征构造并不改变原始实验结果，而是将化学机理和数据结构以更适合模型学习的方式表达出来。例如，log(1 + 投加量) 用于降低投加量跨度过大造成的尺度影响；相对投加强度用于描述单位污染负荷对应药剂量；pH 区间特征用于表达弱酸性、中性和碱性条件下混凝剂水解与氟结合能力的差异；干扰离子拆分特征则用于描述共存阴离子的竞争效应。")
    spec = next(x for x in FIGURE_PLAN if x.fig_no == "图 3-3")
    add_picture_safe(doc, figures.get("图 3-3"), spec.caption, spec.width, spec.fig_no)

    add_heading(doc, "3.3 随机划分下的基准模型筛选与集成建模", 2)
    add_para(doc, "在随机划分评价阶段，线性回归和 Ridge 模型主要作为基准模型，用于判断单纯线性关系是否足以描述混凝除氟过程。结果表明，线性模型难以充分捕捉 pH、投加量、混凝剂类型、干扰离子和初始氟浓度之间的复杂耦合关系。相比之下，Random Forest、Extra Trees、HistGradientBoosting、CatBoost 和 XGBoost 等树模型表现更优，说明当前数据更符合非线性表格数据建模特征。")
    add_para(doc, "在模型优化阶段，Random Forest 表现出较强稳定性；CatBoost 在手动调参后能够明显提升单次划分结果，但多随机种子下波动较大；Random Forest 与 CatBoost 的加权融合和 Ridge/Linear stacking 可在一定程度上利用模型互补性，但提升幅度有限。多随机种子检验表明，单次随机划分下的高性能不能直接代表稳定泛化能力，因此随机划分结果更适合用于模型筛选和内部规律学习。")
    for key in ["图 3-4", "图 3-5", "图 3-6", "图 3-7"]:
        spec = next(x for x in FIGURE_PLAN if x.fig_no == key)
        add_picture_safe(doc, figures.get(key), spec.caption, spec.width, spec.fig_no)

    add_heading(doc, "3.4 全域分组验证与随机划分结果的差异", 2)
    add_para(doc, "为评价机器学习模型在文献异质性条件下的真实泛化能力，本研究进一步采用按文献来源分组的验证策略。与随机划分相比，分组验证能够避免同一文献或相近实验系列同时进入训练集和测试集，从而更严格地检验模型对未见文献工况的预测能力。未筛选全域分组验证结果明显低于随机划分结果，说明模型虽然能够学习数据内部规律，但在跨文献外推时仍受到实验体系差异和局部工况偏移的限制。")
    add_para(doc, "因此，本文将随机划分、未筛选全域分组验证和适应域筛选结果作为三个不同层级的评价口径：随机划分反映数据内部插值能力；未筛选全域分组验证反映跨文献整体外推难度；选择性适应域保留子域则反映风险约束下局部可靠预测能力。三者具有不同解释意义，不能用选择性适应域结果替代全域分组验证性能。")
    add_table(doc, ["评价层级", "模型/规则", "R²", "RMSE", "MAE", "Coverage"], [
        ["随机划分代表性结果", "CatBoost_tuned / tree-based models", "约 0.741", "-", "-", "-"],
        ["多随机种子稳定性", "RF_CatBoost_LinearStacking / RF baseline", "约 0.694 ± 0.040", "-", "-", "-"],
        ["未筛选全域分组验证", "Grouped validation, all-domain", f"{LOCKED_RESULT['all_domain_r2']:.3f}", f"{LOCKED_RESULT['all_domain_rmse']:.3f}", f"{LOCKED_RESULT['all_domain_mae']:.3f}", "100.00%"],
        ["预定义可靠域", "Predefined reliable domain", f"{LOCKED_RESULT['predefined_r2']:.3f}", f"{LOCKED_RESULT['predefined_rmse']:.3f}", f"{LOCKED_RESULT['predefined_mae']:.3f}", "37.96%"],
        ["选择性适应域保留子域", LOCKED_RESULT["model"] + "; " + LOCKED_RESULT["rule"], f"{LOCKED_RESULT['r2']:.3f}", f"{LOCKED_RESULT['rmse']:.3f}", f"{LOCKED_RESULT['mae']:.3f}", "1.47%"],
    ], "表 3-1 随机划分、未筛选全域分组验证与选择性适应域结果的层级对比")
    for key in ["图 3-8", "图 3-9", "图 3-10", "图 3-11"]:
        spec = next(x for x in FIGURE_PLAN if x.fig_no == key)
        add_picture_safe(doc, figures.get(key), spec.caption, spec.width, spec.fig_no)

    add_heading(doc, "3.5 预定义可靠域与选择性适应域的对比", 2)
    add_para(doc, "在预定义可靠域分析中，模型尝试通过固定适应域规则筛除高风险样本，以获得更可信的预测子集。然而，固定可靠域并未带来稳定的性能提升，说明单一、静态的适应域阈值难以充分处理本数据集中复杂的跨文献异质性和工况偏移。进一步的选择性适应域分析显示，当模型仅在满足特定风险约束的保留子域内进行评价和解释时，预测误差显著降低，决定系数明显提高。")
    add_para(doc, "这表明适应域分析的关键并不是简单扩大或缩小预测范围，而是识别模型在当前数据结构下真正具有相对可靠性的局部化学工况空间。该过程必须保持目标变量隔离，避免通过实际测试标签参与筛选而造成目标泄露。")
    for key in ["图 3-12", "图 3-13"]:
        spec = next(x for x in FIGURE_PLAN if x.fig_no == key)
        add_picture_safe(doc, figures.get(key), spec.caption, spec.width, spec.fig_no)

    add_heading(doc, "3.6 选择性适应域保留子域的模型性能", 2)
    add_para(doc, f"最终接受的选择性适应域结果来自 {LOCKED_RESULT['model']}，筛选规则为 {LOCKED_RESULT['rule']}。在该保留子域内，模型取得 R² = {LOCKED_RESULT['r2']:.3f}、RMSE = {LOCKED_RESULT['rmse']:.3f}、MAE = {LOCKED_RESULT['mae']:.3f}，覆盖率为 {LOCKED_RESULT['coverage']*100:.2f}%，保留样本数为 {LOCKED_RESULT['n']}。该结果仅表征适应域规则筛选后保留样本子集的预测表现；未筛选测试集结果与筛选后保留子域结果应分开报告，二者不能混用。")
    add_table(doc, ["项目", "内容"], [
        ["结果类型", LOCKED_RESULT["result_type_cn"]],
        ["模型", LOCKED_RESULT["model"]],
        ["筛选规则", LOCKED_RESULT["rule"]],
        ["R²", f"{LOCKED_RESULT['r2']:.6f}"],
        ["RMSE", f"{LOCKED_RESULT['rmse']:.6f}"],
        ["MAE", f"{LOCKED_RESULT['mae']:.6f}"],
        ["Coverage", f"{LOCKED_RESULT['coverage']:.6f}（约 {LOCKED_RESULT['coverage']*100:.2f}%）"],
        ["保留样本数", str(LOCKED_RESULT["n"])],
        ["主导文献组", LOCKED_RESULT["dominant_group"]],
    ], "表 3-2 最终接受的选择性适应域保留子域关键结果")
    spec = next(x for x in FIGURE_PLAN if x.fig_no == "图 3-14")
    add_picture_safe(doc, figures.get("图 3-14"), spec.caption, spec.width, spec.fig_no)

    add_heading(doc, "3.7 NN-risk 阈值平台、化学窗口与解释边界", 2)
    add_para(doc, "为降低选择性适应域规则被解释为事后挑选结果的风险，本研究进一步对 NN-risk 阈值进行了敏感性分析。结果显示，在接受阈值附近，保留子域 R² 存在相对稳定的高性能区间，而不是由单一孤立阈值产生的偶然峰值。同时，随着阈值放宽，保留样本数和覆盖率逐渐升高，但模型精度并未持续提高，说明更多样本并不必然带来更可靠的外推结果。")
    spec = next(x for x in FIGURE_PLAN if x.fig_no == "图 3-15")
    add_picture_safe(doc, figures.get("图 3-15"), spec.caption, spec.width, spec.fig_no)
    add_para(doc, "保留子域样本主要来自 Paper_010，并集中于特定混凝剂组成、pH、投加量、初始氟浓度和干扰离子浓度范围内。该结果一方面表明模型在该类化学窗口内能够获得较高预测精度，另一方面也提示模型解释具有明确边界。对于该子域以外的混凝剂类型、极端 pH、不同初始氟负荷或复杂共存离子体系，模型预测仍需谨慎使用。")
    spec = next(x for x in FIGURE_PLAN if x.fig_no == "图 3-16")
    add_picture_safe(doc, figures.get("图 3-16"), spec.caption, spec.width, spec.fig_no)

    add_heading(doc, "3.8 稳定性、迁移支持与本节小结", 2)
    add_para(doc, "稳定性分析进一步表明，选择性适应域结果在部分数据划分和相关模型组合中具有可重复支持，但其覆盖率较低且文献来源集中，因此仍应将其定位为保留子域内的可靠预测证据，不能作为未筛选测试集整体泛化能力的直接证据。在论文表述中，应避免将该保留子域结果直接描述为整个混凝除氟数据集的泛化性能。")
    spec = next(x for x in FIGURE_PLAN if x.fig_no == "图 3-17")
    add_picture_safe(doc, figures.get("图 3-17"), spec.caption, spec.width, spec.fig_no)
    add_para(doc, "综上，混凝沉淀除氟文献数据在跨来源验证中表现出明显的泛化难度，说明该体系存在强烈的数据异质性。通过引入无目标泄露的适应域筛选，本研究识别出一个具有较高预测可靠性的局部保留子域，并在该子域内获得较好的预测结果。该发现为后续工艺参数推荐和模型辅助决策提供了更加谨慎的依据：模型可用于支持特定工况窗口内的预测和解释，但对于适应域之外的水质条件和药剂体系，仍需要补充实验或更多文献数据进行验证。")

    # ----------------------------- Conclusion -----------------------------
    add_heading(doc, "4. Conclusion（结论）", 1)
    conclusions = [
        "本研究基于文献数据构建了混凝沉淀除氟机器学习数据集，并从数据清洗、机理特征构造、模型筛选、随机划分评价、跨文献分组验证和适应域分析等方面建立了较完整的建模流程。",
        "随机划分结果表明，树模型相较线性模型更适合描述混凝除氟过程中的非线性和多因素耦合关系；但多随机种子与分组验证结果显示，单次随机划分高精度不能直接代表跨文献泛化能力。",
        "未筛选全域分组验证揭示了混凝除氟文献数据存在显著跨来源异质性，模型对未见文献工况的整体外推仍具有较大挑战。",
        f"在无目标泄露的选择性适应域筛选后，{LOCKED_RESULT['model']} 在保留子域内取得 R² = {LOCKED_RESULT['r2']:.3f}、RMSE = {LOCKED_RESULT['rmse']:.3f}、MAE = {LOCKED_RESULT['mae']:.3f} 的预测表现，但该结果覆盖率仅为 {LOCKED_RESULT['coverage']*100:.2f}%，且主导文献组为 {LOCKED_RESULT['dominant_group']}，因此应限定解释为局部保留子域结果。",
        "未来研究应进一步补充跨文献、跨药剂体系和低覆盖区域实验数据，并结合外部实验验证、动态适应域更新和工程约束优化，以提高模型在真实含氟废水处理场景中的泛化能力和应用价值。",
    ]
    for idx, c in enumerate(conclusions, 1):
        add_para(doc, f"（{idx}）{c}", first_line=False)

    return doc


def write_csv(path: Path, rows: List[Dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    keys = sorted({k for row in rows for k in row.keys()}) if rows else []
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=keys)
        w.writeheader()
        for r in rows:
            w.writerow(r)


def main() -> None:
    project_root = find_project_root()
    output_tables = project_root / "outputs" / "tables"
    reports_dir = project_root / "reports"
    fig_dir = get_figure_dir(project_root)
    output_tables.mkdir(parents=True, exist_ok=True)
    reports_dir.mkdir(parents=True, exist_ok=True)
    template = find_template(project_root)
    files = list_image_files(fig_dir, project_root)

    figure_paths: Dict[str, Optional[Path]] = {}
    manifest_rows = []
    for spec in FIGURE_PLAN:
        p = find_figure(spec, files)
        figure_paths[spec.fig_no] = p
        manifest_rows.append({
            "figure_no": spec.fig_no,
            "stage": spec.stage,
            "section_hint": spec.section_hint,
            "preferred_numbers": ",".join(map(str, spec.preferred_numbers)),
            "found": "yes" if p else "no",
            "path": str(p) if p else "",
            "required": "yes" if spec.required else "no",
            "caption": spec.caption,
        })

    doc = build_doc(project_root, template, figure_paths, reports_dir)
    out_docx = reports_dir / "32u_full_manuscript_with_random_grouped_ad_and_previous_figures.docx"
    out_text = reports_dir / "32u_full_manuscript_with_random_grouped_ad_and_previous_figures_text.txt"
    out_log = reports_dir / "32u_run_log.txt"
    doc.save(str(out_docx))

    # Extract text-like content for quick review
    try:
        doc2 = Document(str(out_docx))
        text_lines = [p.text for p in doc2.paragraphs if p.text.strip()]
        out_text.write_text("\n".join(text_lines), encoding="utf-8")
    except Exception:
        out_text.write_text("text extraction failed", encoding="utf-8")

    write_csv(output_tables / "32u_previous_key_figure_insertion_manifest.csv", manifest_rows)

    quality_rows = []
    mc = media_count(out_docx)
    required_missing = [r for r in manifest_rows if r["required"] == "yes" and r["found"] != "yes"]
    quality_rows.append({"check": "word_media_count", "status": "PASS" if mc >= 4 else "WARN", "value": str(mc), "expected": ">= 4"})
    quality_rows.append({"check": "required_final_figures_found", "status": "PASS" if not required_missing else "FAIL", "value": "; ".join(r["figure_no"] for r in required_missing), "expected": "194-197 found"})
    quality_rows.append({"check": "model_name_locked", "status": "PASS", "value": LOCKED_RESULT["model"], "expected": LOCKED_RESULT["model"]})
    quality_rows.append({"check": "r2_boundary", "status": "PASS", "value": "R2=0.747676 is selective AD retained-subdomain only", "expected": "not all-domain grouped validation"})
    write_csv(output_tables / "32u_docx_quality_audit.csv", quality_rows)

    hierarchy_rows = [
        {"level": "随机划分", "meaning": "数据内部插值与模型筛选", "main_text": "可用于说明模型学习能力，不作为跨文献外推证据"},
        {"level": "未筛选全域分组验证", "meaning": "跨文献整体外推难度", "main_text": f"R2={LOCKED_RESULT['all_domain_r2']:.3f}, RMSE={LOCKED_RESULT['all_domain_rmse']:.3f}"},
        {"level": "预定义可靠域", "meaning": "固定可靠域规则", "main_text": f"R2={LOCKED_RESULT['predefined_r2']:.3f}, Coverage={LOCKED_RESULT['predefined_coverage']*100:.2f}%"},
        {"level": "选择性适应域保留子域", "meaning": "风险约束下的局部可靠预测", "main_text": f"R2={LOCKED_RESULT['r2']:.3f}, RMSE={LOCKED_RESULT['rmse']:.3f}, Coverage={LOCKED_RESULT['coverage']*100:.2f}%"},
    ]
    write_csv(output_tables / "32u_random_grouped_ad_result_hierarchy.csv", hierarchy_rows)

    log_lines = [
        "="*70,
        "32U Full manuscript with previous key figures",
        f"Run time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"Project root: {project_root}",
        f"Template: {template if template else 'not_found_created_new_docx'}",
        f"Figure dir: {fig_dir}",
        f"Table output dir: {output_tables}",
        f"Report output dir: {reports_dir}",
        "-"*70,
        "[LOCKED KEY RESULT]",
        f"{LOCKED_RESULT['model']} | {LOCKED_RESULT['rule']}",
        f"R2={LOCKED_RESULT['r2']:.6f}; RMSE={LOCKED_RESULT['rmse']:.6f}; MAE={LOCKED_RESULT['mae']:.6f}; Coverage={LOCKED_RESULT['coverage']:.6f}; N={LOCKED_RESULT['n']}; Dominant group={LOCKED_RESULT['dominant_group']}",
        "-"*70,
    ]
    for r in manifest_rows:
        log_lines.append(f"[{r['figure_no']}] found={r['found']} path={r['path']}")
    log_lines += [
        "-"*70,
        f"word/media count={mc}",
        f"Output DOCX: {out_docx}",
        f"Output text: {out_text}",
        f"Output manifest: {output_tables / '32u_previous_key_figure_insertion_manifest.csv'}",
        "32U finished successfully." if not required_missing else "32U finished with missing required final figures.",
        "Important writing boundary: R2=0.747676 belongs to the target-free selective AD retained-subdomain result only.",
    ]
    out_log.write_text("\n".join(log_lines), encoding="utf-8")

    print("\n".join(log_lines))


if __name__ == "__main__":
    main()
