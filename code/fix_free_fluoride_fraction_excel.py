#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""只修正最终 Excel 的“游离氟占比”，Python 3.9+，依赖 openpyxl。
放在项目 src/ 后：
    python src/fix_free_fluoride_fraction_excel.py
    python src/fix_free_fluoride_fraction_excel.py --apply
也可从任意位置指定正式数据：
    python fix_free_fluoride_fraction_excel.py --file "D:\项目\data\raw\数据集修改最终.xlsx"
预览不写文件；--apply 才正式修改。所有报告写入数据文件旁的 free_fluoride_audit。
绝不调用模型脚本，不读写 model_ready CSV，不写 outputs。
依据：29c 的数字文献编号 -> Paper_XXX 映射；34a 整组 (1,100] 的尺度规则；
已确认的 11 组、234 条，以及项目镜像中核实的各组条数。
"""
import argparse
from collections import Counter
from datetime import datetime
from decimal import Decimal
import hashlib
import io
import json
import os
from pathlib import Path
import re
import sys
import tempfile
import zipfile

try:
    import openpyxl
except ImportError:
    raise SystemExit("缺少依赖，请在项目 Python 环境运行：python -m pip install openpyxl")

COUNTS = {8: 12, 19: 2, 20: 47, 22: 6, 23: 9, 24: 11,
          25: 16, 26: 18, 28: 37, 30: 58, 32: 18}
SHEET_XML = "xl/worksheets/sheet1.xml"


def require(ok, message):
    if not ok:
        raise RuntimeError(message)


def sha(data):
    return hashlib.sha256(data).hexdigest()


def workbook(data):
    return openpyxl.load_workbook(io.BytesIO(data), data_only=False)


def make_plan(data):
    wb = workbook(data)
    try:
        require(wb.sheetnames == ["Sheet1"], "工作表结构与已核实版本不同，停止。")
        ws = wb["Sheet1"]
        require((ws.max_row, ws.max_column) == (1013, 14),
                "应为 1011 条数据、2 行表头、14 列；当前结构不同，停止。")
        require(ws["B2"].value == "游离氟占比" and ws["N2"].value == "文献编号",
                "列名或列位置与已核实版本不同，停止。")
        groups, changes, original_values = Counter(), [], []
        for row in range(3, 1014):
            group, value = ws.cell(row, 14).value, ws.cell(row, 2).value
            require(type(group) in (int, float) and group == int(group),
                    f"N{row} 不是已确认的数字文献编号，停止。")
            group = int(group)
            groups[group] += 1
            require(type(value) in (int, float), f"B{row} 不是普通数值，停止。")
            number = Decimal(str(value))
            require(number.is_finite(), f"B{row} 不是有限数，停止。")
            original_values.append(number)
            if group in COUNTS:
                require(1 < number <= 100,
                        f"B{row} 不在 (1,100]，可能已修正或数据版本不同；停止，不能重复除以 100。")
                changes.append({
                    "cell": f"B{row}", "excel_row": row,
                    "source_group": f"Paper_{group:03d}",
                    "before": str(number), "after": str(number / 100),
                })
            else:
                require(0 <= number <= 1, f"非目标 B{row} 不在 [0,1]，需要人工核对。")
        require(len(groups) == 27, "来源组数不是 27，停止。")
        require({g: groups[g] for g in COUNTS} == COUNTS,
                "目标各组条数与已确认版本不同，停止。")
        require(len(changes) == 234, "修改单元格数不是 234，停止。")
        new_values = [v / 100 if v > 1 else v for v in original_values]
        ranges = {
            "before_min": str(min(original_values)), "before_max": str(max(original_values)),
            "after_min": str(min(new_values)), "after_max": str(max(new_values)),
        }
        return changes, ranges
    finally:
        wb.close()


def patch_xml(xml, changes):
    # 只替换目标 <v> 数字文本，保留属性、样式、公式及其他字节。
    for item in changes:
        cell = item["cell"].encode("ascii")
        pattern = rb'(<c\b[^>]*\br="' + cell + rb'"[^>]*>)(.*?)(</c>)'
        matches = list(re.finditer(pattern, xml, re.DOTALL))
        require(len(matches) == 1, f"无法唯一定位 {item['cell']}，停止。")
        match = matches[0]
        body = match.group(2)
        require(b"<f" not in body, "目标单元格包含公式，停止。")
        values = list(re.finditer(rb"(<v>)([^<]*)(</v>)", body))
        require(len(values) == 1, "目标不是单一数值节点，停止。")
        val = values[0]
        # XML 可能保存 IEEE 浮点展开文本；按 Excel 读取后的数值核对。
        require(float(val.group(2)) == float(item["before"]), "XML 与读取数值不符，停止。")
        start = match.start(2) + val.start(2)
        end = match.start(2) + val.end(2)
        xml = xml[:start] + item["after"].encode("ascii") + xml[end:]
    return xml


def build_candidate(before, changes):
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(before)) as src, zipfile.ZipFile(buffer, "w") as dst:
        names = src.namelist()
        require(SHEET_XML in names, "缺少已确认的工作表 XML，停止。")
        require(len(names) == len(set(names)), "XLSX 内含重复成员，停止。")
        dst.comment = src.comment
        for info in src.infolist():
            content = src.read(info.filename)
            dst.writestr(info, patch_xml(content, changes) if info.filename == SHEET_XML else content)
    return buffer.getvalue()


def verify(before, after, changes):
    with zipfile.ZipFile(io.BytesIO(before)) as a, zipfile.ZipFile(io.BytesIO(after)) as b:
        require(a.namelist() == b.namelist(), "XLSX 成员列表变化，停止。")
        require(a.comment == b.comment, "ZIP 注释变化，停止。")
        for name in a.namelist():
            expected = patch_xml(a.read(name), changes) if name == SHEET_XML else a.read(name)
            require(expected == b.read(name), f"非授权工作簿内容变化：{name}")
    wa, wb = workbook(before), workbook(after)
    try:
        expected = {c["cell"]: Decimal(c["after"]) for c in changes}
        changed = set()
        for row in wa["Sheet1"]:
            for old in row:
                new = wb["Sheet1"][old.coordinate]
                require(old.data_type == new.data_type and old._style == new._style,
                        f"{old.coordinate} 类型或格式变化。")
                if old.coordinate in expected:
                    require(Decimal(str(new.value)) == expected[old.coordinate], "修正数值不符。")
                    require(old.value != new.value, "目标单元格未发生变化。")
                    changed.add(old.coordinate)
                else:
                    require(old.value == new.value, f"{old.coordinate} 非目标单元格变化。")
        require(changed == set(expected) and len(changed) == 234, "实际修改集合不符。")
    finally:
        wa.close()
        wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--file", type=Path, help="正式 Excel 的完整路径")
    parser.add_argument("--apply", action="store_true", help="校验后替换正式 Excel；默认仅预览")
    args = parser.parse_args()
    path = (args.file or Path(__file__).resolve().parent.parent /
            "data/raw/数据集修改最终.xlsx").resolve()
    require(path.name == "数据集修改最终.xlsx", "文件名必须是 数据集修改最终.xlsx。")
    require(path.parent.name != "sources", "sources 是只读项目参考材料，请指定正式项目原文件。")
    require(path.is_file(), f"找不到正式数据集：{path}\n请通过 --file 指定正确完整路径。")
    backup = path.with_name("数据集修改最终_游离氟修正前备份.xlsx")
    before = path.read_bytes()
    changes, ranges = make_plan(before)
    after = build_candidate(before, changes)
    verify(before, after, changes)
    print(f"文件：{path}\n1011 条记录，27 个来源组；目标 11 组、234 个单元格。")
    for c in changes:
        print(f"{c['cell']}  {c['source_group']}  {c['before']} -> {c['after']}")
    stats = []
    for group, count in COUNTS.items():
        subset = [c for c in changes if c["source_group"] == f"Paper_{group:03d}"]
        stat = {
            "source_group": f"Paper_{group:03d}", "modified_rows": count,
            "before_min": str(min(Decimal(c["before"]) for c in subset)),
            "before_max": str(max(Decimal(c["before"]) for c in subset)),
            "after_min": str(min(Decimal(c["after"]) for c in subset)),
            "after_max": str(max(Decimal(c["after"]) for c in subset)),
        }
        stats.append(stat)
        print(f"{stat['source_group']}: {count} 条；"
              f"{stat['before_min']}–{stat['before_max']} -> {stat['after_min']}–{stat['after_max']}")
    print(f"总体范围：{ranges['before_min']}–{ranges['before_max']} -> "
          f"{ranges['after_min']}–{ranges['after_max']}")
    print("内存候选校验通过：只改 234 个目标数值，其他单元格及工作簿内容不变。")
    if not args.apply:
        print("预览结束，未写入任何文件。正式修改请添加 --apply。")
        return
    if backup.exists():
        require(backup.read_bytes() == before,
                "已有备份与当前原文件不一致，停止且不覆盖备份；请核对数据版本。")
    else:
        with backup.open("xb") as f:
            f.write(before)
            f.flush()
            os.fsync(f.fileno())
    require(backup.read_bytes() == before, "备份完整性校验失败。")
    audit = path.parent / "free_fluoride_audit"
    audit.mkdir(exist_ok=True)
    report = audit / (datetime.now().strftime("%Y%m%d_%H%M%S_%f") + ".json")
    record = {
        "status": "validated_before_replace", "file": str(path), "backup": str(backup),
        "before_sha256": sha(before), "after_sha256": sha(after),
        "records": 1011, "source_groups": 27, "modified_cells": 234,
        "other_cells_unchanged": True, "other_xlsx_contents_unchanged": True,
        "ranges": ranges, "target_group_statistics": stats, "changes": changes,
    }
    report.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary = None
    replaced = False
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as f:
            temporary = Path(f.name)
            f.write(after)
            f.flush()
            os.fsync(f.fileno())
        verify(before, temporary.read_bytes(), changes)
        require(path.read_bytes() == before, "执行期间原文件被其他程序修改，停止。")
        require(backup.read_bytes() == before, "执行期间备份发生变化，停止。")
        os.replace(temporary, path)
        temporary = None
        replaced = True
        verify(before, path.read_bytes(), changes)
        record["status"] = "completed_and_verified"
        report.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        if replaced:
            print("注意：正式文件已经替换，后续校验或报告写入失败；请保留备份并核查，勿重复执行。",
                  file=sys.stderr)
        raise
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    print(f"成功：234 个目标数值已修正，其他内容校验无变化。\n备份：{backup}\n完整报告：{report}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"停止：{exc}", file=sys.stderr)
        sys.exit(1)

